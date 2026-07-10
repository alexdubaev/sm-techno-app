from __future__ import annotations

from datetime import date
from io import BytesIO
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl.cell.rich_text import CellRichText
from openpyxl import Workbook, load_workbook

import stock_sync_api
from stock_sync_web.commercial_offers import (
    CommercialOfferLineInput,
    generate_commercial_offer_workbook,
    read_source_offer_lines,
)
from stock_sync_web.database import WebDatabase
from stock_sync_web.service import WebStockSyncService


class CommercialOfferExcelTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self._temp_dir.name)

    def tearDown(self) -> None:
        self._temp_dir.cleanup()

    def test_source_reader_supports_aliases_russian_numbers_and_skips_empty_rows(self) -> None:
        source_path = self.temp_path / "КП-125 ООО Ромашка.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Позиции"
        sheet.append(["служебная строка"])
        sheet.append(["Код", "Название", "Марка", "Кол-во", "Цена", "Срок", "Комментарий"])
        sheet.append(["CAT-001", "Фильтр", "Caterpillar", "2", "12 500,50", "2-3 дня", "Оригинал"])
        sheet.append([None, None, None, None, None, None, None])
        sheet.append(["JCB-002", "Насос", "JCB", 1, 7000, "5 дней", ""])
        workbook.save(source_path)

        lines = read_source_offer_lines(source_path)

        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0].row_no, 1)
        self.assertEqual(lines[0].article, "CAT-001")
        self.assertEqual(lines[0].name, "Фильтр")
        self.assertEqual(lines[0].brand, "Caterpillar")
        self.assertEqual(lines[0].qty, 2)
        self.assertEqual(lines[0].price_vat, 12500.50)
        self.assertEqual(lines[0].amount_vat, 25001.0)
        self.assertEqual(lines[1].article, "JCB-002")

    def test_generator_writes_offer_rows_and_total_formula_without_extra_sheets(self) -> None:
        template_path = self.temp_path / "template.xlsx"
        output_path = self.temp_path / "out.xlsx"
        create_minimal_template(template_path)

        lines = [
            CommercialOfferLineInput(
                row_no=1,
                article="A-1",
                name="Фильтр",
                brand="CAT",
                qty=2,
                price_vat=1000,
                amount_vat=2000,
                delivery_time="2 дня",
                note="",
                item_id=None,
                warehouse_id=None,
                warehouse_name="",
            ),
            CommercialOfferLineInput(
                row_no=2,
                article="B-2",
                name="Насос",
                brand="JCB",
                qty=3,
                price_vat=500,
                amount_vat=1500,
                delivery_time="5 дней",
                note="OEM",
                item_id=None,
                warehouse_id=None,
                warehouse_name="",
            ),
        ]

        generate_commercial_offer_workbook(
            template_path=template_path,
            output_path=output_path,
            lines=lines,
            offer_number="КП-125",
            client_name="ООО Ромашка",
            offer_date=date(2026, 7, 9),
        )

        workbook = load_workbook(output_path, data_only=False)
        self.assertEqual(workbook.sheetnames, ["КП"])
        sheet = workbook["КП"]
        self.assertEqual(sheet["A7"].value, "Коммерческое предложение № КП-125")
        self.assertEqual(sheet["A9"].value, "Покупатель: ООО Ромашка")
        self.assertEqual(sheet["F9"].value, "Дата: 09.07.2026")
        self.assertEqual(sheet["A13"].value, 1)
        self.assertEqual(sheet["B13"].value, "A-1")
        self.assertEqual(sheet["E13"].number_format, "0")
        self.assertEqual(sheet["G13"].value, '=IF(OR(E13="",F13=""),"",E13*F13)')
        self.assertEqual(sheet["A14"].value, 2)
        self.assertEqual(sheet["G15"].value, "=SUM(G13:G14)")

    def test_generator_moves_template_merged_footer_below_long_offer_rows(self) -> None:
        template_path = self.temp_path / "template_with_footer.xlsx"
        output_path = self.temp_path / "out_long.xlsx"
        create_minimal_template(template_path)
        workbook = load_workbook(template_path)
        sheet = workbook["КП"]
        sheet.merge_cells("E33:F33")
        sheet.merge_cells("A35:I35")
        sheet.merge_cells("A36:I36")
        sheet["A35"] = "Условия оплаты: тест"
        sheet["A36"] = "Условия доставки: тест"
        workbook.save(template_path)

        lines = [
            CommercialOfferLineInput(
                row_no=index,
                article=f"SKU-{index:02d}",
                name=f"Position {index}",
                brand="CAT",
                qty=index,
                price_vat=100,
                amount_vat=index * 100,
                delivery_time="2 дня",
                note="",
                item_id=None,
                warehouse_id=None,
                warehouse_name="",
            )
            for index in range(1, 26)
        ]

        generate_commercial_offer_workbook(
            template_path=template_path,
            output_path=output_path,
            lines=lines,
            offer_number="КП-126",
            client_name="ООО Ромашка",
            offer_date=date(2026, 7, 10),
        )

        workbook = load_workbook(output_path, data_only=False)
        sheet = workbook["КП"]
        merged_ranges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
        total_row = 13 + len(lines)

        self.assertEqual(sheet["B35"].value, "SKU-23")
        self.assertEqual(sheet[f"G{total_row}"].value, f"=SUM(G13:G{total_row - 1})")
        self.assertIn(f"E{total_row}:F{total_row}", merged_ranges)
        self.assertIn(f"A{total_row + 2}:I{total_row + 2}", merged_ranges)
        self.assertNotIn("A35:I35", merged_ranges)

    def test_generator_formats_long_names_date_and_client_name(self) -> None:
        template_path = self.temp_path / "template_formatting.xlsx"
        output_path = self.temp_path / "out_formatting.xlsx"
        create_minimal_template(template_path)
        workbook = load_workbook(template_path)
        sheet = workbook["КП"]
        sheet.column_dimensions["C"].width = 23
        sheet.row_dimensions[13].height = 17.1
        sheet.merge_cells("F9:I9")
        workbook.save(template_path)

        generate_commercial_offer_workbook(
            template_path=template_path,
            output_path=output_path,
            lines=[
                CommercialOfferLineInput(
                    row_no=1,
                    article="LONG-1",
                    name="VALVE GP-SOLENOID (24-VOLT)(COLD START APPL, HYDRAULIC FAN)",
                    brand="CAT",
                    qty=1,
                    price_vat=1000,
                    amount_vat=1000,
                    delivery_time="2 дня",
                    note="",
                    item_id=None,
                    warehouse_id=None,
                    warehouse_name="",
                )
            ],
            offer_number="КП-127",
            client_name="ООО Ромашка",
            offer_date=date(2026, 7, 10),
        )

        workbook = load_workbook(output_path, data_only=False, rich_text=True)
        sheet = workbook["КП"]
        client_cell_value = sheet["A9"].value

        self.assertEqual(sheet["C13"].alignment.wrap_text, True)
        self.assertEqual(sheet["C13"].alignment.vertical, "top")
        self.assertGreaterEqual(sheet.row_dimensions[13].height or 0, 34.0)
        self.assertEqual(sheet["F9"].alignment.horizontal, "right")
        self.assertEqual(str(client_cell_value), "Покупатель: ООО Ромашка")
        self.assertIsInstance(client_cell_value, CellRichText)
        self.assertEqual(client_cell_value[1].text, "ООО Ромашка")
        self.assertEqual(client_cell_value[1].font.b, True)


class CommercialOfferApiTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.db_path = Path(self._temp_dir.name) / "stock_sync.db"
        self.template_path = Path(self._temp_dir.name) / "commercial_offer_template.xlsx"
        create_minimal_template(self.template_path)

        self.db = WebDatabase(db_path=self.db_path)
        self.service = WebStockSyncService(
            db=self.db,
            commercial_offer_template_path=self.template_path,
            commercial_offer_storage_dir=Path(self._temp_dir.name) / "storage",
        )
        self.original_service = stock_sync_api.SERVICE
        stock_sync_api.SERVICE = self.service
        stock_sync_api.app.dependency_overrides[stock_sync_api._get_current_user] = lambda: {
            "id": 1,
            "role": "admin",
        }
        stock_sync_api.app.dependency_overrides[stock_sync_api._get_admin_user] = lambda: {
            "id": 1,
            "role": "admin",
        }
        self.client = TestClient(stock_sync_api.app)

    def tearDown(self) -> None:
        self.client.close()
        stock_sync_api.app.dependency_overrides.clear()
        stock_sync_api.SERVICE = self.original_service
        self._temp_dir.cleanup()

    def test_create_from_draft_lists_downloads_and_marks_sent_without_stock_side_effects(self) -> None:
        initial_item = self.db.create_local_item(
            sku="SKU-001",
            name="Фильтр масляный",
            print_name="Фильтр масляный",
            category_name="CAT",
            group_name="Фильтры",
            price=1250,
            warehouses=[{"warehouse_name": "Основной склад", "quantity": 7}],
        )
        initial_stock = self.service.get_item(int(initial_item["id"]))

        response = self.client.post(
            "/api/commercial-offers/from-draft",
            json={
                "clientName": "ООО Локальный клиент",
                "notes": "Тестовое КП",
                "lines": [
                    {
                        "itemId": initial_item["id"],
                        "article": "SKU-001",
                        "name": "Фильтр масляный",
                        "brand": "CAT",
                        "qty": 2,
                        "priceVat": 1250,
                        "deliveryTime": "2 дня",
                        "note": "Оригинал",
                        "warehouseId": initial_item["warehouses"][0]["warehouse_id"],
                        "warehouseName": "Основной склад",
                    }
                ],
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        offer_id = response.json()["offer"]["id"]

        list_response = self.client.get("/api/commercial-offers")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.json()["items"][0]["id"], offer_id)
        self.assertEqual(list_response.json()["items"][0]["clientName"], "ООО Локальный клиент")

        download_response = self.client.get(f"/api/commercial-offers/{offer_id}/download/output")
        self.assertEqual(download_response.status_code, 200)
        generated = load_workbook(BytesIO(download_response.content), data_only=False)
        self.assertEqual(generated["КП"]["C13"].value, "Фильтр масляный")

        sent_response = self.client.post(
            f"/api/commercial-offers/{offer_id}/mark-sent",
            json={"sentTo": "client@example.com", "notes": "Отправлено по email"},
        )
        self.assertEqual(sent_response.status_code, 200)
        self.assertEqual(sent_response.json()["offer"]["status"], "Отправлено")

        after_stock = self.service.get_item(int(initial_item["id"]))
        self.assertEqual(after_stock["quantity"], initial_stock["quantity"])
        self.assertEqual(self.db.list_orders(include_all=True), [])

    def test_create_from_excel_uses_filename_as_offer_number_and_keeps_source_download(self) -> None:
        source = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Позиции"
        sheet.append(["Артикул", "Наименование", "Бренд", "Количество", "Цена с НДС"])
        sheet.append(["CAT-777", "Ремкомплект", "CAT", 1, 2500])
        workbook.save(source)
        source.seek(0)

        response = self.client.post(
            "/api/commercial-offers/from-excel",
            data={"clientName": "ООО Excel клиент", "notes": ""},
            files={
                "file": (
                    "КП-777 ООО Excel клиент.xlsx",
                    source.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        offer = response.json()["offer"]
        self.assertEqual(offer["number"], "КП-777 ООО Excel клиент")
        self.assertEqual(offer["lineCount"], 1)
        self.assertEqual(offer["totalAmount"], 2500)

        source_response = self.client.get(f"/api/commercial-offers/{offer['id']}/download/source")
        self.assertEqual(source_response.status_code, 200)

    def test_admin_can_delete_commercial_offer_and_user_cannot(self) -> None:
        response = self.client.post(
            "/api/commercial-offers/from-draft",
            json={
                "clientName": "РћРћРћ РЈРґР°Р»РµРЅРёРµ",
                "lines": [
                    {
                        "article": "DEL-1",
                        "name": "РўРµСЃС‚РѕРІР°СЏ РїРѕР·РёС†РёСЏ",
                        "qty": 1,
                        "priceVat": 100,
                    }
                ],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        offer_id = response.json()["offer"]["id"]

        stock_sync_api.app.dependency_overrides[stock_sync_api._get_current_user] = lambda: {
            "id": 2,
            "role": "user",
        }

        forbidden_response = self.client.delete(f"/api/commercial-offers/{offer_id}")
        self.assertEqual(forbidden_response.status_code, 403)

        stock_sync_api.app.dependency_overrides[stock_sync_api._get_current_user] = lambda: {
            "id": 1,
            "role": "admin",
        }

        delete_response = self.client.delete(f"/api/commercial-offers/{offer_id}")
        self.assertEqual(delete_response.status_code, 200, delete_response.text)
        self.assertEqual(delete_response.json(), {"ok": True})

        list_response = self.client.get("/api/commercial-offers")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.json()["items"], [])

        details_response = self.client.get(f"/api/commercial-offers/{offer_id}")
        self.assertEqual(details_response.status_code, 404)


def create_minimal_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "КП"
    sheet["A7"] = "Коммерческое предложение №____"
    sheet["A9"] = "Покупатель: ______________________________________________"
    sheet["F9"] = "Дата: «___» __________ 20___ г."
    headers = [
        "№",
        "Артикул",
        "Наименование",
        "Бренд",
        "Количество",
        "Цена с НДС",
        "Сумма с НДС",
        "Срок поставки",
        "Примечание",
    ]
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=12, column=column_index, value=header)
    for row in range(13, 33):
        sheet.cell(row=row, column=1, value=row - 12)
        sheet.cell(row=row, column=7, value=f'=IF(OR(E{row}="",F{row}=""),"",E{row}*F{row})')
    sheet["E33"] = "Итого с НДС:"
    sheet["G33"] = "=SUM(G13:G32)"
    workbook.save(path)


if __name__ == "__main__":
    unittest.main()
