from __future__ import annotations

from io import BytesIO
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from stock_sync_web.database import WebDatabase
from stock_sync_web.service import WebStockSyncService


class ClientPriceExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory()
        self.db = WebDatabase(Path(self._temp_dir.name) / "stock_sync.db")
        self.service = WebStockSyncService(db=self.db)

    def tearDown(self) -> None:
        self._temp_dir.cleanup()

    def test_client_price_export_uses_warehouse_column_with_row_values(self) -> None:
        first_warehouse = self.db.create_warehouse(name="Склад A")
        second_warehouse = self.db.create_warehouse(name="Склад B")

        self.db.create_local_item(
            sku="SKU-001",
            name="Тестовая позиция",
            print_name="Тестовая позиция",
            category_name="John Deere",
            group_name="John Deere",
            price=1500,
            warehouses=[
                {
                    "warehouse_id": first_warehouse["id"],
                    "warehouse_name": first_warehouse["name"],
                    "quantity": 3,
                },
                {
                    "warehouse_id": second_warehouse["id"],
                    "warehouse_name": second_warehouse["name"],
                    "quantity": 5,
                },
            ],
        )

        content = self.service.export_client_price_bytes()
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook.active

        self.assertEqual(sheet.cell(1, 6).value, "Склад")
        self.assertEqual(sheet.cell(2, 6).value, "Склад A")
        self.assertEqual(sheet.cell(3, 6).value, "Склад B")
        self.assertEqual(sheet.cell(2, 4).value, 3)
        self.assertEqual(sheet.cell(3, 4).value, 5)


if __name__ == "__main__":
    unittest.main()
