from fastapi import FastAPI
from fastapi.testclient import TestClient
from stock_sync_web.stock_valuation import create_stock_valuation_router

class DB:
    def _normalize_stock_sku(self, value): return value.strip().upper()
    def lookup_stock_articles(self, articles):
        return {'001': [{'item_id': 1, 'warehouse_id': 2, 'warehouse_name': 'Склад', 'quantity': 4, 'price': 12}]}
class Service: db = DB()

def test_resolve_preserves_optional_quantity_and_warehouse():
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: Service()))
    response = TestClient(app).post('/api/stock/valuation/resolve', json={'rows': [{'article': '001', 'requestedQty': None, 'sourceRow': 1}]})
    assert response.status_code == 200
    row = response.json()['articleResolution'][0]
    assert row['article'] == '001' and row['requestedQty'] is None
    assert row['variants'][0]['physicalStock'] == 4
    assert response.headers['cache-control'] == 'no-store'

def test_workbook_matches_snapshot_and_neutralizes_formula():
    import base64
    from io import BytesIO
    from openpyxl import load_workbook
    class FormulaDB(DB):
        def lookup_stock_articles(self, articles):
            return {'=DANGEROUS()': [{'item_id': 1, 'warehouse_id': 2, 'warehouse_name': '=WAREHOUSE()', 'quantity': 5, 'price': 11},
                                    {'item_id': 1, 'warehouse_id': 3, 'warehouse_name': 'Второй', 'quantity': 2, 'price': 13}]}
    class FormulaService: db = FormulaDB()
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: FormulaService()))
    response = TestClient(app).post('/api/stock/valuation/resolve', json={'rows': [{'article': '=DANGEROUS()', 'requestedQty': None, 'sourceRow': 1}]})
    assert response.status_code == 200
    result = response.json()
    assert len(result['articleResolution'][0]['variants']) == 2
    workbook = load_workbook(BytesIO(base64.b64decode(result['workbook']['base64'])), data_only=False)
    cells = list(workbook.active.iter_rows())
    assert [cells[1][3].value, cells[2][3].value] == [5, 2]
    assert cells[1][0].value == '=DANGEROUS()' and cells[1][0].data_type == 's'
    assert cells[1][5].value == '=WAREHOUSE()' and cells[1][5].data_type == 's'
    assert workbook.active.max_row == 3


def test_inactive_actor_cannot_resolve():
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user', 'is_active': False}, lambda: Service()))
    response = TestClient(app).post('/api/stock/valuation/resolve', json={'rows': [{'article': '001', 'requestedQty': None, 'sourceRow': 1}]})
    assert response.status_code == 403

def test_intake_cleans_source_after_extractor_error(monkeypatch, tmp_path):
    import tempfile
    import stock_sync_web.stock_valuation as module
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))
    class BrokenExtractor:
        def extract(self, request, *, temp_root=None): raise RuntimeError('sensitive input')
    monkeypatch.setattr(module, 'CodexDocumentExtractor', BrokenExtractor)
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: Service()))
    response = TestClient(app).post('/api/stock/valuation/intake', files={'file': ('list.txt', b'001\n', 'text/plain')})
    assert response.status_code == 400
    assert 'sensitive input' not in response.text
    assert list((tmp_path/'sm-techno-stock-valuation-v1').iterdir()) == []

def test_real_database_lookup_has_no_business_writes(tmp_path):
    from stock_sync_web.database import WebDatabase
    from stock_sync_web.service import WebStockSyncService
    db = WebDatabase(tmp_path / 'synthetic.db')
    service = WebStockSyncService(db=db, commercial_offer_storage_dir=tmp_path/'offers', document_storage_dir=tmp_path/'documents')
    db.create_local_item(sku='00123', name='Синтетический товар', print_name='Синтетический товар', category_name='Тест', group_name='Тест', price=42,
                         warehouses=[{'warehouse_name':'Первый', 'quantity':2}, {'warehouse_name':'Второй', 'quantity':5}])
    with db.connect() as connection:
        tables = [r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        before = {table: list(map(tuple, connection.execute('SELECT * FROM "' + table + '"'))) for table in tables}
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: service))
    client = TestClient(app)
    response = client.post('/api/stock/valuation/resolve', json={'rows': [
        {'article': '00123', 'requestedQty': None, 'sourceRow': 1},
        {'article': '00123', 'requestedQty': None, 'sourceRow': 2},
        {'article': 'no-such-sku', 'requestedQty': None, 'sourceRow': 3}]})
    assert response.status_code == 200
    resolved = response.json()['articleResolution']
    assert [r['status'] for r in resolved] == ['matched', 'duplicate', 'not_found']
    assert sorted(v['physicalStock'] for v in resolved[0]['variants']) == [2, 5]
    with db.connect() as connection:
        after = {table: list(map(tuple, connection.execute('SELECT * FROM "' + table + '"'))) for table in tables}
    assert before == after
    assert not any(p.is_file() for p in (tmp_path/'offers').rglob('*')) if (tmp_path/'offers').exists() else True

def test_invalid_payload_is_redacted_and_not_cached():
    app = FastAPI()
    app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: Service()))
    response = TestClient(app).post('/api/stock/valuation/resolve', json={'rows': [{'article': 12345, 'requestedQty': None, 'sourceRow': 1}]})
    assert response.status_code == 422
    assert '12345' not in response.text
    assert response.headers['cache-control'] == 'no-store'


def test_scavenger_preserves_active_and_foreign_workspaces(monkeypatch, tmp_path):
    import os
    import tempfile
    import time
    import stock_sync_web.stock_valuation as module
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))
    with module._stock_workspace() as active:
        root = active.parent
        orphan = root / ('stock-' + 'a' * 32)
        orphan.mkdir(); (orphan / '.lease').write_bytes(b'sm-techno-stock-valuation-v1'); (orphan / 'source.txt').write_text('private')
        foreign = root / ('stock-' + 'b' * 32); foreign.mkdir(); (foreign / '.lease').write_bytes(b'foreign'); (foreign / 'keep').write_text('keep')
        old = time.time() - 1200
        os.utime(orphan, (old, old)); os.utime(active, (old, old)); os.utime(foreign, (old, old))
        module._scavenge(root)
        assert active.exists() and foreign.exists() and not orphan.exists()
    assert not active.exists() and foreign.exists()


def test_extractor_uses_supplied_stock_root(monkeypatch, tmp_path):
    import json
    import stock_sync_web.assistant_extraction as extraction
    monkeypatch.setattr(extraction, 'discover_codex_cli', lambda configured: tmp_path/'codex')
    image = tmp_path/'image.png'; image.write_bytes(b'synthetic')
    seen = {}
    def runner(args, prompt, env, workspace, timeout):
        seen['workspace'] = workspace
        seen['image'] = workspace/'page-1.png'
        assert seen['image'].read_bytes() == b'synthetic'
        result = {'payload': {'rows': [{'article':'001', 'requested_qty':None, 'source_row':1}]}, 'warnings':[], 'missing_fields':[]}
        return '\n'.join(map(json.dumps, [{'type':'item.completed', 'item':{'type':'agent_message', 'text':json.dumps(result)}}, {'type':'turn.completed'}]))
    root = tmp_path/'owned'; root.mkdir()
    result = extraction.CodexDocumentExtractor(runner=runner).extract(extraction.ExtractionRequest('articles', ['001'], [image], []), temp_root=root)
    assert result.payload['rows'][0]['article'] == '001'
    assert seen['workspace'].parent == root and not seen['workspace'].exists()

def test_intake_transient_result_and_owned_runtime(monkeypatch, tmp_path):
    import tempfile
    import stock_sync_web.stock_valuation as module
    from stock_sync_web.assistant_extraction import ExtractionResult
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))
    observed = {}
    class FakeExtractor:
        def extract(self, request, *, temp_root=None):
            observed['root'] = temp_root
            assert request.kind == 'articles' and request.texts == ['00123']
            assert (temp_root/'source.txt').read_text() == '00123'
            return ExtractionResult('articles', {'rows':[{'article':'00123', 'requested_qty':None, 'source_row':1}]})
    monkeypatch.setattr(module, 'CodexDocumentExtractor', FakeExtractor)
    app = FastAPI(); app.include_router(create_stock_valuation_router(lambda: {'id': 1, 'role': 'user'}, lambda: Service()))
    response = TestClient(app).post('/api/stock/valuation/intake', files={'file':('list.txt', b'00123', 'text/plain')})
    assert response.status_code == 200
    assert response.json()['rows'] == [{'article':'00123', 'requestedQty':None, 'sourceRow':1}]
    assert not observed['root'].exists()
