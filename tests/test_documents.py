from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from urllib.parse import quote

from fastapi.testclient import TestClient

import stock_sync_api
from stock_sync_web.database import WebDatabase
from stock_sync_web.documents import (
    DocumentLineInput,
    build_document_context,
    generate_document_docx,
)
from stock_sync_web.service import WebStockSyncService
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


VALID_CLIENT_CARD = {
    "legal_type": "legal_entity",
    "document_name": "OOO Romashka",
    "full_name": "Obschestvo s ogranichennoy otvetstvennostyu Romashka",
    "inn": "7707083893",
    "kpp": "770701001",
    "is_buyer": True,
    "is_supplier": False,
    "is_inactive": False,
    "bank_name_or_bik": "044525225 PAO Sberbank",
    "bank_name": "PAO Sberbank",
    "bank_bik": "044525225",
    "bank_account": "40702810900000000001",
    "correspondent_account": "30101810400000000225",
    "contact_person": "Anna Petrova",
    "email": "client@example.ru",
    "phone": "+7 999 000-00-00",
    "legal_address": "Moscow, Legal st. 1",
    "actual_address": "Moscow, Actual st. 2",
    "ogrn": "1027700132195",
    "signer_position": "Director",
    "signer_name": "Ivan Ivanov",
    "signer_basis": "Charter",
    "notes": "",
}


class DocumentGenerationTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self._temp_dir.name)

    def tearDown(self) -> None:
        self._temp_dir.cleanup()

    def test_contract_template_uses_client_and_document_fields(self) -> None:
        template_path = self.temp_path / "contract_template.docx"
        output_path = self.temp_path / "contract_output.docx"
        create_contract_template(template_path)

        context = build_document_context(
            client=VALID_CLIENT_CARD,
            document_number="D-17",
            document_date="2026-07-14",
        )
        generate_document_docx(
            template_path=template_path,
            output_path=output_path,
            context=context,
        )

        text = read_docx_text(output_path)

        self.assertIn("Contract D-17 from 14.07.2026", text)
        self.assertIn("Client: OOO Romashka", text)
        self.assertIn("INN/KPP: 7707083893 / 770701001", text)
        self.assertIn("Signer: Director Ivan Ivanov, basis: Charter", text)

    def test_contract_template_asset_uses_sm_techno_contract_wording(self) -> None:
        template_path = Path("assets/templates/contract_template.docx")
        output_path = self.temp_path / "contract_asset_output.docx"

        context = build_document_context(
            client=VALID_CLIENT_CARD,
            document_number="D-17",
            document_date="2026-07-14",
        )
        generate_document_docx(
            template_path=template_path,
            output_path=output_path,
            context=context,
        )

        text = read_docx_text(output_path)

        self.assertIn("ДОГОВОР ПОСТАВКИ № D-17", text)
        self.assertIn('ООО "СМ ТЕХНО"', text)
        self.assertIn("OOO Romashka", text)
        self.assertIn("Director Ivan Ivanov", text)
        self.assertNotIn("Director Obschestvo", text)

    def test_client_context_builds_genitive_signer_and_short_name(self) -> None:
        context = build_document_context(
            client={
                **VALID_CLIENT_CARD,
                "document_name": "АГРОЗУМ ООО",
                "signer_position": "Генеральный директор",
                "signer_name": "Столяров Сергей Михайлович",
                "signer_basis": "Устава",
            },
            document_number="D-18",
            document_date="2026-07-14",
        )

        client = context["client"]

        self.assertEqual(client["signer_position_genitive"], "генерального директора")
        self.assertEqual(client["signer_name_genitive"], "Столярова Сергея Михайловича")
        self.assertEqual(client["signer_short_name"], "Столяров С.М.")

    def test_contract_template_asset_uses_genitive_intro_and_short_signature(self) -> None:
        template_path = Path("assets/templates/contract_template.docx")
        output_path = self.temp_path / "contract_asset_agrozum.docx"

        context = build_document_context(
            client={
                **VALID_CLIENT_CARD,
                "document_name": "АГРОЗУМ ООО",
                "signer_position": "Генеральный директор",
                "signer_name": "Столяров Сергей Михайлович",
                "signer_basis": "Устава",
            },
            document_number="D-18",
            document_date="2026-07-14",
        )
        generate_document_docx(
            template_path=template_path,
            output_path=output_path,
            context=context,
        )

        text = read_docx_text(output_path)
        buyer_requisites = Document(output_path).tables[0].rows[1].cells[1].text

        self.assertIn("в лице генерального директора Столярова Сергея Михайловича", text)
        self.assertIn("___________________/ Столяров С.М.", text)
        self.assertNotIn("Генеральный директор Столяров Сергей Михайлович", text)
        self.assertLessEqual(max_consecutive_blank_lines(buyer_requisites), 2)

    def test_contract_template_keeps_requisites_section_together_without_empty_gaps(self) -> None:
        document = Document("assets/templates/contract_template.docx")
        blocks = list(iter_document_blocks(document))
        section_index = next(
            index
            for index, block in enumerate(blocks)
            if isinstance(block, Paragraph) and block.text.strip().startswith("13.")
        )
        section_heading = blocks[section_index]
        requisites_table = blocks[section_index + 1]

        self.assertIsInstance(section_heading, Paragraph)
        self.assertIsInstance(requisites_table, Table)
        self.assertTrue(section_heading.paragraph_format.keep_with_next)
        self.assertFalse(
            isinstance(blocks[section_index - 1], Paragraph)
            and not blocks[section_index - 1].text.strip()
        )
        self.assertFalse(
            section_index + 2 < len(blocks)
            and isinstance(blocks[section_index + 2], Paragraph)
            and not blocks[section_index + 2].text.strip()
        )
        assert isinstance(requisites_table, Table)
        self.assertTrue(all(row_has_cant_split(row) for row in requisites_table.rows))

    def test_generated_contract_polishes_requisites_section_layout(self) -> None:
        template_path = self.temp_path / "unpolished_contract_template.docx"
        output_path = self.temp_path / "polished_contract_output.docx"
        create_unpolished_requisites_template(template_path)

        context = build_document_context(
            client={
                **VALID_CLIENT_CARD,
                "document_name": "АГРОЗУМ ООО",
                "signer_position": "Генеральный директор",
                "signer_name": "Столяров Сергей Михайлович",
                "signer_basis": "Устава",
            },
            document_number="D-19",
            document_date="2026-07-14",
        )
        generate_document_docx(
            template_path=template_path,
            output_path=output_path,
            context=context,
        )

        document = Document(output_path)
        blocks = list(iter_document_blocks(document))
        section_index = next(
            index
            for index, block in enumerate(blocks)
            if isinstance(block, Paragraph) and block.text.strip().startswith("13.")
        )
        section_heading = blocks[section_index]
        requisites_table = blocks[section_index + 1]

        self.assertIsInstance(section_heading, Paragraph)
        self.assertIsInstance(requisites_table, Table)
        self.assertTrue(section_heading.paragraph_format.keep_with_next)
        self.assertFalse(
            isinstance(blocks[section_index - 1], Paragraph)
            and not blocks[section_index - 1].text.strip()
        )
        self.assertFalse(
            section_index + 2 < len(blocks)
            and isinstance(blocks[section_index + 2], Paragraph)
            and not blocks[section_index + 2].text.strip()
        )
        assert isinstance(requisites_table, Table)
        self.assertEqual(len(requisites_table.rows), 2)
        self.assertTrue(all(row_has_cant_split(row) for row in requisites_table.rows))
        self.assertLessEqual(max_consecutive_blank_lines(requisites_table.rows[1].cells[1].text), 2)

    def test_specification_template_repeats_commercial_offer_lines(self) -> None:
        template_path = self.temp_path / "specification_template.docx"
        output_path = self.temp_path / "specification_output.docx"
        create_specification_template(template_path)

        context = build_document_context(
            client=VALID_CLIENT_CARD,
            document_number="SP-4",
            document_date="2026-07-14",
            lines=[
                DocumentLineInput(row_no=1, article="CAT-001", name="Filter", qty=2, price=1250),
                DocumentLineInput(row_no=2, article="JCB-002", name="Pump", qty=1, price=5000),
            ],
        )
        generate_document_docx(
            template_path=template_path,
            output_path=output_path,
            context=context,
        )

        text = read_docx_text(output_path)

        self.assertIn("Specification SP-4", text)
        self.assertIn("OOO Romashka", text)
        self.assertIn("CAT-001", text)
        self.assertIn("Filter", text)
        self.assertIn("JCB-002", text)
        self.assertIn("Pump", text)
        self.assertIn("7 500.00", text)


class DocumentDatabaseMigrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self) -> None:
        self._temp_dir.cleanup()

    def test_migration_adds_document_fields_and_documents_table(self) -> None:
        old_db_path = Path(self._temp_dir.name) / "old_stock_sync.db"
        with sqlite3.connect(old_db_path) as conn:
            conn.execute(
                """
                CREATE TABLE crm_clients (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    contact_person TEXT,
                    email TEXT,
                    phone TEXT,
                    notes TEXT,
                    linked_counterparty_id INTEGER,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                INSERT INTO crm_clients(name, contact_person, email, phone, notes, created_at, updated_at)
                VALUES('Old client', 'Anna', 'old@example.ru', '+7', 'legacy', '2026-07-10T10:00:00', '2026-07-10T10:00:00')
                """
            )

        migrated = WebDatabase(db_path=old_db_path)
        row = migrated.get_crm_client(1)

        self.assertIsNotNone(row)
        assert row is not None
        self.assertEqual(row["ogrn"], "")
        self.assertEqual(row["bank_name"], "")
        self.assertEqual(row["bank_bik"], "")
        self.assertEqual(row["correspondent_account"], "")
        self.assertEqual(row["signer_position"], "")
        self.assertEqual(row["signer_name"], "")
        self.assertEqual(row["signer_basis"], "")

        with migrated.connect() as conn:
            document_columns = {
                item["name"] for item in conn.execute("PRAGMA table_info(documents)").fetchall()
            }

        self.assertIn("document_type", document_columns)
        self.assertIn("commercial_offer_id", document_columns)
        self.assertIn("missing_fields", document_columns)


class DocumentApiTest(unittest.TestCase):
    def setUp(self) -> None:
        self._temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.temp_path = Path(self._temp_dir.name)
        self.db_path = self.temp_path / "stock_sync.db"
        self.contract_template_path = self.temp_path / "contract_template.docx"
        self.specification_template_path = self.temp_path / "specification_template.docx"
        create_contract_template(self.contract_template_path)
        create_specification_template(self.specification_template_path)

        self.db = WebDatabase(db_path=self.db_path)
        self.service = WebStockSyncService(
            db=self.db,
            document_template_paths={
                "contract": self.contract_template_path,
                "specification": self.specification_template_path,
            },
            document_storage_dir=self.temp_path / "storage" / "documents",
        )
        self.original_service = stock_sync_api.SERVICE
        stock_sync_api.SERVICE = self.service
        stock_sync_api.app.dependency_overrides[stock_sync_api._get_current_user] = lambda: {
            "id": 1,
            "role": "admin",
        }
        self.client = TestClient(stock_sync_api.app)

    def tearDown(self) -> None:
        self.client.close()
        stock_sync_api.app.dependency_overrides.clear()
        stock_sync_api.SERVICE = self.original_service
        self._temp_dir.cleanup()

    def test_create_list_download_specification_from_commercial_offer(self) -> None:
        crm_client = self.db.create_crm_client_card(VALID_CLIENT_CARD)
        offer_id = self.db.create_commercial_offer(
            number="KP-99",
            client_source="local",
            counterparty_id=None,
            crm_client_id=int(crm_client["id"]),
            client_name=crm_client["document_name"],
            offer_date="2026-07-14",
            source_filename=None,
            source_path=None,
            output_path="storage/commercial_offers/test.xlsx",
            notes="",
            lines=[
                {
                    "row_no": 1,
                    "article": "CAT-001",
                    "name": "Filter",
                    "brand": "CAT",
                    "qty": 2,
                    "price_vat": 1250,
                    "amount_vat": 2500,
                    "delivery_time": "2 days",
                    "note": "",
                    "warehouse_id": None,
                    "warehouse_name": "",
                }
            ],
            created_by_user_id=1,
        )

        response = self.client.post(
            "/api/documents",
            json={
                "documentType": "specification",
                "number": "SP-99",
                "documentDate": "2026-07-14",
                "clientSource": "local",
                "clientId": int(crm_client["id"]),
                "commercialOfferId": offer_id,
                "notes": "Created from KP",
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        document_payload = response.json()["document"]
        self.assertEqual(document_payload["documentType"], "specification")
        self.assertEqual(document_payload["clientName"], "OOO Romashka")
        self.assertEqual(document_payload["commercialOfferId"], offer_id)
        self.assertEqual(document_payload["missingFields"], [])

        list_response = self.client.get("/api/documents")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.json()["items"][0]["id"], document_payload["id"])

        download_response = self.client.get(f"/api/documents/{document_payload['id']}/download")
        self.assertEqual(download_response.status_code, 200)
        self.assertEqual(
            download_response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        text = read_docx_text(BytesIO(download_response.content))
        self.assertIn("Specification SP-99", text)
        self.assertIn("CAT-001", text)

    def test_download_contract_uses_sm_techno_filename(self) -> None:
        crm_client = self.db.create_crm_client_card(VALID_CLIENT_CARD)

        response = self.client.post(
            "/api/documents",
            json={
                "documentType": "contract",
                "number": "D-42",
                "documentDate": "2026-07-14",
                "clientSource": "local",
                "clientId": int(crm_client["id"]),
                "commercialOfferId": None,
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        document_payload = response.json()["document"]
        download_response = self.client.get(f"/api/documents/{document_payload['id']}/download")

        self.assertEqual(download_response.status_code, 200)
        expected_filename = "Договор_ООО_СМ_ТЕХНО_OOO Romashka_D-42.docx"
        self.assertIn(
            f"filename*=UTF-8''{quote(expected_filename)}",
            download_response.headers["content-disposition"],
        )

    def test_contract_uses_signer_position_selected_in_generator(self) -> None:
        crm_client = self.db.create_crm_client_card(
            {
                **VALID_CLIENT_CARD,
                "signer_position": "",
                "signer_name": "Ivan Ivanov",
                "signer_basis": "Charter",
            }
        )

        response = self.client.post(
            "/api/documents",
            json={
                "documentType": "contract",
                "number": "D-43",
                "documentDate": "2026-07-14",
                "clientSource": "local",
                "clientId": int(crm_client["id"]),
                "commercialOfferId": None,
                "signerPosition": "Генеральный директор",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        document_payload = response.json()["document"]
        self.assertNotIn("signer_position", document_payload["missingFields"])

        download_response = self.client.get(f"/api/documents/{document_payload['id']}/download")
        self.assertEqual(download_response.status_code, 200)
        text = read_docx_text(BytesIO(download_response.content))
        self.assertIn("Генеральный директор Ivan Ivanov", text)


def create_contract_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("Contract {{ document.number }} from {{ document.date }}")
    document.add_paragraph("Client: {{ client.document_name }}")
    document.add_paragraph("INN/KPP: {{ client.inn }} / {{ client.kpp }}")
    document.add_paragraph("OGRN: {{ client.ogrn }}")
    document.add_paragraph("Signer: {{ client.signer_position }} {{ client.signer_name }}, basis: {{ client.signer_basis }}")
    document.save(path)


def create_specification_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("Specification {{ document.number }}")
    document.add_paragraph("Client: {{ client.document_name }}")
    table = document.add_table(rows=4, cols=6)
    for index, title in enumerate(["No", "Article", "Name", "Qty", "Price", "Amount"]):
        table.rows[0].cells[index].text = title
    table.rows[1].cells[0].text = "{%tr for line in spec.lines %}"
    table.rows[2].cells[0].text = "{{ line.row_no }}"
    table.rows[2].cells[1].text = "{{ line.article }}"
    table.rows[2].cells[2].text = "{{ line.name }}"
    table.rows[2].cells[3].text = "{{ line.qty }}"
    table.rows[2].cells[4].text = "{{ line.price }}"
    table.rows[2].cells[5].text = "{{ line.amount }}"
    table.rows[3].cells[0].text = "{%tr endfor %}"
    document.add_paragraph("Total: {{ spec.total_amount }}")
    document.save(path)


def create_unpolished_requisites_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("12.9. Previous section.")
    document.add_paragraph("")
    document.add_paragraph("13. Юридические адреса и реквизиты сторон")
    document.add_paragraph("")
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Поставщик:"
    table.rows[0].cells[1].text = "Покупатель:"
    table.rows[1].cells[0].text = "ООО « СМ ТЕХНО »\n___________________/ Дыбаев А.А."
    table.rows[1].cells[1].text = "{{ client.document_name }}\n___________________/ {{ client.signer_short_name }}"
    table.rows[2].cells[0].text = ""
    table.rows[2].cells[1].text = ""
    document.add_paragraph("")
    document.save(path)


def read_docx_text(source: Path | BytesIO) -> str:
    document = Document(source)
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def iter_document_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def row_has_cant_split(row) -> bool:
    tr_pr = row._tr.trPr
    return tr_pr is not None and tr_pr.find(qn("w:cantSplit")) is not None


def max_consecutive_blank_lines(value: str) -> int:
    longest = 0
    current = 0
    for line in value.splitlines():
        if line.strip():
            current = 0
            continue
        current += 1
        longest = max(longest, current)
    return longest


if __name__ == "__main__":
    unittest.main()
