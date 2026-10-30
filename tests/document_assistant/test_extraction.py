from io import BytesIO
import json
from pathlib import Path
import pytest
from PIL import Image
from openpyxl import Workbook
from stock_sync_web.assistant_extraction import (
    CodexDocumentExtractor, DocumentSourceReader, ExtractionError,
    ExtractionRequest, extraction_schema, validate_payload,
)
from tests.document_assistant.test_storage_flow import docx_bytes


def test_table_order_blank_cells_and_leading_zero_articles(tmp_path):
    from docx import Document
    doc=Document(); doc.add_paragraph('START'); table=doc.add_table(rows=2,cols=3)
    table.cell(0,0).text='Артикул'; table.cell(0,2).text='Количество'
    table.cell(1,0).text='00123'; table.cell(1,2).text='2'; doc.add_paragraph('END')
    doc.sections[0].header.paragraphs[0].text='HEADER REQUISITES'
    doc.sections[0].footer.paragraphs[0].text='FOOTER REQUISITES'
    path=tmp_path/'input.docx'; doc.save(path)
    texts,images=DocumentSourceReader().read(path,tmp_path)
    assert texts==['START\nАртикул\t\tКоличество\n00123\t\t2\nEND\nHEADER REQUISITES\nFOOTER REQUISITES'] and not images
    book=Workbook(); sheet=book.active; sheet.append(['Артикул','Количество']); sheet.append([123,2]); sheet['A2'].number_format='00000'
    path=tmp_path/'input.xlsx'; book.save(path)
    assert '00123\t2' in DocumentSourceReader().read(path,tmp_path)[0][0]


@pytest.mark.parametrize('suffix,content',[('.txt','00123 2'),('.csv','Артикул;Количество\n00123;2'),('.tsv','Артикул\tКоличество\n00123\t2')])
def test_text_encodings_preserve_codes(tmp_path,suffix,content):
    path=tmp_path/('input'+suffix); path.write_bytes(content.encode('cp1251'))
    texts,images=DocumentSourceReader().read(path,tmp_path)
    assert '00123' in texts[0] and not images


@pytest.mark.parametrize('suffix,fmt',[('.png','PNG'),('.jpeg','JPEG'),('.webp','WEBP'),('.tiff','TIFF'),('.heic','HEIF')])
def test_photos_become_explicit_visual_pages(tmp_path,suffix,fmt):
    if fmt=='HEIF':
        from pillow_heif import register_heif_opener
        register_heif_opener()
    source=tmp_path/('photo'+suffix)
    Image.new('RGB',(200,100),'white').save(source,format=fmt)
    texts,images=DocumentSourceReader().read(source,tmp_path)
    assert not texts and len(images)==1 and images[0].suffix=='.png'
    with Image.open(images[0]) as rendered: assert rendered.size==(200,100)


def test_scanned_pdf_renders_each_page_for_visual_extraction(tmp_path):
    source=tmp_path/'scan.pdf'; Image.new('RGB',(200,100),'white').save(source,'PDF')
    texts,images=DocumentSourceReader().read(source,tmp_path)
    assert not texts and len(images)==1


def test_model_schema_rejects_tool_events_untrusted_price_and_bad_decimals():
    result={'payload':{'rows':[{'article':'00123','requested_qty':'2','source_row':1}]},'warnings':[],'missing_fields':[]}
    events='\n'.join(json.dumps(v) for v in [
        {'type':'thread.started'}, {'type':'turn.started'},
        {'type':'item.completed','item':{'type':'agent_message','text':json.dumps(result)}}, {'type':'turn.completed'}])
    assert CodexDocumentExtractor.parse_events(events,'articles')['payload']['rows'][0]['article']=='00123'
    with pytest.raises(ExtractionError): CodexDocumentExtractor.parse_events('{"type":"item.started","item":{"type":"command_execution"}}\n'+events,'articles')
    with pytest.raises(ValueError): validate_payload('articles',{'rows':[{'article':'00123','requested_qty':'2','source_row':1,'price':'1'}]})
    with pytest.raises(ValueError): validate_payload('articles',{'rows':[{'article':'00123','requested_qty':'NaN','source_row':1}]})
    with pytest.raises(ValueError): validate_payload('articles',{'rows':[{'article':123,'requested_qty':'2','source_row':1}]})


def test_selected_material_only_and_environment_allowlist(tmp_path,monkeypatch):
    import stock_sync_web.assistant_extraction as module
    monkeypatch.setattr(module,'discover_codex_cli',lambda configured:Path('C:/verified/codex.exe'))
    monkeypatch.setenv('ONEC_PASSWORD','synthetic-secret'); monkeypatch.setenv('OPENAI_API_KEY','synthetic-other-secret')
    seen={}
    def runner(args,prompt,env,workspace,timeout):
        seen.update(args=args,prompt=prompt,env=env,workspace=workspace,files=list(workspace.iterdir()))
        result={'payload':{'rows':[{'article':'00123','requested_qty':None,'source_row':1}]},'warnings':[],'missing_fields':[]}
        return '\n'.join(map(json.dumps,[{'type':'item.completed','item':{'type':'agent_message','text':json.dumps(result)}},{'type':'turn.completed'}]))
    result=CodexDocumentExtractor(runner=runner).extract(ExtractionRequest('articles',['00123'],[],[8]))
    assert result.payload['rows'][0]['article']=='00123'
    assert 'ONEC_PASSWORD' not in seen['env'] and 'OPENAI_API_KEY' not in seen['env']
    assert not seen['workspace'].is_relative_to(Path.cwd())
    assert {p.name for p in seen['files']}=={'output.schema.json','temp'}
    assert result.provenance==[{'assetId':8,'method':'gpt-6.1-terra','reviewed':False}]


def test_malformed_encrypted_and_oversized_sources_fail(tmp_path):
    from pypdf import PdfWriter
    from stock_sync_web.assistant_storage import UploadLimits
    path=tmp_path/'bad.docx'; path.write_bytes(b'not a zip')
    with pytest.raises(ValueError): DocumentSourceReader().read(path,tmp_path)
    path=tmp_path/'locked.pdf'; writer=PdfWriter(); writer.add_blank_page(100,100); writer.encrypt('secret'); writer.write(path)
    with pytest.raises(ValueError): DocumentSourceReader().read(path,tmp_path)
    path=tmp_path/'large.txt'; path.write_bytes(b'12345')
    with pytest.raises(ExtractionError): DocumentSourceReader(UploadLimits(max_bytes=4)).read(path,tmp_path)


def test_legacy_xls_fixture_has_real_cells_and_preserves_text_sku(tmp_path):
    path=Path(__file__).parent/'fixtures/synthetic.xls'
    texts,images=DocumentSourceReader().read(path,tmp_path)
    assert '00123\t2' in texts[0] and not images


def test_legacy_office_missing_converter_fails_explicitly(tmp_path):
    from stock_sync_web.assistant_extraction import IsolatedOfficeConverter
    reader=DocumentSourceReader(converter=IsolatedOfficeConverter(image=''))
    for suffix in ('doc','rtf','odt'):
        with pytest.raises(ExtractionError) as error:
            reader.read(Path(__file__).parent/('fixtures/synthetic.'+suffix),tmp_path)
        assert error.value.code=='converter_unavailable'


def test_cli_timeout_and_nonzero_never_expose_process_diagnostics(tmp_path):
    import sys
    from stock_sync_web.assistant_extraction import _runtime_environment
    with pytest.raises(ExtractionError) as error:
        CodexDocumentExtractor._run([sys.executable,'-c','import time; time.sleep(5)'],'',_runtime_environment(tmp_path),tmp_path,0.1)
    assert error.value.code=='runtime_timeout'
    with pytest.raises(ExtractionError) as error:
        CodexDocumentExtractor._run([sys.executable,'-c','import sys; print("SYNTHETIC_PRIVATE_PATH",file=sys.stderr); sys.exit(2)'],'',_runtime_environment(tmp_path),tmp_path,10)
    assert error.value.code=='runtime_failed' and 'SYNTHETIC_PRIVATE_PATH' not in str(error.value)
