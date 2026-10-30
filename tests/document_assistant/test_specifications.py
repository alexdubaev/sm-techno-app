from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
import pytest

from tests.document_assistant.test_storage_flow import api, upload, docx_bytes
from tests.document_assistant.test_clients_contracts import FIELDS, requisites, confirm


INVOICE = {'invoice_number':'254','date':'2026-10-30','buyer_inn':FIELDS['inn'],'buyer_kpp':FIELDS['kpp'],
    'buyer_document_name':FIELDS['document_name'],'currency':'RUB','total_amount':'310000.00','vat_amount':'55901.64',
    'amount_in_words':'Триста десять тысяч рублей 00 копеек','lines':[{'row_no':1,'article':'SYN-795','name':'SYN-795 Тестовая шестерня',
    'qty':'1','unit':'шт','unit_price_incl_vat':'310000.00','unit_price_excl_vat':'254098.36','vat_amount':'55901.64','total_amount':'310000.00'}]}


def invoice_bytes(payload=None, extra_rows=False):
    p=payload or INVOICE; d=Document()
    d.add_paragraph('Счет на оплату № '+p['invoice_number']+' от 30.10.2026')
    d.add_paragraph('Поставщик: ООО Продавец, ИНН 7804702073, КПП 780401001')
    d.add_paragraph('Покупатель: '+p['buyer_document_name']+', ИНН '+p['buyer_inn']+', КПП '+p['buyer_kpp'])
    d.add_paragraph('Основной договор: НЕ ИСПОЛЬЗОВАТЬ 999')
    t=d.add_table(rows=0,cols=9)
    for values in [['']*9,['','№','Товар (Услуга)','Кол-во','Ед.','Цена с НДС','Цена без ндс','Сумма НДС','Сумма'],['']+list(map(str,range(1,9)))]:
        for c,v in zip(t.add_row().cells,values):c.text=v
    for line in p['lines']:
        for c,v in zip(t.add_row().cells,['',str(line['row_no']),line['name'],line['qty'],line['unit'],line['unit_price_incl_vat'],line['unit_price_excl_vat'],line['vat_amount'],line['total_amount']]):c.text=v
    if extra_rows:
        from decimal import Decimal
        r=t.add_row();r.cells[2].merge(r.cells[4]).text=str(sum(Decimal(x['qty']) for x in p['lines']))
        r.cells[5].merge(r.cells[8]).text=p['total_amount']
        t.add_row()
    d.add_paragraph('Итого: '+p['total_amount'])
    d.add_paragraph('В том числе НДС: '+p['vat_amount'])
    d.add_paragraph('Всего к оплате: '+p['total_amount'])
    d.add_paragraph('Всего наименований '+str(len(p['lines']))+', на сумму '+p['total_amount']+' руб.')
    d.add_paragraph(p['amount_in_words'])
    b=BytesIO();d.save(b);return b.getvalue()


def setup_invoice(api,monkeypatch,payload=None,source=None):
    http,service,users,cid=api
    with service.db.transaction(immediate=True) as c:c.execute('UPDATE crm_clients SET inn=?,kpp=? WHERE id=?',(FIELDS['inn'],FIELDS['kpp'],cid))
    j,_=requisites(http,users,monkeypatch,cid=cid)
    r=confirm(http,users,j,'generate_contract',{**FIELDS,'contract_number':'Д-37','contract_date':'2026-10-30'})
    assert r.status_code==200,r.text
    from stock_sync_web.assistant_extraction import ExtractionResult
    p=payload or INVOICE
    monkeypatch.setattr('stock_sync_web.assistant_extraction.CodexDocumentExtractor.extract',lambda self,request:ExtractionResult('invoice',json.loads(json.dumps(p))))
    a=upload(http,users,cid,source or invoice_bytes(),role='source_invoice',name='invoice.docx').json()
    j=http.post('/api/document-assistant/jobs',headers=users['alice']['headers'],json={'kind':'invoice','clientId':cid,'assetIds':[a['id']]}).json()
    j=http.get('/api/document-assistant/jobs/'+j['id'],headers=users['alice']['headers']).json()
    assert j['status']=='review',j
    return j,a


def generate(api,job,key='generate-specification-0001',cid=None):
    http,_,users,client=api
    return http.post(f'/api/crm/clients/{cid or client}/specifications',headers=users['boris']['headers'],json={'invoiceJobId':job['id'],'operationKey':key})


def test_invoice_to_specification_final_retains_original_and_exact_sources(api,monkeypatch):
    http,service,users,cid=api;j,a=setup_invoice(api,monkeypatch)
    r=generate(api,j);assert r.status_code==200,r.text
    spec=r.json();assert spec['number']=='1' and spec['invoiceAssetId']==a['id'] and spec['contractVersionId']
    body=http.get(f"/api/documents/{spec['id']}/download",headers=users['alice']['headers']).content
    d=Document(BytesIO(body));assert len(d.tables[1].columns)==8
    assert [c.text for c in d.tables[1].rows[1].cells]==['1','SYN-795 Тестовая шестерня','1','шт','310 000,00','254 098,36','55 901,64','310 000,00']
    text='\n'.join(p.text for p in d.paragraphs)
    assert 'Д-37' in text and '999' not in text and INVOICE['amount_in_words'] in text
    assert FIELDS['document_name'] in text
    assert generate(api,j).json()['id']==spec['id']
    final=docx_bytes('Менеджер дописал условия')
    r=http.post(f"/api/crm/clients/{cid}/documents/{spec['id']}/versions",headers=users['alice']['headers'],data={'role':'final','expectedVersion':1},files={'file':('final.docx',final)})
    assert r.status_code==200,r.text
    assert http.get(f"/api/document-assets/{spec['assets'][0]['id']}/download",headers=users['boris']['headers']).content==body
    assert http.get(f"/api/documents/{spec['id']}/download",headers=users['boris']['headers']).content==final
    again=generate(api,j).json();assert again['id']==spec['id'] and again['version']==2
    from stock_sync_web.assistant_repository import AssistantRepository
    saved=AssistantRepository(service.db).get_document(spec['id']);snapshot=json.loads(saved['legal_snapshot_json'])
    assert snapshot['source_sha256'] and snapshot['legal_fields']['signer_name']==FIELDS['signer_name']


@pytest.mark.parametrize('change',[{'total_amount':'310001.00'},{'buyer_inn':'7712345678'},{'amount_in_words':'Один рубль 00 копеек'},{'lines':[]}])
def test_untrusted_model_cannot_publish_changed_invoice(api,monkeypatch,change):
    j,_=setup_invoice(api,monkeypatch,{**INVOICE,**change})
    r=generate(api,j);assert r.status_code in (409,422),r.text
    http,_,users,cid=api
    summary=http.get(f'/api/crm/clients/{cid}/documents',headers=users['alice']['headers']).json()
    assert len(summary['documents'])==1 and summary['nextSpecNumber']==1


def test_archive_sequence_parallel_and_replay_are_atomic(api,monkeypatch):
    http,_,users,cid=api;j,_=setup_invoice(api,monkeypatch)
    a=upload(http,users,cid,docx_bytes('Архивная спецификация №5')).json()
    r=http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'specification','assetId':a['id'],'number':'5','date':'2026-10-29'})
    assert r.status_code==200,r.text
    with ThreadPoolExecutor(max_workers=2) as w:results=list(w.map(lambda k:generate(api,j,k),['parallel-spec-0001','parallel-spec-0002']))
    assert [r.status_code for r in results]==[200,200],[r.text for r in results]
    assert {r.json()['number'] for r in results}=={'6','7'}
    assert generate(api,j,'parallel-spec-0001').json()['id']==results[0].json()['id']


def test_changed_canonical_file_is_read_and_foreign_job_rejected(api,monkeypatch):
    http,service,users,cid=api;j,_=setup_invoice(api,monkeypatch)
    first=generate(api,j).json()
    other=service.db.create_crm_client_card({'document_name':'Чужой клиент'})['id']
    assert generate(api,j,'foreign-client-001',other).status_code==404
    replacement=docx_bytes('Нечитаемый договор без юридических данных')
    r=http.post(f"/api/crm/clients/{cid}/documents/{first['contractId']}/versions",headers=users['alice']['headers'],data={'role':'contract_version','expectedVersion':1},files={'file':('changed.docx',replacement)})
    assert r.status_code==200,r.text
    failed=generate(api,j,'changed-contract-001');assert failed.status_code in (409,422),failed.text
    assert generate(api,j).json()['contractVersionId']==first['contractVersionId']


def test_rounding_service_rows_and_missing_image_relationship_do_not_lose_goods(api,monkeypatch):
    source=invoice_bytes(extra_rows=True)
    # A missing nontext logo target must not prevent safe body XML extraction.
    original=ZipFile(BytesIO(source));b=BytesIO()
    with ZipFile(b,'w') as z:
        for n in original.namelist():
            data=original.read(n)
            if n=='word/_rels/document.xml.rels':data=data.replace(b'</Relationships>',b'<Relationship Id="rLogo" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/missing.png"/></Relationships>')
            z.writestr(n,data)
    j,_=setup_invoice(api,monkeypatch,source=b.getvalue())
    assert generate(api,j).status_code==200


def test_source_invoice_date_cannot_be_changed_by_model(api,monkeypatch):
    j,_=setup_invoice(api,monkeypatch,{**INVOICE,'date':'2025-10-30'})
    r=generate(api,j);assert r.status_code==422,r.text


def test_source_rounding_multiple_goods_and_new_contract_version_keep_old_snapshot(api,monkeypatch):
    from stock_sync_web.assistant_repository import AssistantRepository
    http,service,users,cid=api
    p=json.loads(json.dumps(INVOICE));p['lines'][0].update(qty='3',unit_price_incl_vat='100.00',unit_price_excl_vat='81.97',vat_amount='54.10',total_amount='300.00')
    p['lines'].append({**p['lines'][0],'row_no':2,'name':'SYN-001 Тестовый товар','article':'SYN-001','qty':'1','unit_price_incl_vat':'10.00','unit_price_excl_vat':'10.00','vat_amount':'0.00','total_amount':'10.00'})
    p.update(total_amount='310.00',vat_amount='54.10',amount_in_words='Триста десять рублей 00 копеек')
    j,_=setup_invoice(api,monkeypatch,p,invoice_bytes(p,True))
    first=generate(api,j);assert first.status_code==200,first.text
    first=first.json();repo=AssistantRepository(service.db)
    snapshot=repo.get_document(first['id'])['legal_snapshot_json']
    d=Document(BytesIO(http.get(f"/api/documents/{first['id']}/download",headers=users['alice']['headers']).content))
    assert len(d.tables[1].rows)==3 and d.tables[1].rows[1].cells[6].text=='54,10'
    q,_=requisites(http,users,monkeypatch,{**FIELDS,'signer_name':'Петров Пётр Петрович'},cid)
    updated=confirm(http,users,q,'generate_contract',{**FIELDS,'signer_name':'Петров Пётр Петрович','contract_number':'Д-37','contract_date':'2026-10-30'},key='new-signer-contract')
    assert updated.status_code==200,updated.text
    second=generate(api,j,'second-spec-new-version');assert second.status_code==200,second.text
    assert second.json()['contractVersionId']!=first['contractVersionId']
    new=json.loads(repo.get_document(second.json()['id'])['legal_snapshot_json'])
    assert new['legal_fields']['signer_name']=='Петров Пётр Петрович'
    assert repo.get_document(first['id'])['legal_snapshot_json']==snapshot


def test_extractor_separate_article_and_printed_date_preserve_complete_source_name(api,monkeypatch):
    p=json.loads(json.dumps(INVOICE));p['lines'][0]['name']='Тестовая шестерня';p['date']='30.10.2026'
    j,_=setup_invoice(api,monkeypatch,p)
    r=generate(api,j);assert r.status_code==200,r.text
    http,_,users,_=api
    d=Document(BytesIO(http.get(f"/api/documents/{r.json()['id']}/download",headers=users['alice']['headers']).content))
    assert d.tables[1].rows[1].cells[1].text=='SYN-795 Тестовая шестерня'


def test_response_failure_after_commit_preserves_published_file_and_replay(api,monkeypatch):
    from stock_sync_web.assistant_repository import AssistantRepository
    from stock_sync_web.assistant_service import DocumentAssistantService
    http,service,users,cid=api;j,_=setup_invoice(api,monkeypatch)
    original=AssistantRepository.get_document;failed=[]
    def once(self,*args,**kwargs):
        row=original(self,*args,**kwargs)
        if row['document_type']=='specification' and not failed:
            failed.append(True);raise OSError('Synthetic response read failure after commit')
        return row
    monkeypatch.setattr(AssistantRepository,'get_document',once)
    assert generate(api,j).status_code==503
    replay=generate(api,j);assert replay.status_code==200,replay.text
    assert http.get(f"/api/documents/{replay.json()['id']}/download",headers=users['alice']['headers']).status_code==200
    assets=AssistantRepository(service.db).list_client_assets(cid,replay.json()['id'])
    assert DocumentAssistantService(service).store.exists(assets[0]['storage_key'])


@pytest.mark.parametrize('number,expected',[('5','6'),('неизвестен',None)])
def test_preexisting_legacy_archive_numbers_are_reconciled_before_allocation(api,monkeypatch,number,expected):
    http,service,users,cid=api;j,_=setup_invoice(api,monkeypatch)
    output=service.document_exports_dir/'legacy-spec.docx';output.write_bytes(docx_bytes('Прежняя спецификация'))
    legacy_id=service.db.create_document(document_type='specification',number=number,client_source='local',counterparty_id=None,crm_client_id=cid,
        commercial_offer_id=None,client_name=FIELDS['document_name'],document_date='2026-10-01',output_path=service._store_path(output),notes='',missing_fields='[]',created_by_user_id=users['alice']['id'])
    from stock_sync_web.assistant_repository import AssistantRepository
    canonical=AssistantRepository(service.db).get_client_document_state(cid)['contract_document_id']
    with service.db.transaction(immediate=True) as conn:conn.execute('UPDATE documents SET contract_document_id=? WHERE id=?',(canonical,legacy_id))
    r=generate(api,j)
    if expected:assert r.status_code==200 and r.json()['number']==expected,r.text
    else:assert r.status_code==409,r.text


def test_imported_ready_contract_is_extracted_from_saved_bytes_and_keeps_original(api,monkeypatch):
    from stock_sync_web.assistant_extraction import ExtractionResult
    http,_,users,cid=api;j,_=setup_invoice(api,monkeypatch)
    summary=http.get(f'/api/crm/clients/{cid}/documents',headers=users['alice']['headers']).json()
    old=summary['contract'];content=http.get(f"/api/documents/{old['id']}/download",headers=users['alice']['headers']).content
    assert http.delete(f"/api/documents/{old['id']}",headers=users['admin']['headers']).status_code==200
    asset=upload(http,users,cid,content).json()
    imported=http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'contract','assetId':asset['id'],'number':'Д-37','date':'2026-10-30'})
    assert imported.status_code==200,imported.text
    seen=[]
    def extract(self,request):
        seen.extend(request.texts)
        assert request.kind=='contract' and request.source_asset_ids==[asset['id']]
        return ExtractionResult('contract',{'number':'Д-37','date':'30.10.2026','legal_fields':dict(FIELDS)})
    monkeypatch.setattr('stock_sync_web.assistant_extraction.CodexDocumentExtractor.extract',extract)
    spec=generate(api,j);assert spec.status_code==200,spec.text
    assert any(FIELDS['document_name'] in text for text in seen)
    assert spec.json()['contractId']==imported.json()['id'] and spec.json()['contractVersionId']==asset['id']
    assert http.get(f"/api/document-assets/{asset['id']}/download",headers=users['boris']['headers']).content==content
