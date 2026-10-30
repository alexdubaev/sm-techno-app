from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
import pytest
from fastapi.testclient import TestClient
from docx import Document
import stock_sync_api
from stock_sync_web.database import WebDatabase
from stock_sync_web.service import WebStockSyncService


def docx_bytes(text):
    stream = BytesIO(); document = Document(); document.add_paragraph(text); document.save(stream)
    return stream.getvalue()


@pytest.fixture
def api(tmp_path, monkeypatch):
    db = WebDatabase(tmp_path / 'test.db')
    service = WebStockSyncService(db=db, commercial_offer_storage_dir=tmp_path/'storage/offers', document_storage_dir=tmp_path/'storage/documents')
    monkeypatch.setattr(stock_sync_api, 'SERVICE', service)
    users = {}
    for name, role in [('alice','user'),('boris','user'),('admin','admin')]:
        uid = db.create_user(username=name, password='isolated-password', role=role, full_name=name)
        token = service.login_app_user(name, 'isolated-password')['token']
        users[name] = {'id':uid,'headers':{'Authorization':f'Bearer {token}'}}
    client = db.create_crm_client_card({'document_name':'ООО Тест документов', 'full_name':'Общество Тест документов', 'notes':'PRIVATE CANARY', 'crm_owner_user_id':users['alice']['id']})
    with TestClient(stock_sync_api.app) as http:
        yield http, service, users, int(client['id'])


def upload(http, users, cid, content, role='archive', name='original.docx', user='alice'):
    return http.post(f'/api/crm/clients/{cid}/document-assets', headers=users[user]['headers'], data={'role':role}, files={'file':(name,content,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})


def test_cross_manager_import_final_versions_and_admin_delete(api):
    http, service, users, cid = api
    original = docx_bytes('Original contract')
    response = upload(http,users,cid,original)
    assert response.status_code == 200, response.text
    asset = response.json(); assert 'storage_key' not in response.text and 'storageKey' not in response.text
    response = http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['boris']['headers'],json={'type':'contract','assetId':asset['id'],'number':'Д-17','date':'2026-10-30'})
    assert response.status_code == 200,response.text
    document = response.json(); did = document['id']
    summary = http.get(f'/api/crm/clients/{cid}/documents', headers=users['boris']['headers'])
    assert summary.status_code == 200,summary.text
    assert 'PRIVATE CANARY' not in summary.text
    assert summary.json()['contract']['id'] == did
    assert http.get(f'/api/document-assets/{asset["id"]}/download',headers=users['boris']['headers']).content == original
    assert http.delete(f'/api/documents/{did}',headers=users['alice']['headers']).status_code == 403
    final = docx_bytes('Reviewed final contract')
    result = http.post(f'/api/crm/clients/{cid}/documents/{did}/versions',headers=users['boris']['headers'],data={'role':'contract_version','expectedVersion':1},files={'file':('final.docx',final,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert result.status_code == 200,result.text
    assert result.json()['version'] == 2
    assert http.get(f'/api/document-assets/{asset["id"]}/download',headers=users['boris']['headers']).content == original
    assert http.get(f'/api/documents/{did}/download',headers=users['boris']['headers']).content == final
    assert http.delete(f'/api/documents/{did}',headers=users['admin']['headers']).status_code == 200
    assert http.get(f'/api/document-assets/{asset["id"]}/download',headers=users['admin']['headers']).status_code == 404
    assert not list((service.document_storage_dir/'client_assets').glob('*/*'))


def test_article_offer_uses_all_exact_stock_variants_and_replays(api):
    http, service, users, cid = api
    item = service.db.create_local_item(sku='00123', name='Синтетическая позиция', print_name='Документное имя',
        category_name='CAT', group_name='Тест', price=90, warehouses=[
            {'warehouse_name':'Активный A','quantity':2}, {'warehouse_name':'Активный B','quantity':5}])
    with service.db.transaction(immediate=True) as conn:
        conn.execute("UPDATE item_warehouse_balances SET price=CASE warehouse_id WHEN (SELECT id FROM warehouses WHERE name='Активный A') THEN 100 ELSE 250 END WHERE item_id=?",(item['id'],))
        conn.execute("INSERT INTO warehouses(name,is_active,created_at,updated_at) VALUES('Неактивный склад',0,'2026-10-30','2026-10-30')")
        inactive=conn.execute("SELECT id FROM warehouses WHERE name='Неактивный склад'").fetchone()[0]
        conn.execute("INSERT INTO item_warehouse_balances(item_id,warehouse_id,quantity,price,updated_at) VALUES(?,?,9,999,'2026-10-30')",(item['id'],inactive))
        warehouse=conn.execute("SELECT id FROM warehouses WHERE name='Активный A'").fetchone()[0]
        conn.execute("INSERT INTO items(sku,sku_normalized,name,print_name,category_name,price,is_local,created_at,updated_at) VALUES('AMB-1','amb-1','Ambiguous one','Ambiguous one','CAT',33,1,'2026-10-30','2026-10-30')")
        ambiguous_one=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.execute("INSERT INTO items(sku,sku_normalized,name,print_name,category_name,price,is_local,created_at,updated_at) VALUES('amb-1',NULL,'Ambiguous two','Ambiguous two','CAT',44,1,'2026-10-30','2026-10-30')")
        ambiguous_two=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.execute("INSERT INTO item_warehouse_balances(item_id,warehouse_id,quantity,price,updated_at) VALUES(?,?,2,33,'2026-10-30')",(ambiguous_one,warehouse))
        conn.execute("INSERT INTO item_warehouse_balances(item_id,warehouse_id,quantity,price,updated_at) VALUES(?,?,3,44,'2026-10-30')",(ambiguous_two,warehouse))
    service.db.create_local_item(sku='ZERO-1',name='Нулевой остаток',print_name='Нулевой остаток',category_name='CAT',group_name='Тест',price=10,quantity=0)
    service.db.create_local_item(sku='MISSING-70',name='Похожий артикул',print_name='Похожий артикул',category_name='CAT',group_name='Тест',price=10,warehouses=[{'warehouse_name':'Активный A','quantity':1}])
    service.db.create_local_item(sku='ALL-2',name='Полный остаток',print_name='Полный остаток',category_name='CAT',group_name='Тест',price=31,warehouses=[
        {'warehouse_name':'Активный A','quantity':3},{'warehouse_name':'Активный B','quantity':7}])
    with service.db.transaction(immediate=True) as conn:
        conn.execute("UPDATE item_warehouse_balances SET price=31 WHERE item_id=(SELECT id FROM items WHERE sku='ALL-2')")
    from stock_sync_web.assistant_repository import AssistantRepository
    from stock_sync_web.assistant_service import DocumentAssistantService
    repo=AssistantRepository(service.db)
    assistant=DocumentAssistantService(service)
    staged=assistant.store.stage_upload(BytesIO(b'00123 4'),'articles.txt','text/plain',assistant.limits)
    stored=assistant.store.publish(staged,assistant.store.new_key())
    source=repo.create_asset(cid,'attachment',stored,users['alice']['id'])
    job=repo.create_job(users['alice']['id'],'articles',cid,[source['id']])
    job=repo.transition_job(job['id'],job['version'],'running')
    rows=[{'article':'00123','requested_qty':'4','source_row':1},{'article':'00123','requested_qty':'4','source_row':2},
        {'article':'MISSING-7','requested_qty':None,'source_row':3},{'article':'ZERO-1','requested_qty':None,'source_row':4},
        {'article':'AMB-1','requested_qty':None,'source_row':5},{'article':'ALL-2','requested_qty':None,'source_row':6}]
    job=repo.transition_job(job['id'],job['version'],'review',result={'payload':{'rows':rows},'warnings':[],'missing_fields':[],'provenance':[]})
    review=http.post(f"/api/document-assistant/jobs/{job['id']}/confirm",headers=users['alice']['headers'],json={'version':job['version'],'operationKey':'article-review-0001','action':'review','fields':job['result']['payload']})
    assert review.status_code==200,review.text
    before=service.db.list_orders(include_all=True)
    with service.db.connect() as conn:
        stock_before=[tuple(row) for row in conn.execute('SELECT item_id,warehouse_id,quantity,price FROM item_warehouse_balances ORDER BY item_id,warehouse_id')]
        reservations_before=[tuple(row) for row in conn.execute('SELECT * FROM order_reservations ORDER BY id')]
    def create_offer(_):
        return http.post(f'/api/crm/clients/{cid}/article-offers',headers=users['alice']['headers'],json={'articlesJobId':job['id'],'operationKey':'article-offer-0001'})
    with ThreadPoolExecutor(max_workers=2) as workers:
        concurrent=list(workers.map(create_offer,range(2)))
    assert all(response.status_code==200 for response in concurrent),[response.text for response in concurrent]
    assert concurrent[0].json()['offer']['id']==concurrent[1].json()['offer']['id']
    created=concurrent[0]
    assert created.status_code==200,created.text
    offer=created.json(); lines=offer['lines']
    assert [(line['article'],line['qty'],line['priceVat'],line['warehouseName']) for line in lines]==[
        ('00123',2,100,'Активный A'),('00123',4,250,'Активный B'),('ALL-2',3,31,'Активный A'),('ALL-2',7,31,'Активный B')]
    assert offer['offer']['clientName']=='ООО Тест документов'
    assert [row['status'] for row in offer['articleResolution']]==['matched','duplicate','not_found','zero_stock','ambiguous','matched']
    assert http.get(f"/api/document-assets/{source['id']}/download",headers=users['boris']['headers']).content==b'00123 4'
    assert http.get(f"/api/commercial-offers/{offer['offer']['id']}",headers=users['boris']['headers']).status_code==404
    assert all(item['id']!=offer['offer']['id'] for item in http.get('/api/commercial-offers',headers=users['boris']['headers']).json()['items'])
    assert http.get(f"/api/commercial-offers/{offer['offer']['id']}",headers=users['admin']['headers']).status_code==200
    output=http.get(f"/api/commercial-offers/{offer['offer']['id']}/download/output",headers=users['alice']['headers'])
    assert output.status_code==200
    from openpyxl import load_workbook
    sheet=load_workbook(BytesIO(output.content),data_only=False)['КП']
    assert sheet['B13'].value=='00123' and sheet['F13'].value==100 and sheet['I13'].value=='Активный A · наличие: 2'
    assert sheet['B14'].value=='00123' and sheet['F14'].value==250 and sheet['I14'].value=='Активный B · наличие: 5'
    assert sheet['G13'].value=='=IF(OR(E13="",F13=""),"",E13*F13)'
    assert sheet['E15'].value==3 and sheet['E16'].value==7
    with service.db.connect() as conn:
        stock_after=[tuple(row) for row in conn.execute('SELECT item_id,warehouse_id,quantity,price FROM item_warehouse_balances ORDER BY item_id,warehouse_id')]
        reservations_after=[tuple(row) for row in conn.execute('SELECT * FROM order_reservations ORDER BY id')]
    assert stock_after==stock_before and reservations_after==reservations_before
    replay=http.post(f'/api/crm/clients/{cid}/article-offers',headers=users['alice']['headers'],json={'articlesJobId':job['id'],'operationKey':'article-offer-0001'})
    assert replay.status_code==200 and replay.json()['offer']['id']==offer['offer']['id']
    assert service.db.list_orders(include_all=True)==before


def test_article_review_rejects_conflicting_duplicate_quantities(api):
    http,service,users,cid=api
    from stock_sync_web.assistant_repository import AssistantRepository
    from stock_sync_web.assistant_service import DocumentAssistantService
    repo=AssistantRepository(service.db);assistant=DocumentAssistantService(service)
    staged=assistant.store.stage_upload(BytesIO(b'00123'),'articles.txt','text/plain',assistant.limits)
    asset=repo.create_asset(cid,'attachment',assistant.store.publish(staged,assistant.store.new_key()),users['alice']['id'])
    job=repo.create_job(users['alice']['id'],'articles',cid,[asset['id']])
    job=repo.transition_job(job['id'],job['version'],'running')
    rows=[{'article':'00123','requested_qty':'2','source_row':1},{'article':'00123','requested_qty':None,'source_row':2}]
    job=repo.transition_job(job['id'],job['version'],'review',result={'payload':{'rows':rows},'warnings':[],'missing_fields':[],'provenance':[]})
    response=http.post(f"/api/document-assistant/jobs/{job['id']}/confirm",headers=users['alice']['headers'],json={
        'version':job['version'],'operationKey':'article-conflict-0001','action':'review','fields':{'rows':rows}})
    assert response.status_code==422 and 'строках 1 и 2' in response.json()['detail']
    assert repo.get_job(job['id'])['status']=='review'

def test_legacy_generation_cannot_create_second_canonical_contract(api):
    http, service, users, cid=api
    source=upload(http,users,cid,docx_bytes('Original')).json()
    assert http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'contract','assetId':source['id'],'number':'17','date':'2026-10-30'}).status_code==200
    response=http.post('/api/documents',headers=users['boris']['headers'],json={'documentType':'contract','clientSource':'local','clientId':cid,'number':'18','documentDate':'2026-10-30'})
    assert response.status_code==409,response.text

@pytest.mark.parametrize('suffix,body,mime',[('doc',b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'+b'junk'*40,'application/msword')])
def test_legacy_office_requires_real_document_structure(api,suffix,body,mime):
    http,service,users,cid=api
    result=http.post(f'/api/crm/clients/{cid}/document-assets',headers=users['alice']['headers'],data={'role':'attachment'},files={'file':('fake.'+suffix,body,mime)})
    assert result.status_code==422,result.text
    assert not list((service.document_storage_dir/'client_assets').glob('*/*'))

def test_legacy_download_checks_immutable_asset_content(api):
    http,service,users,cid=api
    asset=upload(http,users,cid,docx_bytes('Original')).json()
    response=http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'contract','assetId':asset['id'],'number':'17','date':'2026-10-30'})
    did=response.json()['id']
    from stock_sync_web.assistant_repository import AssistantRepository
    from stock_sync_web.assistant_service import DocumentAssistantService
    raw=AssistantRepository(service.db).get_asset(asset['id'])
    DocumentAssistantService(service).store.path_for_key(raw['storage_key']).write_bytes(b'changed')
    result=http.get(f'/api/documents/{did}/download',headers=users['boris']['headers'])
    assert result.status_code==409,result.text
    assert str(service.document_storage_dir) not in result.text

def test_unassigned_drafts_are_actor_scoped_even_with_known_asset_id(api):
    http,service,users,cid=api
    from stock_sync_web.assistant_service import DocumentAssistantService
    from stock_sync_web.assistant_repository import AssistantRepository
    assistant=DocumentAssistantService(service); repo=AssistantRepository(service.db)
    staged=assistant.store.stage_upload(BytesIO(b'Synthetic draft'),'draft.txt','text/plain',assistant.limits)
    stored=assistant.store.publish(staged,assistant.store.new_key())
    asset=repo.create_asset(None,'requisites',stored,users['alice']['id'])
    with pytest.raises(PermissionError):
        assistant.read_asset({'id':users['boris']['id'],'role':'user'},asset['id'])
    assert assistant.read_asset({'id':users['alice']['id'],'role':'user'},asset['id'])[0]==b'Synthetic draft'
    with pytest.raises(Exception,match='не найден'):
        repo.create_job(users['boris']['id'],'requisites',None,[asset['id']])


def test_archive_sequence_allocation_is_concurrent_and_replay_bound_to_source(api):
    from concurrent.futures import ThreadPoolExecutor
    from stock_sync_web.assistant_repository import AssistantRepository,DocumentConflict
    http,service,users,cid=api; repo=AssistantRepository(service.db)
    original=upload(http,users,cid,docx_bytes('Original')).json()
    contract=http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'contract','assetId':original['id'],'number':'17','date':'2026-10-30'}).json()
    archived=upload(http,users,cid,docx_bytes('Archived spec')).json()
    assert http.post(f'/api/crm/clients/{cid}/documents/import',headers=users['alice']['headers'],json={'type':'specification','assetId':archived['id'],'number':'5','date':'2026-10-30','contractId':contract['id']}).status_code==200
    invoice=upload(http,users,cid,docx_bytes('Invoice'),role='source_invoice').json()
    job=repo.create_job(users['alice']['id'],'invoice',cid,[invoice['id']]); job=repo.transition_job(job['id'],job['version'],'running'); job=repo.transition_job(job['id'],job['version'],'review',result={'synthetic':True})
    state=repo.get_client_document_state(cid)
    payload={'client_name':'ООО Тест документов','document_date':'2026-10-30','output_path':contract['assets'][0]['originalFilename'],'legal_snapshot':{'number':'17'}}
    def allocate(index):
        key=f'specification-{index}'
        repo.begin_operation(key,users['alice']['id'],job['id'],'generate_specification')
        return repo.allocate_specification(cid,contract['id'],state['contract_version_asset_id'],invoice['id'],key,payload)
    with ThreadPoolExecutor(max_workers=4) as workers: results=list(workers.map(allocate,range(8)))
    assert sorted(row['spec_sequence'] for row in results)==list(range(6,14))
    replay=allocate(0); assert replay['id']==results[0]['id']
    assert repo.get_client_document_state(cid)['next_spec_number']==14
    with pytest.raises(DocumentConflict):
        repo.allocate_specification(cid,contract['id'],state['contract_version_asset_id'],archived['id'],'specification-0',payload)

def test_publication_rollback_and_legacy_version_delete_leave_no_orphans(api):
    from stock_sync_web.assistant_repository import AssistantRepository
    http,service,users,cid=api
    with patch.object(AssistantRepository,'create_asset',side_effect=PermissionError('private-path-do-not-expose')):
        result=upload(http,users,cid,docx_bytes('Rejected upload'))
        assert result.status_code==503 and 'private-path-do-not-expose' not in result.text
    assert not AssistantRepository(service.db).list_client_assets(cid)
    assert not list((service.document_storage_dir/'client_assets').glob('*/*'))
    document=http.post('/api/documents',headers=users['alice']['headers'],json={'documentType':'contract','clientSource':'local','clientId':cid,'number':'17','documentDate':'2026-10-30'}).json()['document']
    old_path=service._resolve_document_stored_path(service.db.get_document(document['id'])['output_path'])
    response=http.post(f'/api/crm/clients/{cid}/documents/{document["id"]}/versions',headers=users['boris']['headers'],data={'role':'contract_version','expectedVersion':1},files={'file':('final.docx',docx_bytes('Final'),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code==200,response.text
    assets=AssistantRepository(service.db).list_client_assets(cid,document['id'])
    assert len(assets)==2
    from stock_sync_web.assistant_service import DocumentAssistantService
    assistant=DocumentAssistantService(service)
    saved=assistant.read_asset({'id':users['boris']['id'],'role':'user'},assets[0]['id'])[0]
    assert saved.startswith(b'PK')
    with patch.object(AssistantRepository,'delete_document',side_effect=OSError('failed-db-delete')):
        response=http.delete(f'/api/documents/{document["id"]}',headers=users['admin']['headers'])
        assert response.status_code==503
    assert assistant.read_asset({'id':users['boris']['id'],'role':'user'},assets[0]['id'])[0]==saved
    assert http.delete(f'/api/documents/{document["id"]}',headers=users['admin']['headers']).status_code==200
    assert not old_path.exists(), 'Legacy original was copied into immutable versions, but old unreferenced path leaked'
    assert not list(assistant.store.staging.iterdir())

def test_legacy_onec_document_reuses_linked_card_without_new_client(api):
    http,service,users,cid=api
    with service.db.transaction() as conn:
        conn.execute("INSERT INTO counterparties(id,onec_key,name,full_name,updated_at) VALUES(901,'synthetic-901','Program name','Legal name','2026-10-30')")
        conn.execute('UPDATE crm_clients SET linked_counterparty_id=901 WHERE id=?',(cid,))
    output=service.document_exports_dir/'legacy.docx'; output.write_bytes(docx_bytes('Old linked contract'))
    did=service.db.create_document(document_type='contract',number='42',client_source='onec',counterparty_id=901,crm_client_id=None,commercial_offer_id=None,client_name='Legal name',document_date='2026-10-30',output_path=service._store_path(output),notes='',missing_fields='[]',created_by_user_id=users['alice']['id'])
    summary=http.get(f'/api/crm/clients/{cid}/documents',headers=users['boris']['headers'])
    assert summary.status_code==200,summary.text
    assert summary.json()['contract']['id']==did
    journal=http.get('/api/documents',headers=users['boris']['headers']).json()['items']
    assert journal[0]['crmClientId']==cid
    assert len(service.db.list_crm_clients())==1
