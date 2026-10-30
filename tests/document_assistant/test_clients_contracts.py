from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest

from stock_sync_desktop.onec_api import OneCClientError
from tests.document_assistant.test_storage_flow import api, docx_bytes


FIELDS = {'document_name':'ООО Синтетический покупатель','full_name':'Общество Синтетический покупатель',
    'inn':'7701234567','kpp':'770101001','signer_name':'Иванов Иван Иванович',
    'signer_position':'Директор','signer_basis':'Устава','legal_address':'Тестовый адрес'}


class OneCDouble:
    def __init__(self, existing=None, failure=None):
        self.existing=existing; self.failure=failure; self.posts=[]; self.lookups=[];self.updates=[]

    def find_counterparty_by_identity(self, **identity):
        self.lookups.append(identity)
        return self.existing

    def create_counterparty(self, row):
        self.posts.append(dict(row))
        if self.failure: raise self.failure
        return {'Ref_Key':'synthetic-ref-'+row['inn'],'Description':'PROGRAM NAME',
            'НаименованиеПолное':row['document_name'],'ИНН':row['inn'],'КПП':row.get('kpp') or ''}

    def update_counterparty_from_card(self, ref_key, row):
        self.updates.append((ref_key,dict(row)))
        return self.existing


def requisites(http, users, monkeypatch, fields=None, cid=None):
    from stock_sync_web.assistant_extraction import ExtractionResult
    payload=fields or FIELDS
    monkeypatch.setattr('stock_sync_web.assistant_extraction.CodexDocumentExtractor.extract',
        lambda self,request:ExtractionResult('requisites',dict(payload),[{'assetId':request.source_asset_ids[0],'method':'double','reviewed':False}]))
    if cid is None:
        asset=http.post('/api/document-assistant/intake',headers=users['alice']['headers'],data={'kind':'requisites'},
            files={'file':('requisites.docx',docx_bytes('Синтетические реквизиты'))}).json()
    else:
        asset=http.post(f'/api/crm/clients/{cid}/document-assets',headers=users['alice']['headers'],data={'role':'requisites'},
            files={'file':('requisites.docx',docx_bytes('Синтетические реквизиты'))}).json()
    response=http.post('/api/document-assistant/jobs',headers=users['alice']['headers'],json={
        'kind':'requisites','assetIds':[asset['id']],**({'clientId':cid} if cid else {})})
    assert response.status_code==200,response.text
    return http.get('/api/document-assistant/jobs/'+response.json()['id'],headers=users['alice']['headers']).json(),asset


def confirm(http, users, job, action='create_client', fields=None, key='create-client-action-0001', user='alice'):
    return http.post('/api/document-assistant/jobs/'+job['id']+'/confirm',headers=users[user]['headers'],json={
        'version':job['version'],'operationKey':key,'action':action,'fields':fields or job['result']['payload']})


def test_reviewed_creation_uses_real_client_service_once_and_binds_sources(api,monkeypatch):
    http,service,users,_=api
    remote=OneCDouble(); monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,asset=requisites(http,users,monkeypatch)
    response=confirm(http,users,job)
    assert response.status_code==200,response.text
    result=response.json(); cid=result['clientId']
    assert result['resultType']=='client' and result['syncStatus']=='synced'
    assert service.db.get_crm_client(cid)['linked_counterparty_id'] is not None
    assert remote.posts[0]['document_name']==FIELDS['document_name'] and len(remote.posts)==1
    assert result['job']['result']['extracted_payload']==FIELDS
    assert result['job']['result']['review']['actorUserId']==users['alice']['id']
    assert http.get(f'/api/document-assets/{asset["id"]}/download',headers=users['boris']['headers']).status_code==200
    replay=confirm(http,users,job)
    assert replay.status_code==200 and replay.json()==result and len(remote.posts)==1
    assert confirm(http,users,job,key='new-key-same-job').status_code==409


def test_generated_contract_snapshot_and_formatting_preserve_one_canonical_root(api,monkeypatch):
    http,service,users,cid=api
    job,_=requisites(http,users,monkeypatch,cid=cid)
    fields={**FIELDS,'contract_number':'Д-37','contract_date':'2026-10-30'}
    response=confirm(http,users,job,'generate_contract',fields)
    assert response.status_code==200,response.text
    did=response.json()['resultId']
    summary=http.get(f'/api/crm/clients/{cid}/documents',headers=users['boris']['headers']).json()
    assert summary['contract']['id']==did and summary['contract']['version']==1
    generated=http.get(f'/api/documents/{did}/download',headers=users['boris']['headers']).content
    document=Document(BytesIO(generated)); intro=document.paragraphs[3]
    assert 'Иванова Ивана Ивановича' in intro.text and FIELDS['document_name'] in intro.text
    assert '{{' not in '\n'.join(p.text for p in document.paragraphs)
    for run in intro.runs:
        if not run.text.strip(): continue
        is_name=run.text in [FIELDS['document_name'],'Иванова Ивана Ивановича','ООО "СМ ТЕХНО"','Дыбаева Александра Александровича']
        assert bool(run.bold)==is_name,(run.text,run.bold)
    persisted=service.db.get_crm_client(cid)
    assert persisted['document_name']=='ООО Тест документов' and not persisted.get('signer_name')
    from stock_sync_web.assistant_repository import AssistantRepository
    snapshot=json.loads(AssistantRepository(service.db).get_document(did)['legal_snapshot_json'])
    assert snapshot['document_name']==FIELDS['document_name'] and snapshot['signer_name']==FIELDS['signer_name']
    newer,_=requisites(http,users,monkeypatch,cid=cid)
    changed=confirm(http,users,newer,'generate_contract',{**fields,'contract_number':'OTHER'},key='different-number')
    assert changed.status_code==409
    result=confirm(http,users,newer,'generate_contract',{**fields,'bank_name':'Новый тестовый банк'},key='new-contract-version')
    assert result.status_code==200,result.text
    assert result.json()['resultId']==did
    summary=http.get(f'/api/crm/clients/{cid}/documents',headers=users['boris']['headers']).json()
    assert len(summary['documents'])==1 and summary['contract']['version']==2
    old=summary['contract']['assets'][0]
    assert http.get(f'/api/document-assets/{old["id"]}/download',headers=users['boris']['headers']).content==generated
    assert http.get(f'/api/crm/clients/{cid}',headers=users['boris']['headers']).status_code==403


@pytest.mark.parametrize('controls',[{'contract_number':{}},{'contract_date':[]},{'is_buyer':'false'},{'legal_type':[]},{'owner_user_id':1}])
def test_creation_controls_are_validated_before_side_effects(api,monkeypatch,controls):
    http,service,users,_=api
    remote=OneCDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    result=confirm(http,users,job,fields={**FIELDS,**controls})
    assert result.status_code==422,result.text
    assert not remote.posts and len(service.db.list_crm_clients())==1


@pytest.mark.parametrize('mode',['local','cached','remote'])
def test_exact_existing_identity_reuses_minimal_documents_card_without_private_merge(api,monkeypatch,mode):
    http,service,users,_=api
    remote=OneCDouble()
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    local_id=None
    if mode=='local':
        local=service.db.create_crm_client_card({**FIELDS,'notes':'OTHER MANAGER PRIVATE'},owner_user_id=users['boris']['id']);local_id=local['id']
    elif mode=='cached':
        service.db.upsert_counterparties([{'onec_key':'cached-ref','name':'PROGRAM','full_name':'ООО Из справочника','inn':FIELDS['inn'],'kpp':FIELDS['kpp']}])
    else:
        remote.existing={'Ref_Key':'remote-ref','Description':'PROGRAM','НаименованиеПолное':'ООО Удалённое','ИНН':FIELDS['inn'],'КПП':FIELDS['kpp']}
    job,_=requisites(http,users,monkeypatch)
    result=confirm(http,users,job)
    assert result.status_code==200,result.text
    cid=result.json()['clientId']
    assert not remote.posts and len(service.db.list_crm_clients())==2
    assert 'OTHER MANAGER PRIVATE' not in result.text
    if local_id:
        assert cid==local_id and service.db.get_crm_client(cid)['crm_owner_user_id']==users['boris']['id']
        assert service.db.get_crm_client(cid)['linked_counterparty_id'] is None
        assert http.get(f'/api/crm/clients/{cid}',headers=users['alice']['headers']).status_code==403


def test_concurrent_confirmations_different_jobs_and_keys_do_not_duplicate_remote_post(api,monkeypatch):
    http,service,users,_=api
    remote=OneCDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    jobs=[requisites(http,users,monkeypatch)[0] for _ in range(2)]
    with ThreadPoolExecutor(max_workers=2) as workers:
        results=list(workers.map(lambda pair:confirm(http,users,pair[1],key=f'parallel-confirm-{pair[0]}'),enumerate(jobs)))
    assert [r.status_code for r in results]==[200,200],[r.text for r in results]
    assert results[0].json()['clientId']==results[1].json()['clientId'] and len(remote.posts)==1
    job,_=requisites(http,users,monkeypatch,fields={**FIELDS,'inn':'7712345678'})
    with ThreadPoolExecutor(max_workers=2) as workers:
        responses=list(workers.map(lambda _:confirm(http,users,job,key='parallel-same-job'),range(2)))
    assert [r.status_code for r in responses]==[200,200] and responses[0].json()==responses[1].json()
    assert len(remote.posts)==2


def test_unknown_remote_post_is_durable_and_not_blindly_retried(api,monkeypatch):
    http,service,users,_=api
    remote=OneCDouble(failure=OneCClientError('network result unknown'))
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    response=confirm(http,users,job);assert response.status_code==200,response.text
    cid=response.json()['clientId'];assert response.json()['syncStatus']=='sync_error'
    assert confirm(http,users,job).json()==response.json() and len(remote.posts)==1
    retry=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    assert retry.status_code==409,retry.text
    assert len(remote.posts)==1


def test_missing_credentials_retains_existing_safe_retry_flow(api,monkeypatch):
    http,service,users,_=api
    def unavailable(**kwargs): raise ValueError('Настройте личные реквизиты подключения 1С.')
    monkeypatch.setattr(service,'build_user_client',unavailable)
    job,_=requisites(http,users,monkeypatch)
    response=confirm(http,users,job);assert response.status_code==200,response.text
    assert response.json()['syncStatus']=='sync_error'
    remote=OneCDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    cid=response.json()['clientId']
    retried=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    assert retried.status_code==200,retried.text
    assert retried.json()['sync']['status']=='synced' and len(remote.posts)==1


def test_retry_after_missing_credentials_also_guards_an_unknown_post(api,monkeypatch):
    http,service,users,_=api
    def unavailable(**kwargs): raise ValueError('Нет личных реквизитов подключения.')
    monkeypatch.setattr(service,'build_user_client',unavailable)
    job,_=requisites(http,users,monkeypatch)
    response=confirm(http,users,job);assert response.status_code==200,response.text
    remote=OneCDouble(failure=OneCClientError('unknown retry outcome'))
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    cid=response.json()['clientId']
    first=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    assert first.status_code==200 and first.json()['sync']['status']=='sync_error',first.text
    again=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    assert again.status_code==409,again.text
    assert len(remote.posts)==1


def test_ip_empty_kpp_and_optional_contract_fields_do_not_block_creation(api,monkeypatch):
    http,service,users,_=api
    remote=OneCDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    fields={'document_name':'ИП Синтетический','inn':'770123456789','kpp':None}
    job,_=requisites(http,users,monkeypatch,fields=fields)
    result=confirm(http,users,job,'generate_contract',{**fields,'legal_type':'individual_entrepreneur','contract_number':'ИП-1','contract_date':'2026-10-30'})
    assert result.status_code==200,result.text
    assert len(remote.posts)==1 and not remote.posts[0].get('kpp')
    cid=result.json()['clientId'];assert not service.db.get_crm_client(cid)['kpp']
    from stock_sync_web.assistant_repository import AssistantRepository
    missing=json.loads(AssistantRepository(service.db).get_document(result.json()['resultId'])['missing_fields'])
    assert 'signer_name' in missing and 'bank_account' in missing and 'kpp' not in missing


def test_manual_private_submission_fields_never_enter_shared_review_or_contract_snapshot(api,monkeypatch):
    http,service,users,_=api
    remote=OneCDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    private={'notes':'PRIVATE NOTE CANARY','email_note':'PRIVATE EMAIL CANARY','phone_note':'PRIVATE PHONE CANARY'}
    result=confirm(http,users,job,fields={**FIELDS,**private,'is_inactive':False,'bank_name_or_bik':'044525225'})
    assert result.status_code==200,result.text
    cid=result.json()['clientId'];saved=service.db.get_crm_client(cid)
    assert all(saved[key]==value for key,value in private.items())
    assert remote.posts[0]['notes']==private['notes']
    shared=http.get('/api/document-assistant/jobs/'+job['id'],headers=users['boris']['headers'])
    assert shared.status_code==200
    listing=http.get(f'/api/document-assistant/jobs?clientId={cid}',headers=users['boris']['headers'])
    documents=http.get(f'/api/crm/clients/{cid}/documents',headers=users['boris']['headers'])
    for response in [result,shared,listing,documents]:
        assert all(value not in response.text for value in private.values())
    bad=confirm(http,users,shared.json(),'generate_contract',fields={**FIELDS,**private,'contract_number':'37','contract_date':'2026-10-30'},key='private-not-legal')
    assert bad.status_code==422
    assert confirm(http,users,job,fields={**FIELDS,**private,'notes':'changed'}).status_code==409


def test_partial_remote_success_keeps_link_and_prevents_another_post(api,monkeypatch):
    from stock_sync_desktop.onec_api import OneCCounterpartySyncError
    http,service,users,_=api
    partial={'Ref_Key':'partial-created-ref','НаименованиеПолное':FIELDS['document_name'],'ИНН':FIELDS['inn'],'КПП':FIELDS['kpp']}
    remote=OneCDouble(failure=OneCCounterpartySyncError('bank step failed',partial))
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    result=confirm(http,users,job)
    assert result.status_code==200 and result.json()['syncStatus']=='sync_error',result.text
    cid=result.json()['clientId'];assert service.db.get_crm_client(cid)['linked_counterparty_id'] is not None
    assert http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers']).status_code==409
    assert confirm(http,users,job).json()==result.json() and len(remote.posts)==1


def test_reviewed_legal_names_are_literal_docx_text(api,monkeypatch):
    http,service,users,cid=api
    fields={**FIELDS,'document_name':'ООО «Тест & Партнёры <Один>»'}
    job,_=requisites(http,users,monkeypatch,fields=fields,cid=cid)
    result=confirm(http,users,job,'generate_contract',{**fields,'contract_number':'L-37','contract_date':'2026-10-30'})
    assert result.status_code==200,result.text
    content=http.get(f'/api/documents/{result.json()["resultId"]}/download',headers=users['alice']['headers']).content
    document=Document(BytesIO(content))
    assert fields['document_name'] in document.paragraphs[3].text


def test_partial_known_ref_retains_existing_supplementary_update_retry_without_post(api,monkeypatch):
    from stock_sync_desktop.onec_api import OneCCounterpartySyncError
    http,service,users,_=api
    known={'Ref_Key':'partial-known-ref','НаименованиеПолное':FIELDS['document_name'],'ИНН':FIELDS['inn'],'КПП':FIELDS['kpp']}
    remote=OneCDouble(failure=OneCCounterpartySyncError('bank step failed',known))
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    created=confirm(http,users,job);assert created.status_code==200,created.text
    cid=created.json()['clientId']
    remote.existing=known
    retried=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    assert retried.status_code==200,retried.text
    assert retried.json()['sync']['status']=='synced'
    assert len(remote.posts)==1 and len(remote.updates)==1 and remote.updates[0][0]=='partial-known-ref'


def test_partial_known_ref_cannot_post_when_remote_identity_points_elsewhere(api,monkeypatch):
    from stock_sync_desktop.onec_api import OneCCounterpartySyncError
    http,service,users,_=api
    known={'Ref_Key':'partial-known-ref','НаименованиеПолное':FIELDS['document_name'],'ИНН':FIELDS['inn'],'КПП':FIELDS['kpp']}
    remote=OneCDouble(failure=OneCCounterpartySyncError('bank step failed',known))
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    created=confirm(http,users,job);assert created.status_code==200,created.text
    cid=created.json()['clientId'];remote.existing={**known,'Ref_Key':'different-ref'}
    assert http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers']).status_code==409
    remote.existing=None
    assert http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers']).status_code==409
    assert len(remote.posts)==1 and not remote.updates


@pytest.mark.parametrize('guarded_retry',[False,True])
def test_ambiguous_remote_identity_never_chooses_a_card_or_updates_a_guarded_reference(api,monkeypatch,guarded_retry):
    from stock_sync_desktop.onec_api import OneCClient,OneCCounterpartySyncError
    from urllib.parse import parse_qs,urlsplit
    http,service,users,_=api
    known={'Ref_Key':'known-reference','НаименованиеПолное':FIELDS['document_name'],'ИНН':FIELDS['inn'],'КПП':FIELDS['kpp']}
    if guarded_retry:
        first=OneCDouble(failure=OneCCounterpartySyncError('bank step failed',known))
        monkeypatch.setattr(service,'build_user_client',lambda **kw:first)
        job,_=requisites(http,users,monkeypatch)
        created=confirm(http,users,job);assert created.status_code==200,created.text
        cid=created.json()['clientId']
    class ODataDouble(OneCClient):
        def __init__(self):
            super().__init__('http://onec.example','test-user','test-password')
            self.writes=[];self.request_limits=[]
        def _request(self,method,endpoint_or_url,payload=None):
            assert method=='GET','Only synthetic read is allowed'
            top=int(parse_qs(urlsplit(endpoint_or_url).query).get('$top',['2'])[0])
            self.request_limits.append(top)
            return {'value':[known,{**known,'Ref_Key':'another-reference'}][:top]}
        def update_counterparty_from_card(self,ref_key,row):
            self.writes.append(ref_key);return known
        def create_counterparty(self,row):
            self.writes.append('NEW POST');return known
    remote=ODataDouble();monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    if guarded_retry:
        result=http.post(f'/api/clients/{cid}/send-to-onec',headers=users['alice']['headers'])
    else:
        job,_=requisites(http,users,monkeypatch)
        result=confirm(http,users,job)
    assert result.status_code==409,result.text
    assert not remote.writes and remote.request_limits==[2]
    # Legacy callers keep the exact existing one-result request/return semantics.
    legacy=remote.find_counterparty_by_identity(legal_type='legal_entity',inn=FIELDS['inn'],kpp=FIELDS['kpp'])
    assert legacy==known and remote.request_limits[-1]==1


@pytest.mark.parametrize('malformed',[{'Ref_Key':''},{'ИНН':'7700000000'},{'КПП':'770102002'},{'ИНН':7701234567}])
def test_strict_remote_identity_rejects_malformed_candidates_before_client_creation(api,monkeypatch,malformed):
    from stock_sync_desktop.onec_api import OneCClient
    http,service,users,_=api
    candidate={'Ref_Key':'candidate-ref','ИНН':FIELDS['inn'],'КПП':FIELDS['kpp'],**malformed}
    remote=OneCClient('http://onec.example','test-user','test-password')
    monkeypatch.setattr(remote,'_request',lambda *args,**kwargs:{'value':[candidate]})
    monkeypatch.setattr(service,'build_user_client',lambda **kw:remote)
    job,_=requisites(http,users,monkeypatch)
    response=confirm(http,users,job)
    assert response.status_code==409,response.text
    assert len(service.db.list_crm_clients())==1
