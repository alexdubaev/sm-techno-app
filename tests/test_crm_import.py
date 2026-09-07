from io import BytesIO
import sqlite3

import pytest
from openpyxl import Workbook

from stock_sync_web.crm_repository import CrmRepository
from stock_sync_web.database import WebDatabase


def workbook(clients, contacts=()):
    book = Workbook()
    sheet = book.active
    sheet.title = "Клиенты"
    headers = ["Компания", "ИНН", "КПП", "Город", "Сайт", "Основной контакт", "Телефон", "Почта", "Telegram", "MAX", "Комментарий", "__crm_client_id", "__color_key", "__export_version"]
    sheet.append(headers)
    for row in clients:
        sheet.append([row.get(key) for key in headers])
    sheet = book.create_sheet("Контакты")
    headers = ["Компания", "Контактное лицо", "Телефон", "Почта", "Основной контакт", "__crm_client_id", "__crm_contact_id"]
    sheet.append(headers)
    for row in contacts:
        sheet.append([row.get(key) for key in headers])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def legacy_workbook(clients, contacts=()):
    """Pre-round-trip export shape: no hidden metadata columns."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Клиенты"
    client_headers = ["Компания", "ИНН", "КПП", "Город", "Сайт", "Основной контакт", "Телефон", "Почта", "Telegram", "MAX", "Комментарий"]
    sheet.append(client_headers)
    for row in clients:
        sheet.append([row.get(key) for key in client_headers])
    sheet = book.create_sheet("Контакты")
    contact_headers = ["Компания", "Контактное лицо", "Телефон", "Почта", "Основной контакт"]
    sheet.append(contact_headers)
    for row in contacts:
        sheet.append([row.get(key) for key in contact_headers])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
def setup(tmp_path):
    db = WebDatabase(tmp_path / "import.db")
    owner = db.create_user(username="owner", password="password", role="user")
    other = db.create_user(username="other", password="password", role="admin")
    repo = CrmRepository(db)
    tab = repo.ensure_work_tab_for_actor(actor_id=owner, owner_id=owner)
    return db, repo, owner, other, tab["id"]


def run(setup, content, *, final=False, **kwargs):
    _, repo, owner, _, tab = setup
    params = dict(actor_id=owner, owner_id=owner, content=content, target_tab_id=tab)
    params.update(kwargs)
    method = repo.import_excel_for_actor if final else repo.preview_excel_import_for_actor
    return method(**params)


def rows(db, table):
    with db.connect() as conn:
        return [dict(row) for row in conn.execute(f"SELECT * FROM {table} ORDER BY rowid")]


def snapshot(db):
    return {table: rows(db, table) for table in ("crm_tabs", "crm_clients", "crm_contacts", "crm_assignments", "crm_row_preferences", "crm_sync_jobs", "counterparties")}


def test_preview_is_read_only_and_final_creates_local_card_contact_and_colour(setup, monkeypatch):
    db, _, owner, _, _ = setup
    content = workbook([{"Компания": "Альфа", "ИНН": "0012345678", "__color_key": "green"}], [{"Компания": "Альфа", "Контактное лицо": "Анна", "Телефон": "+7 (900) 123-45-67", "Основной контакт": "Да"}])
    before = snapshot(db)
    original = db.transaction
    def forbidden():
        pytest.fail("preview opened a write transaction")
    monkeypatch.setattr(db, "transaction", forbidden)
    preview = run(setup, content, target_tab_id=None, new_tab_name="Импорт")
    assert (preview["clientsToCreate"], preview["contactsToCreate"], preview["clientsToAssign"]) == (1, 1, 1)
    assert snapshot(db) == before
    monkeypatch.setattr(db, "transaction", original)
    result = run(setup, content, final=True, target_tab_id=None, new_tab_name="Импорт")
    assert result["targetTab"]["name"] == "Импорт"
    card = rows(db, "crm_clients")[0]
    assert (card["inn"], card["sync_status"], card["linked_counterparty_id"], card["crm_owner_user_id"]) == ("0012345678", "local", None, owner)
    assert rows(db, "crm_row_preferences")[0]["color_key"] == "green"
    second = run(setup, content, final=True, target_tab_id=result["targetTab"]["id"])
    assert (second["clientsToCreate"], second["contactsToCreate"], second["unchangedClients"]) == (0, 0, 1)
    assert len(rows(db, "crm_contacts")) == 1


def test_name_and_phone_only_client_is_imported_as_local_card(setup):
    db, _, owner, _, _ = setup
    result = run(setup, workbook([{"Компания": "Только имя и телефон", "Телефон": "+7 (900) 123-45-67"}]), final=True)
    assert result["clientsToCreate"] == 1
    card = rows(db, "crm_clients")[0]
    assert (card["document_name"], card["phone"], card["crm_owner_user_id"], card["sync_status"], card["linked_counterparty_id"]) == (
        "Только имя и телефон", "+7 (900) 123-45-67", owner, "local", None,
    )


def test_import_preserves_telegram_and_max_contacts(setup):
    """Ignoring optional messenger columns would make a round-trip lose contact links."""
    db, _, _, _, _ = setup
    result = run(setup, workbook([{
        "Компания": "Мессенджер-клиент",
        "Телефон": "+7 900 000-00-01",
        "Telegram": "@messenger_client",
        "MAX": "https://max.ru/messenger_client",
    }]), final=True)

    assert result["clientsToCreate"] == 1
    card = rows(db, "crm_clients")[0]
    assert (card["telegram"], card["max_link"]) == ("@messenger_client", "https://max.ru/messenger_client")


def test_legacy_workbook_without_hidden_columns_is_supported(setup):
    db, repo, owner, _, _ = setup
    existing = repo.create_local_client(actor_id=owner, values={"document_name": "Старый экспорт", "inn": "00123", "phone": "111"})
    result = run(setup, legacy_workbook([{"Компания": "Новое имя", "ИНН": "00123", "Телефон": "222"}]), final=True)
    assert result["clientsToUpdate"] == 1
    assert result["clientsToCreate"] == 0
    assert rows(db, "crm_clients")[0]["id"] == existing["id"]
    assert rows(db, "crm_clients")[0]["document_name"] == "Новое имя"


def test_hidden_id_has_priority_and_blank_cells_preserve_fields(setup):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "Старое", "inn": "123", "phone": "999", "city": "Москва"})
    other = repo.create_local_client(actor_id=owner, values={"document_name": "Другое", "inn": "456"})
    result = run(setup, workbook([{"__crm_client_id": card["id"], "Компания": "Новое", "ИНН": "456"}]), final=True)
    assert result["clientsToUpdate"] == 1
    actual = rows(db, "crm_clients")
    assert (actual[0]["document_name"], actual[0]["phone"], actual[0]["city"]) == ("Новое", "999", "Москва")
    assert actual[0]["name"] == "Новое"
    assert actual[1]["id"] == other["id"]


def test_inn_kpp_phone_and_email_matching_without_technical_identifiers(setup):
    db, repo, owner, _, _ = setup
    for kpp in ("1", "2"):
        repo.create_local_client(actor_id=owner, values={"document_name": "Филиал", "inn": "123", "kpp": kpp})
    repo.create_local_client(actor_id=owner, values={"document_name": "  Альфа  Сервис ", "phone": "+7 (900) 123-45-67"})
    preview = run(setup, legacy_workbook([
        {"Компания": "Филиал новый", "ИНН": "123", "КПП": "2"},
        {"Компания": "альфа  сервис", "Телефон": "79001234567"},
        {"Компания": "Переименованная Альфа", "Телефон": "79001234567"},
    ]))
    assert preview["clientsToUpdate"] == 2
    assert preview["clientsToCreate"] == 1
    assert not preview["errors"]


@pytest.mark.parametrize("client,code", [({"Компания": "Без реквизитов"}, "insufficient_identity"), ({"ИНН": "123"}, "ambiguous_client")])
def test_identity_errors_block_all_final_writes(setup, client, code):
    db, repo, owner, _, _ = setup
    for name in ("Один", "Два"):
        repo.create_local_client(actor_id=owner, values={"document_name": name, "inn": "123"})
    content = workbook([client, {"Компания": "Валидный", "ИНН": "999"}])
    before = snapshot(db)
    preview = run(setup, content)
    assert code in [e["code"] for e in preview["errors"]]
    with pytest.raises(ValueError):
        run(setup, content, final=True, target_tab_id=None, new_tab_name="Не создавать")
    assert snapshot(db) == before


def test_duplicate_workbook_rows_are_conflicts(setup):
    preview = run(setup, workbook([{"Компания": "А", "ИНН": "123"}, {"Компания": "Б", "ИНН": "123"}]))
    assert preview["duplicateConflicts"] == 1


def test_linked_cards_and_their_contacts_are_skipped_without_onec_access(setup, monkeypatch):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "Связанный", "inn": "123"})
    with db.transaction() as conn:
        conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (999, 'k', '1C', 'now')")
        conn.execute("UPDATE crm_clients SET linked_counterparty_id = 999 WHERE id = ?", (card["id"],))
    before = snapshot(db)
    original = db.connect
    observed = []
    def guarded_connect():
        conn = original()
        def guard(action, table, column, *_):
            if table in {"counterparties", "crm_sync_jobs", "crm_client_sync_state", "crm_sync_conflicts"} or (action == sqlite3.SQLITE_UPDATE and column == "linked_counterparty_id"):
                observed.append((action, table, column))
                return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK
        conn.set_authorizer(guard)
        return conn
    monkeypatch.setattr(db, "connect", guarded_connect)
    content = workbook([{"__crm_client_id": card["id"], "Компания": "Изменить", "__color_key": "red"}], [{"__crm_client_id": card["id"], "Контактное лицо": "Контакт", "Телефон": "123"}])
    result = run(setup, content, final=True)
    assert result["skippedOneCLinked"] == 1
    assert not observed
    monkeypatch.setattr(db, "connect", original)
    assert snapshot(db) == before


def test_include_existing_false_updates_but_does_not_move_assignment(setup):
    db, repo, owner, _, tab = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    repo.assign_client_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], tab_id=tab)
    content = workbook([{"ИНН": "123", "Компания": "Б"}])
    result = run(setup, content, final=True, target_tab_id=None, new_tab_name="Новая", include_existing_clients=False)
    assert result["clientsToAssign"] == 0
    assert rows(db, "crm_assignments")[0]["tab_id"] == tab
    run(setup, content, final=True, target_tab_id=result["targetTab"]["id"])
    assert len(rows(db, "crm_assignments")) == 1
    assert rows(db, "crm_assignments")[0]["tab_id"] == result["targetTab"]["id"]


def test_contact_id_updates_and_preserves_blank_phone(setup):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    contact = repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Анна", phone="123")
    result = run(setup, workbook([{"__crm_client_id": card["id"]}], [{"__crm_client_id": card["id"], "__crm_contact_id": contact["id"], "Контактное лицо": "Анна новая", "Почта": "a@example.com"}]), final=True)
    assert result["contactsToUpdate"] == 1
    assert rows(db, "crm_contacts")[0]["phone"] == "123"


@pytest.mark.parametrize("target,name", [(None, None), (1, "Новая"), (0, None), (99999, None), (None, "Клиенты 1С")])
def test_invalid_targets_are_rejected(setup, target, name):
    with pytest.raises(ValueError):
        run(setup, workbook([]), target_tab_id=target, new_tab_name=name)


def test_owner_permissions_and_foreign_tab_and_card_are_guarded(setup):
    _, repo, owner, other, tab = setup
    content = workbook([])
    with pytest.raises(PermissionError):
        run(setup, content, actor_id=other)
    with pytest.raises(ValueError):
        run(setup, content, actor_id=other, owner_id=other, target_tab_id=tab)
    card = repo.create_local_client(actor_id=other, values={"document_name": "Чужой", "inn": "123"})
    preview = run(setup, workbook([{"__crm_client_id": card["id"], "Компания": "Взлом", "ИНН": "123"}]))
    assert preview["errors"]


def test_database_failure_rolls_back_every_import_write(setup):
    db, _, _, _, _ = setup
    with db.transaction() as conn:
        conn.execute("CREATE TRIGGER fail_contact BEFORE INSERT ON crm_contacts BEGIN SELECT RAISE(ABORT, 'injected failure'); END")
    before = snapshot(db)
    content = workbook([{"Компания": "А", "ИНН": "123", "__color_key": "red"}], [{"Компания": "А", "Контактное лицо": "Анна", "Телефон": "123"}])
    with pytest.raises(sqlite3.IntegrityError, match="injected failure"):
        run(setup, content, final=True, target_tab_id=None, new_tab_name="Откат")
    assert snapshot(db) == before


def test_malformed_workbook_is_rejected(setup):
    with pytest.raises(ValueError):
        run(setup, b"not xlsx")


def test_unmatched_kpp_falls_back_to_unique_inn(setup):
    _, repo, owner, _, _ = setup
    repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123", "kpp": "old"})
    preview = run(setup, workbook([{"Компания": "А", "ИНН": "123", "КПП": "new"}]))
    assert (preview["clientsToCreate"], preview["clientsToUpdate"]) == (0, 1)


def test_colour_restores_for_existing_target_even_when_assignment_inclusion_is_false(setup):
    db, repo, owner, _, tab = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    repo.assign_client_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], tab_id=tab)
    run(setup, workbook([{"ИНН": "123", "__color_key": "green"}]), final=True, include_existing_clients=False)
    assert rows(db, "crm_row_preferences")[0]["color_key"] == "green"


def test_deleted_source_ids_can_fall_back_to_identity_including_contacts(setup):
    db, _, _, _, _ = setup
    content = workbook([{"Компания": "А", "ИНН": "123", "__crm_client_id": 777}], [{"Компания": "А", "__crm_client_id": 777, "__crm_contact_id": 888, "Контактное лицо": "Анна", "Телефон": "123"}])
    result = run(setup, content, final=True)
    assert result["contactsToCreate"] == 1
    assert rows(db, "crm_contacts")[0]["crm_client_id"] == rows(db, "crm_clients")[0]["id"]


def test_new_and_updated_import_sql_never_enters_onec_or_outbox_paths(setup, monkeypatch):
    db, repo, owner, _, _ = setup
    repo.create_local_client(actor_id=owner, values={"document_name": "Старая", "inn": "123"})
    original = db.connect
    statements = []
    def traced_connect():
        conn = original()
        conn.set_trace_callback(statements.append)
        return conn
    monkeypatch.setattr(db, "connect", traced_connect)
    content = workbook([{"Компания": "Обновлённая", "ИНН": "123"}, {"Компания": "Новая", "ИНН": "456"}], [{"Компания": "Новая", "Контактное лицо": "Анна", "Телефон": "123"}])
    run(setup, content)
    result = run(setup, content, final=True)
    assert (result["clientsToCreate"], result["clientsToUpdate"], result["contactsToCreate"]) == (1, 1, 1)
    prohibited = ("counterparties", "crm_sync", "outbox", "linked_counterparty_id")
    assert not [sql for sql in statements if any(name in sql.casefold() for name in prohibited)]
    assert all(c["sync_status"] == "local" and c["linked_counterparty_id"] is None for c in rows(db, "crm_clients"))


def test_contact_ambiguity_and_wrong_client_id_block_import(setup):
    _, repo, owner, _, _ = setup
    first = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    second = repo.create_local_client(actor_id=owner, values={"document_name": "Б", "inn": "456"})
    for _ in range(2):
        contact = repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=first["id"], name="Анна", email="a@example.com")
    preview = run(setup, workbook([{"__crm_client_id": first["id"]}], [{"__crm_client_id": first["id"], "Контактное лицо": " АННА ", "Почта": "A@EXAMPLE.COM"}]))
    assert preview["errors"][0]["code"] == "ambiguous_contact"
    preview = run(setup, workbook([{"__crm_client_id": second["id"]}], [{"__crm_client_id": second["id"], "__crm_contact_id": contact["id"], "Контактное лицо": "Анна", "Телефон": "123"}]))
    assert preview["errors"][0]["code"] == "inaccessible_contact"


def test_formula_and_unknown_export_version_are_row_errors(setup):
    preview = run(setup, workbook([{"Компания": "=1+1", "ИНН": "123", "__export_version": "2"}]))
    assert {e["code"] for e in preview["errors"]} == {"formula_not_allowed", "unsupported_version"}
    assert all(e["sheet"] == "Клиенты" and e["row"] == 2 for e in preview["errors"])


def test_failure_rolls_back_existing_updates_moves_and_colours(setup):
    db, repo, owner, _, tab = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "Старое", "inn": "123"})
    repo.assign_client_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], tab_id=tab)
    with db.transaction() as conn:
        conn.execute("CREATE TRIGGER fail_contact_update BEFORE INSERT ON crm_contacts BEGIN SELECT RAISE(ABORT, 'injected failure'); END")
    before = snapshot(db)
    content = workbook([{"Компания": "Новое", "ИНН": "123", "__color_key": "blue"}], [{"Компания": "Новое", "Контактное лицо": "Анна", "Телефон": "123"}])
    with pytest.raises(sqlite3.IntegrityError, match="injected failure"):
        run(setup, content, final=True, target_tab_id=None, new_tab_name="Откат обновлений")
    assert snapshot(db) == before


def test_legacy_export_roundtrip_keeps_phone_and_contacts_without_duplicates(setup):
    from stock_sync_web.crm_export import build_crm_export_xlsx
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "00123", "phone": "+79001234567"})
    contact = repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Анна", phone="+79001234567")
    content = build_crm_export_xlsx(client_rows=[card], contact_rows=[dict(contact, client_id=card["id"])])
    result = run(setup, content, final=True)
    assert result["clientsToCreate"] == result["contactsToCreate"] == 0
    assert rows(db, "crm_clients")[0]["phone"] == "+79001234567"
    assert len(rows(db, "crm_contacts")) == 1


def test_corrupt_xml_is_a_workbook_validation_error(setup):
    from zipfile import ZipFile
    original = workbook([{"Компания": "А", "ИНН": "123"}])
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<broken" if name == "xl/worksheets/sheet1.xml" else source.read(name))
    with pytest.raises(ValueError):
        run(setup, output.getvalue())


def test_repeated_stale_source_id_is_rejected_without_writes(setup):
    db, _, _, _, _ = setup
    before = snapshot(db)
    content = workbook([{"Компания": "А", "ИНН": "123", "__crm_client_id": 777}, {"Компания": "Б", "ИНН": "456", "__crm_client_id": 777}], [{"Компания": "А", "__crm_client_id": 777, "Контактное лицо": "Анна", "Телефон": "123"}])
    preview = run(setup, content)
    assert preview["duplicateConflicts"] == 1
    assert preview["errors"][0]["code"] == "duplicate_source_client_id"
    with pytest.raises(ValueError):
        run(setup, content, final=True, target_tab_id=None, new_tab_name="Не создавать")
    assert snapshot(db) == before


def test_distinct_stale_ids_keep_contacts_with_correct_companies(setup):
    db, _, _, _, _ = setup
    content = workbook([{"Компания": "А", "ИНН": "123", "__crm_client_id": 777}, {"Компания": "Б", "ИНН": "456", "__crm_client_id": 888}], [{"Компания": "Б", "__crm_client_id": 888, "Контактное лицо": "Борис", "Телефон": "222"}, {"Компания": "А", "__crm_client_id": 777, "Контактное лицо": "Анна", "Телефон": "111"}])
    run(setup, content, final=True)
    company_names = {c["id"]: c["document_name"] for c in rows(db, "crm_clients")}
    assert {c["name"]: company_names[c["crm_client_id"]] for c in rows(db, "crm_contacts")} == {"Анна": "А", "Борис": "Б"}


@pytest.mark.parametrize("original,changed,repeated", [
    ({"document_name": "А", "inn": "123"}, {"ИНН": "456"}, {"ИНН": "456"}),
    ({"document_name": "А", "phone": "111"}, {"Компания": "Б", "Телефон": "222"}, {"Компания": "б", "Телефон": "2-2-2"}),
])
def test_client_duplicates_match_planned_identity_and_block_all_writes(setup, original, changed, repeated):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values=original)
    content = workbook([dict(changed, __crm_client_id=card["id"]), repeated])
    before = snapshot(db)
    preview = run(setup, content)
    assert preview["clientsToCreate"] == 0
    assert preview["errors"][0]["code"] == "duplicate_client"
    with pytest.raises(ValueError):
        run(setup, content, final=True)
    assert snapshot(db) == before


def test_contact_duplicates_match_planned_identity_and_block_all_writes(setup):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    contact = repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Анна", phone="111")
    content = workbook([{"__crm_client_id": card["id"]}], [{"__crm_client_id": card["id"], "__crm_contact_id": contact["id"], "Контактное лицо": "Борис", "Телефон": "222"}, {"__crm_client_id": card["id"], "Контактное лицо": "борис", "Телефон": "2-2-2"}])
    before = snapshot(db)
    preview = run(setup, content)
    assert preview["contactsToCreate"] == 0
    assert preview["errors"][0]["code"] == "duplicate_contact"
    with pytest.raises(ValueError):
        run(setup, content, final=True)
    assert snapshot(db) == before


@pytest.mark.parametrize("existing", [False, True])
def test_last_explicit_primary_contact_is_stable_on_reimport(setup, existing):
    db, repo, owner, _, _ = setup
    if existing:
        card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
        repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Анна", phone="111")
        repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Борис", phone="222", is_primary=True)
    content = workbook([{"Компания": "А", "ИНН": "123"}], [{"Компания": "А", "Контактное лицо": "Анна", "Телефон": "111", "Основной контакт": "Да"}, {"Компания": "А", "Контактное лицо": "Борис", "Телефон": "222", "Основной контакт": "Да"}])
    for index in range(3):
        result = run(setup, content, final=True)
        assert {c["name"]: c["is_primary"] for c in rows(db, "crm_contacts")} == {"Анна": 0, "Борис": 1}
        if index:
            assert result["contactsToUpdate"] == result["contactsToCreate"] == 0


def test_explicit_primary_wins_over_existing_primary_with_blank_source_cell(setup):
    db, repo, owner, _, _ = setup
    card = repo.create_local_client(actor_id=owner, values={"document_name": "А", "inn": "123"})
    repo.add_contact_for_actor(actor_id=owner, owner_id=owner, client_id=card["id"], name="Борис", phone="222", is_primary=True)
    content = workbook([{"Компания": "А", "ИНН": "123"}], [{"Компания": "А", "Контактное лицо": "Анна", "Телефон": "111", "Основной контакт": "Да"}, {"Компания": "А", "Контактное лицо": "Борис", "Телефон": "222"}])
    for _ in range(3):
        run(setup, content, final=True)
        assert {c["name"]: c["is_primary"] for c in rows(db, "crm_contacts")} == {"Анна": 1, "Борис": 0}
