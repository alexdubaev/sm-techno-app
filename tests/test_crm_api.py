from __future__ import annotations

import tempfile
import unittest
import os
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

os.environ.setdefault("SM_TECHNO_INITIAL_ADMIN_PASSWORD", "crm-api-test-password")

import stock_sync_api
from stock_sync_web.crm_repository import CrmRepository
from stock_sync_web.database import WebDatabase
from stock_sync_web.service import WebStockSyncService
from stock_sync_desktop.onec_api import OneCClientError


class CrmApiTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.service = WebStockSyncService(db=WebDatabase(Path(self.temp_dir.name) / "crm-api.db"))
        self.service.bootstrap()
        self.owner_id = self.service.db.create_user(username="owner", password="password", role="user")
        self.other_id = self.service.db.create_user(username="other", password="password", role="user")
        self.admin_id = self.service.db.create_user(username="admin-api", password="password", role="admin")
        self.original_service = stock_sync_api.SERVICE
        stock_sync_api.SERVICE = self.service
        self.current_user = {"id": self.owner_id, "role": "user", "username": "owner"}
        stock_sync_api.app.dependency_overrides[stock_sync_api._get_current_user] = lambda: self.current_user
        self.client = TestClient(stock_sync_api.app)

    def tearDown(self) -> None:
        self.client.close()
        stock_sync_api.app.dependency_overrides.clear()
        stock_sync_api.SERVICE = self.original_service
        self.temp_dir.cleanup()

    def as_user(self, user_id: int, role: str = "user") -> None:
        self.current_user = {"id": user_id, "role": role, "username": "test"}

    def create_open_sync_conflict(self) -> tuple[int, dict[str, object]]:
        created = self.client.post("/api/crm/clients", json={"documentName": "Базовое имя"}).json()
        client_id = created["client"]["id"]
        self.service.db.update_crm_client_sync_state(client_id, sync_status="synced", synced=True)
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET document_name = ? WHERE id = ?", ("Локальное имя", client_id))
        self.service.db.merge_crm_client_fields_from_counterparty(
            client_id,
            {"document_name": "Имя из 1С", "email": "", "phone": ""},
        )
        conflict = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts").json()["items"][0]
        return client_id, conflict

    def link_primary_client(self, client_id: int, counterparty_id: int) -> None:
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (?, ?, ?, ?)",
                (counterparty_id, f"onec-{counterparty_id}", f"Компания {counterparty_id}", "2026-09-04T00:00:00"),
            )
            conn.execute(
                "UPDATE crm_clients SET linked_counterparty_id = ?, sync_status = 'synced' WHERE id = ?",
                (counterparty_id, client_id),
            )

    def test_local_card_is_created_in_own_work_tab_without_onec(self) -> None:
        response = self.client.post("/api/crm/clients", json={"documentName": "Новый лид"})

        self.assertEqual(201, response.status_code)
        data = response.json()
        self.assertEqual("Новый лид", data["client"]["documentName"])
        self.assertEqual("local", data["client"]["syncStatus"])
        self.assertEqual("В работе", data["assignment"]["tabName"])
        self.assertIsNone(data["client"]["linkedCounterpartyId"])

    def test_non_admin_cannot_supply_another_owner_context(self) -> None:
        response = self.client.get(f"/api/crm/tabs?ownerId={self.other_id}")

        self.assertEqual(403, response.status_code)

    def test_admin_must_explicitly_choose_owner_and_can_read_its_workspace(self) -> None:
        repo = CrmRepository(self.service.db)
        card = repo.create_local_client(actor_id=self.owner_id, values={"document_name": "Лид владельца"})
        work = repo.ensure_work_tab_for_actor(actor_id=self.owner_id, owner_id=self.owner_id)
        repo.assign_client_for_actor(actor_id=self.owner_id, owner_id=self.owner_id, client_id=card["id"], tab_id=work["id"])
        self.as_user(self.admin_id, "admin")

        missing_context = self.client.get("/api/crm/tabs")
        response = self.client.get(f"/api/crm/tabs?ownerId={self.owner_id}")

        self.assertEqual(400, missing_context.status_code)
        self.assertEqual(200, response.status_code)
        self.assertEqual(self.owner_id, response.json()["ownerId"])
        self.assertEqual("В работе", response.json()["items"][0]["name"])

    def test_personal_operations_and_admin_archive_are_server_guarded(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид"}).json()
        client_id = created["client"]["id"]
        tab_id = created["assignment"]["tabId"]
        contact = self.client.post(f"/api/crm/clients/{client_id}/contacts", json={"name": "Ирина", "isPrimary": True})
        event = self.client.post(f"/api/crm/clients/{client_id}/events", json={"kind": "call", "body": "Позвонили"})
        reminder = self.client.post(f"/api/crm/clients/{client_id}/reminders", json={"dueAt": "2026-09-05T10:00:00"})
        denied = self.client.post(f"/api/crm/clients/{client_id}/archive", json={"reason": "Нет"})

        self.assertEqual(201, contact.status_code)
        self.assertEqual(201, event.status_code)
        self.assertEqual(201, reminder.status_code)
        self.assertEqual(403, denied.status_code)
        self.as_user(self.admin_id, "admin")
        archived = self.client.post(f"/api/crm/clients/{client_id}/archive?ownerId={self.owner_id}", json={"reason": "Дубликат"})
        audit = self.client.get(f"/api/crm/clients/{client_id}/audit?ownerId={self.owner_id}")

        self.assertEqual(200, archived.status_code)
        self.assertEqual(200, audit.status_code)
        self.assertEqual(self.admin_id, audit.json()["items"][0]["actorUserId"])
        self.assertEqual(self.owner_id, audit.json()["items"][0]["ownerUserId"])
        self.assertEqual(tab_id, created["assignment"]["tabId"])

    def test_owner_can_create_personal_tab_and_move_local_lead(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид для переноса"}).json()
        client_id = created["client"]["id"]

        tab_response = self.client.post("/api/crm/tabs", json={"name": "Перезвонить"})

        self.assertEqual(201, tab_response.status_code)
        tab = tab_response.json()["tab"]
        moved = self.client.post(f"/api/crm/clients/{client_id}/move", json={"tabId": tab["id"]})

        self.assertEqual(200, moved.status_code)
        self.assertEqual("Перезвонить", moved.json()["assignment"]["tabName"])

    def test_local_lead_is_not_readable_by_another_user_and_admin_needs_context(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Закрытый лид"}).json()
        client_id = created["client"]["id"]

        self.as_user(self.other_id)
        denied = self.client.get(f"/api/crm/clients/{client_id}")
        self.assertEqual(403, denied.status_code)

        self.as_user(self.admin_id, "admin")
        missing_context = self.client.get(f"/api/crm/clients/{client_id}")
        self.assertEqual(400, missing_context.status_code)

    def test_owner_can_read_the_current_sync_conflict_values(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Базовое имя"}).json()
        client_id = created["client"]["id"]
        self.service.db.update_crm_client_sync_state(client_id, sync_status="synced", synced=True)
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET document_name = ? WHERE id = ?", ("Локальное имя", client_id))
        self.service.db.merge_crm_client_fields_from_counterparty(
            client_id,
            {"document_name": "Имя из 1С", "email": "", "phone": ""},
        )

        listed = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts")

        self.assertEqual(200, listed.status_code, listed.text)
        self.assertEqual(
            [{"fieldName": "documentName", "localValue": "Локальное имя", "remoteValue": "Имя из 1С", "sourceVersion": 1}],
            [{key: item[key] for key in ("fieldName", "localValue", "remoteValue", "sourceVersion")} for item in listed.json()["items"]],
        )

    def test_other_owner_cannot_read_or_resolve_a_sync_conflict(self) -> None:
        client_id, conflict = self.create_open_sync_conflict()

        self.as_user(self.other_id)
        listed = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts")
        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve",
            json={"choice": "remote", "expectedUpdatedAt": conflict["updatedAt"]},
        )

        self.assertEqual(403, listed.status_code)
        self.assertEqual(403, resolved.status_code)

    def test_non_author_cannot_read_or_resolve_a_linked_card_sync_conflict(self) -> None:
        """A shared primary-list row does not make its field conflict public."""
        created = self.client.post("/api/crm/clients", json={"documentName": "Общий конфликт"}).json()
        client_id = created["client"]["id"]
        self.link_primary_client(client_id, 865)
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET document_name = ? WHERE id = ?", ("Локальное значение", client_id))
        self.service.db.merge_crm_client_fields_from_counterparty(
            client_id,
            {"document_name": "Значение из 1С", "email": "", "phone": ""},
        )
        conflict = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts").json()["items"][0]

        self.as_user(self.other_id)
        listed = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts")
        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve",
            json={"choice": "remote", "expectedUpdatedAt": conflict["updatedAt"]},
        )

        self.assertEqual(403, listed.status_code)
        self.assertEqual(403, resolved.status_code)

    def test_stale_sync_conflict_version_keeps_the_conflict_open(self) -> None:
        client_id, conflict = self.create_open_sync_conflict()

        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve",
            json={"choice": "remote", "expectedUpdatedAt": "устаревшая-версия"},
        )
        listed = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts")

        self.assertEqual(409, resolved.status_code)
        self.assertIn("уже изменён", resolved.json()["detail"])
        self.assertEqual([conflict["id"]], [item["id"] for item in listed.json()["items"]])

    def test_owner_can_resolve_current_sync_conflict_with_remote_value(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Базовое имя"}).json()
        client_id = created["client"]["id"]
        self.service.db.update_crm_client_sync_state(client_id, sync_status="synced", synced=True)
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET document_name = ? WHERE id = ?", ("Локальное имя", client_id))
        self.service.db.merge_crm_client_fields_from_counterparty(
            client_id,
            {"document_name": "Имя из 1С", "email": "", "phone": ""},
        )
        conflict = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts").json()["items"][0]

        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve",
            json={"choice": "remote", "expectedUpdatedAt": conflict["updatedAt"]},
        )

        self.assertEqual(200, resolved.status_code, resolved.text)
        self.assertEqual("Имя из 1С", resolved.json()["client"]["documentName"])
        self.assertEqual("synced", resolved.json()["client"]["syncStatus"])
        self.assertEqual([], self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts").json()["items"])
        audit = self.client.get(f"/api/crm/clients/{client_id}/audit")
        with self.service.db.connect() as conn:
            conflict_row = conn.execute("SELECT status, resolved_value_json, resolved_by_user_id FROM crm_sync_conflicts WHERE id = ?", (conflict["id"],)).fetchone()
            snapshot = conn.execute("SELECT last_synced_snapshot FROM crm_sync_state WHERE crm_client_id = ?", (client_id,)).fetchone()
        self.assertEqual(("resolved_remote", '"Имя из 1С"', self.owner_id), tuple(conflict_row))
        self.assertEqual("Имя из 1С", json.loads(snapshot["last_synced_snapshot"])["document_name"])
        self.assertEqual([(self.owner_id, self.owner_id, "resolve_sync_conflict", "remote")], [
            (item["actorUserId"], item["ownerUserId"], item["action"], item["reason"])
            for item in audit.json()["items"]
        ])

    def test_owner_can_keep_local_value_when_safe_remote_write_is_unavailable(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Базовое имя"}).json()
        client_id = created["client"]["id"]
        self.service.db.update_crm_client_sync_state(client_id, sync_status="synced", synced=True)
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET document_name = ? WHERE id = ?", ("Локальное имя", client_id))
        self.service.db.merge_crm_client_fields_from_counterparty(
            client_id,
            {"document_name": "Имя из 1С", "email": "", "phone": ""},
        )
        conflict = self.client.get(f"/api/crm/clients/{client_id}/sync-conflicts").json()["items"][0]

        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve",
            json={"choice": "local", "expectedUpdatedAt": conflict["updatedAt"]},
        )

        self.assertEqual(200, resolved.status_code, resolved.text)
        self.assertEqual("Локальное имя", resolved.json()["client"]["documentName"])
        self.assertEqual("blocked_capability", resolved.json()["client"]["syncStatus"])
        audit = self.client.get(f"/api/crm/clients/{client_id}/audit")
        with self.service.db.connect() as conn:
            conflict_row = conn.execute("SELECT status, resolved_value_json, resolved_by_user_id FROM crm_sync_conflicts WHERE id = ?", (conflict["id"],)).fetchone()
            snapshot = conn.execute("SELECT last_synced_snapshot FROM crm_sync_state WHERE crm_client_id = ?", (client_id,)).fetchone()
        self.assertEqual(("resolved_local", '"Локальное имя"', self.owner_id), tuple(conflict_row))
        self.assertEqual("Имя из 1С", json.loads(snapshot["last_synced_snapshot"])["document_name"])
        self.assertEqual([(self.owner_id, self.owner_id, "resolve_sync_conflict", "local")], [
            (item["actorUserId"], item["ownerUserId"], item["action"], item["reason"])
            for item in audit.json()["items"]
        ])

    def test_admin_resolution_audit_keeps_the_actual_actor_and_selected_owner(self) -> None:
        client_id, conflict = self.create_open_sync_conflict()
        self.as_user(self.admin_id, "admin")

        resolved = self.client.post(
            f"/api/crm/clients/{client_id}/sync-conflicts/{conflict['id']}/resolve?ownerId={self.owner_id}",
            json={"choice": "remote", "expectedUpdatedAt": conflict["updatedAt"]},
        )
        audit = self.client.get(f"/api/crm/clients/{client_id}/audit?ownerId={self.owner_id}")

        self.assertEqual(200, resolved.status_code, resolved.text)
        self.assertEqual([(self.admin_id, self.owner_id, "resolve_sync_conflict", "remote")], [
            (item["actorUserId"], item["ownerUserId"], item["action"], item["reason"])
            for item in audit.json()["items"]
        ])

    def test_owner_can_save_palette_preference_for_own_row(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид с цветом"}).json()
        client_id = created["client"]["id"]
        tab_id = created["assignment"]["tabId"]

        response = self.client.put(
            f"/api/crm/clients/{client_id}/row-preference",
            json={"tabId": tab_id, "colorKey": "blue", "position": 5},
        )

        self.assertEqual(200, response.status_code)
        self.assertEqual("blue", response.json()["preference"]["colorKey"])
        self.assertEqual(5, response.json()["preference"]["position"])
        listed = self.client.get(f"/api/crm/clients?tabId={tab_id}")
        self.assertEqual("blue", listed.json()["items"][0]["rowPreference"]["colorKey"])

    def test_primary_color_update_preserves_existing_position(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Связанная компания"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (601, 'onec-601', 'Связанная компания', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 601, sync_status = 'synced' WHERE id = ?", (client_id,))
            conn.execute(
                "INSERT INTO crm_primary_row_preferences(owner_user_id, crm_client_id, color_key, position, updated_at) VALUES (?, ?, ?, ?, ?)",
                (self.owner_id, client_id, "blue", 5000, "2026-09-04T00:00:00"),
            )

        response = self.client.put(
            f"/api/crm/clients/{client_id}/primary-row-preference",
            json={"colorKey": "pink", "expectedOrderVersion": 1},
        )
        listed = self.client.get("/api/crm/clients?primaryOnly=true")

        self.assertEqual(200, response.status_code)
        self.assertEqual("pink", response.json()["preference"]["colorKey"])
        self.assertEqual(5000, response.json()["preference"]["position"])
        self.assertEqual("pink", listed.json()["items"][0]["primaryRowPreference"]["colorKey"])

    def test_primary_color_update_accepts_latest_list_order_version_for_another_row(self) -> None:
        first = self.client.post("/api/crm/clients", json={"documentName": "Первая связанная"}).json()["client"]["id"]
        second = self.client.post("/api/crm/clients", json={"documentName": "Вторая связанная"}).json()["client"]["id"]
        self.link_primary_client(first, 605)
        self.link_primary_client(second, 606)

        first_update = self.client.put(
            f"/api/crm/clients/{first}/primary-row-preference",
            json={"colorKey": "blue", "expectedOrderVersion": 0},
        )
        current = self.client.get("/api/crm/clients?primaryOnly=true")
        second_update = self.client.put(
            f"/api/crm/clients/{second}/primary-row-preference",
            json={"colorKey": "pink", "expectedOrderVersion": current.json()["orderVersion"]},
        )
        updated = self.client.get("/api/crm/clients?primaryOnly=true")

        self.assertEqual(200, first_update.status_code)
        self.assertEqual(1, current.json()["orderVersion"])
        self.assertEqual(200, second_update.status_code)
        self.assertEqual(2, updated.json()["orderVersion"])
        self.assertEqual(["blue", "pink"], [item["primaryRowPreference"]["colorKey"] for item in updated.json()["items"]])

    def test_primary_reorder_uses_neighbors_and_rejects_stale_version_without_erasing_color(self) -> None:
        first = self.client.post("/api/crm/clients", json={"documentName": "Первая"}).json()["client"]["id"]
        second = self.client.post("/api/crm/clients", json={"documentName": "Вторая"}).json()["client"]["id"]
        third = self.client.post("/api/crm/clients", json={"documentName": "Третья"}).json()["client"]["id"]
        for counterparty_id, client_id in enumerate((first, second, third), start=611):
            self.link_primary_client(client_id, counterparty_id)

        initial = self.client.get("/api/crm/clients?primaryOnly=true")
        colored = self.client.put(
            f"/api/crm/clients/{third}/primary-row-preference",
            json={"colorKey": "orange", "expectedOrderVersion": 0},
        )
        reordered = self.client.post(
            "/api/crm/primary/reorder",
            json={
                "clientId": third,
                "beforeClientId": second,
                "afterClientId": first,
                "expectedOrderVersion": 1,
            },
        )
        stale = self.client.post(
            "/api/crm/primary/reorder",
            json={
                "clientId": second,
                "beforeClientId": None,
                "afterClientId": None,
                "expectedOrderVersion": 1,
            },
        )
        current = self.client.get("/api/crm/clients?primaryOnly=true")

        self.assertEqual(200, initial.status_code)
        self.assertEqual(200, colored.status_code)
        self.assertEqual(200, reordered.status_code)
        self.assertEqual([first, third, second], reordered.json()["clientIds"])
        self.assertEqual(2, reordered.json()["orderVersion"])
        self.assertEqual(409, stale.status_code)
        self.assertIn("Конфликт версии", stale.json()["detail"])
        self.assertEqual([first, third, second], [item["id"] for item in current.json()["items"]])
        self.assertEqual("orange", current.json()["items"][1]["primaryRowPreference"]["colorKey"])
        self.assertEqual(2, current.json()["orderVersion"])

    def test_personal_color_uses_tab_order_version_without_rewriting_positions(self) -> None:
        first = self.client.post("/api/crm/clients", json={"documentName": "Первый"}).json()
        second = self.client.post("/api/crm/clients", json={"documentName": "Второй"}).json()
        tab_id = first["assignment"]["tabId"]
        first_id, second_id = first["client"]["id"], second["client"]["id"]
        before = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"]
        version = max(item["rowPreference"]["orderVersion"] for item in before)
        positions = {item["id"]: item["rowPreference"]["position"] for item in before}
        reordered = self.client.post(f"/api/crm/tabs/{tab_id}/reorder", json={"clientId": second_id, "beforeClientId": first_id, "afterClientId": None, "expectedOrderVersion": version})
        self.assertEqual(200, reordered.status_code)
        stale = self.client.put(f"/api/crm/clients/{first_id}/row-preference", json={"tabId": tab_id, "colorKey": "red", "expectedOrderVersion": version})
        self.assertEqual(409, stale.status_code)
        current = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"]
        self.assertEqual([second_id, first_id], [item["id"] for item in current])
        reordered_positions = {item["id"]: item["rowPreference"]["position"] for item in current}
        self.assertNotEqual(reordered_positions, positions)
        success_version = max(item["rowPreference"]["orderVersion"] for item in current)
        success = self.client.put(f"/api/crm/clients/{first_id}/row-preference", json={"tabId": tab_id, "colorKey": "red", "expectedOrderVersion": success_version})
        self.assertEqual(200, success.status_code)
        after = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"]
        self.assertEqual("red", next(item for item in after if item["id"] == first_id)["rowPreference"]["colorKey"])
        self.assertEqual({item["id"]: item["rowPreference"]["position"] for item in after}, {item["id"]: item["rowPreference"]["position"] for item in current})

    def test_personal_color_materializes_missing_preference_with_current_version(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Без предпочтения"}).json()
        client_id, tab_id = created["client"]["id"], created["assignment"]["tabId"]
        initial = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"][0]["rowPreference"]
        with self.service.db.transaction() as conn:
            conn.execute("DELETE FROM crm_row_preferences WHERE owner_user_id = ? AND tab_id = ? AND crm_client_id = ?", (self.owner_id, tab_id, client_id))
        response = self.client.put(f"/api/crm/clients/{client_id}/row-preference", json={"tabId": tab_id, "colorKey": "blue", "expectedOrderVersion": initial["orderVersion"]})
        self.assertEqual(200, response.status_code, response.text)
        preference = response.json()["preference"]
        self.assertEqual("blue", preference["colorKey"])
        self.assertEqual(0, preference["position"])
        self.assertGreater(preference["orderVersion"], initial["orderVersion"])

    def test_personal_color_materialization_preserves_missing_row_list_position(self) -> None:
        first = self.client.post("/api/crm/clients", json={"documentName": "Альфа"}).json()
        second = self.client.post("/api/crm/clients", json={"documentName": "Бета"}).json()
        third = self.client.post("/api/crm/clients", json={"documentName": "Гамма"}).json()
        tab_id = first["assignment"]["tabId"]
        first_id = first["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute(
                "DELETE FROM crm_row_preferences WHERE owner_user_id = ? AND tab_id = ? AND crm_client_id = ?",
                (self.owner_id, tab_id, first_id),
            )
        before = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"]
        before_ids = [item["id"] for item in before]
        before_positions = {
            item["id"]: item["rowPreference"]["position"]
            for item in before
            if item["id"] != first_id
        }
        version = max(item["rowPreference"]["orderVersion"] for item in before if item["rowPreference"] is not None)

        response = self.client.put(
            f"/api/crm/clients/{first_id}/row-preference",
            json={"tabId": tab_id, "colorKey": "blue", "expectedOrderVersion": version},
        )

        self.assertEqual(200, response.status_code, response.text)
        after = self.client.get(f"/api/crm/clients?tabId={tab_id}").json()["items"]
        self.assertEqual(before_ids, [item["id"] for item in after])
        self.assertEqual(
            before_positions,
            {item["id"]: item["rowPreference"]["position"] for item in after if item["id"] != first_id},
        )

    def test_primary_preferences_require_linked_card_and_owner_scope(self) -> None:
        local = self.client.post("/api/crm/clients", json={"documentName": "Локальная карточка"}).json()["client"]["id"]
        linked = self.client.post("/api/crm/clients", json={"documentName": "Общая компания"}).json()["client"]["id"]
        self.link_primary_client(linked, 621)

        local_denied = self.client.put(
            f"/api/crm/clients/{local}/primary-row-preference",
            json={"colorKey": "pink", "expectedOrderVersion": 0},
        )
        owner_color = self.client.put(
            f"/api/crm/clients/{linked}/primary-row-preference",
            json={"colorKey": "pink", "expectedOrderVersion": 0},
        )
        self.as_user(self.other_id)
        owner_denied = self.client.put(
            f"/api/crm/clients/{linked}/primary-row-preference?ownerId={self.owner_id}",
            json={"colorKey": "pink", "expectedOrderVersion": 0},
        )
        other_color = self.client.put(
            f"/api/crm/clients/{linked}/primary-row-preference",
            json={"colorKey": "cyan", "expectedOrderVersion": 0},
        )
        self.as_user(self.owner_id)
        owner_list = self.client.get("/api/crm/clients?primaryOnly=true")
        self.as_user(self.admin_id, "admin")
        admin_without_owner = self.client.put(
            f"/api/crm/clients/{linked}/primary-row-preference",
            json={"colorKey": "pink", "expectedOrderVersion": 0},
        )
        admin_for_owner = self.client.put(
            f"/api/crm/clients/{linked}/primary-row-preference?ownerId={self.owner_id}",
            json={"colorKey": "pink", "expectedOrderVersion": 1},
        )

        self.assertEqual(400, local_denied.status_code)
        self.assertEqual(200, owner_color.status_code)
        self.assertEqual(403, owner_denied.status_code)
        self.assertEqual(200, other_color.status_code)
        self.assertEqual("pink", owner_list.json()["items"][0]["primaryRowPreference"]["colorKey"])
        self.assertEqual(400, admin_without_owner.status_code)
        self.assertEqual(403, admin_for_owner.status_code)

    def test_owner_can_reorder_between_neighbors_with_version(self) -> None:
        first = self.client.post("/api/crm/clients", json={"documentName": "Первый"}).json()
        second = self.client.post("/api/crm/clients", json={"documentName": "Второй"}).json()
        third = self.client.post("/api/crm/clients", json={"documentName": "Третий"}).json()
        tab_id = first["assignment"]["tabId"]

        response = self.client.post(
            f"/api/crm/tabs/{tab_id}/reorder",
            json={
                "clientId": third["client"]["id"],
                "beforeClientId": second["client"]["id"],
                "afterClientId": first["client"]["id"],
                "expectedOrderVersion": 0,
            },
        )

        self.assertEqual(200, response.status_code)
        self.assertEqual(1, response.json()["orderVersion"])
        self.assertEqual([first["client"]["id"], third["client"]["id"], second["client"]["id"]], response.json()["clientIds"])

        stale = self.client.post(
            f"/api/crm/tabs/{tab_id}/reorder",
            json={
                "clientId": second["client"]["id"],
                "beforeClientId": None,
                "afterClientId": None,
                "expectedOrderVersion": 0,
            },
        )

        self.assertEqual(409, stale.status_code)
        self.assertIn("Конфликт версии", stale.json()["detail"])

    def test_admin_can_leave_linked_client_only_in_primary_list(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Связанная компания"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (501, 'onec-501', 'Связанная компания', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 501, sync_status = 'synced' WHERE id = ?", (client_id,))

        self.as_user(self.admin_id, "admin")
        response = self.client.delete(f"/api/crm/clients/{client_id}/assignment?ownerId={self.owner_id}")

        self.assertEqual(200, response.status_code)
        self.as_user(self.owner_id)
        detail = self.client.get(f"/api/crm/clients/{client_id}")
        self.assertIsNone(detail.json()["client"]["assignment"])

    def test_only_admin_can_archive_and_restore_local_card_in_selected_workspace(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид на архив"}).json()
        client_id = created["client"]["id"]

        denied = self.client.post(f"/api/crm/clients/{client_id}/local-archive", json={"reason": "Дубликат"})
        self.assertEqual(403, denied.status_code)

        self.as_user(self.admin_id, "admin")
        missing_context = self.client.post(f"/api/crm/clients/{client_id}/local-archive", json={"reason": "Дубликат"})
        stale_archive = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Дубликат", "expectedVersion": 0},
        )

        self.assertEqual(400, missing_context.status_code)
        self.assertEqual(409, stale_archive.status_code)
        self.assertEqual("local", self.service.db.get_crm_client(client_id)["sync_status"])
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Дубликат", "expectedVersion": 1},
        )
        self.assertEqual(200, archived.status_code)
        self.assertEqual(2, archived.json()["version"])
        self.assertEqual("archived", self.service.db.get_crm_client(client_id)["sync_status"])

        stale_restore = self.client.post(
            f"/api/crm/clients/{client_id}/local-restore?ownerId={self.owner_id}",
            json={"expectedVersion": 1},
        )

        self.assertEqual(409, stale_restore.status_code)
        self.assertEqual("archived", self.service.db.get_crm_client(client_id)["sync_status"])
        restored = self.client.post(
            f"/api/crm/clients/{client_id}/local-restore?ownerId={self.owner_id}",
            json={"expectedVersion": 2},
        )
        self.assertEqual(200, restored.status_code)
        self.assertEqual(3, restored.json()["version"])
        self.assertEqual("local", self.service.db.get_crm_client(client_id)["sync_status"])

    def test_local_restore_does_not_reopen_an_assignment_only_archive(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид с архивным назначением"}).json()
        client_id = created["client"]["id"]

        self.as_user(self.admin_id, "admin")
        assignment_archive = self.client.post(
            f"/api/crm/clients/{client_id}/archive?ownerId={self.owner_id}",
            json={"reason": "Скрыть из личного списка"},
        )
        local_restore = self.client.post(
            f"/api/crm/clients/{client_id}/local-restore?ownerId={self.owner_id}",
            json={"expectedVersion": 1},
        )

        self.assertEqual(200, assignment_archive.status_code)
        self.assertEqual(400, local_restore.status_code)
        self.assertEqual("local", self.service.db.get_crm_client(client_id)["sync_status"])
        self.as_user(self.owner_id)
        detail = self.client.get(f"/api/crm/clients/{client_id}")
        self.assertIsNone(detail.json()["client"]["assignment"])

    def test_assignment_restore_does_not_reopen_a_local_card_archive(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Локально архивированный лид"}).json()
        client_id = created["client"]["id"]

        self.as_user(self.admin_id, "admin")
        local_archive = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Дубликат", "expectedVersion": 1},
        )
        assignment_restore = self.client.post(
            f"/api/crm/clients/{client_id}/restore?ownerId={self.owner_id}",
        )

        self.assertEqual(200, local_archive.status_code)
        self.assertEqual(400, assignment_restore.status_code)
        self.assertEqual("archived", self.service.db.get_crm_client(client_id)["sync_status"])
        self.as_user(self.owner_id)
        detail = self.client.get(f"/api/crm/clients/{client_id}")
        self.assertIsNone(detail.json()["client"]["assignment"])

    def test_card_update_rejects_stale_version_without_overwriting_data(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Версионный лид"}).json()
        client_id = created["client"]["id"]

        initial = self.client.get(f"/api/crm/clients/{client_id}")
        updated = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"documentName": "Актуальное название", "expectedVersion": 1},
        )
        stale = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"documentName": "Устаревшее название", "expectedVersion": 1},
        )

        self.assertEqual(200, initial.status_code)
        self.assertEqual(1, initial.json()["client"]["version"])
        self.assertEqual(200, updated.status_code)
        self.assertEqual("Актуальное название", updated.json()["client"]["documentName"])
        self.assertEqual(2, updated.json()["client"]["version"])
        self.assertEqual(409, stale.status_code)
        self.assertIn("Конфликт версии", stale.json()["detail"])
        self.assertEqual(
            "Актуальное название",
            self.client.get(f"/api/crm/clients/{client_id}").json()["client"]["documentName"],
        )

    def test_card_update_coalesces_one_pending_sync_job_for_shared_fields(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Очередь лид"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (801, 'onec-801', 'Очередь лид', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 801, sync_status = 'synced' WHERE id = ?", (client_id,))
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (502, 'onec-502', 'Очередь лид', '2026-09-04T00:00:00')"
            )
            conn.execute(
                "UPDATE crm_clients SET linked_counterparty_id = 502, sync_status = 'synced' WHERE id = ?",
                (client_id,),
            )

        first = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"documentName": "Очередь лид 1", "expectedVersion": 1},
        )
        second = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"email": "latest@example.test", "expectedVersion": 2},
        )
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT crm_client_id, author_user_id, operation, payload, status FROM crm_sync_jobs WHERE crm_client_id = ?",
                (client_id,),
            ).fetchall()
            card = conn.execute("SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual(200, first.status_code)
        self.assertEqual(200, second.status_code)
        self.assertEqual(1, len(jobs))
        self.assertEqual(client_id, jobs[0]["crm_client_id"])
        self.assertEqual(self.owner_id, jobs[0]["author_user_id"])
        self.assertEqual("update", jobs[0]["operation"])
        self.assertEqual("pending", jobs[0]["status"])
        self.assertEqual("Очередь лид 1", json.loads(jobs[0]["payload"])["document_name"])
        self.assertEqual("latest@example.test", json.loads(jobs[0]["payload"])["email"])
        self.assertEqual("pending", card["sync_status"])

    def test_due_sync_worker_blocks_automatic_update_without_proven_conditional_write(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Безопасная очередь"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (802, 'onec-802', 'Безопасная очередь', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 802, sync_status = 'synced' WHERE id = ?", (client_id,))
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (503, 'onec-503', 'Безопасная очередь', '2026-09-04T00:00:00')"
            )
            conn.execute(
                "UPDATE crm_clients SET linked_counterparty_id = 503, sync_status = 'synced' WHERE id = ?",
                (client_id,),
            )
        self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"documentName": "Безопасная очередь 2", "expectedVersion": 1},
        )

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()
            card = conn.execute("SELECT sync_status, sync_error FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual("blocked_capability", job["status"])
        self.assertEqual("blocked_capability", card["sync_status"])
        self.assertIn("условной записи", card["sync_error"])

    def test_due_sync_worker_recovers_a_stale_running_job_after_restart(self) -> None:
        """A worker crash must not leave an outbox job permanently claimed."""
        created = self.client.post("/api/crm/clients", json={"documentName": "Перезапуск очереди"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (804, 'onec-804', 'Перезапуск очереди', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 804, sync_status = 'synced' WHERE id = ?", (client_id,))
        self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"email": "restart@example.test", "expectedVersion": 1},
        )
        with self.service.db.transaction() as conn:
            conn.execute(
                "UPDATE crm_sync_jobs SET status = 'running', claimed_at = ? WHERE crm_client_id = ?",
                ("1970-01-01T00:00:00+00:00", client_id),
            )

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual("blocked_capability", job["status"])

    def test_new_edit_stays_pending_when_an_older_claimed_job_finishes(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Гонка очереди"}).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute("INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (803, 'onec-803', 'Гонка очереди', '2026-09-04T00:00:00')")
            conn.execute("UPDATE crm_clients SET linked_counterparty_id = 803, sync_status = 'synced' WHERE id = ?", (client_id,))
        self.client.patch(f"/api/crm/clients/{client_id}", json={"email": "first@example.test", "expectedVersion": 1})
        claimed = self.service.db.claim_next_crm_sync_job()
        self.client.patch(f"/api/crm/clients/{client_id}", json={"phone": "70000000000", "expectedVersion": 2})
        self.service.db.block_crm_sync_job(int(claimed["id"]), message="Старая попытка")
        with self.service.db.connect() as conn:
            jobs = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ? ORDER BY id", (client_id,)).fetchall()
            card = conn.execute("SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual(["blocked_capability", "pending"], [job["status"] for job in jobs])
        self.assertEqual("pending", card["sync_status"])

    def test_stale_create_completion_keeps_newer_local_edit_pending(self) -> None:
        """A create response must link the card without clearing a newer local edit."""
        class EditingDuringCreateOneC:
            def __init__(self, client: TestClient, test_case: unittest.TestCase, client_id: int) -> None:
                self.client = client
                self.test_case = test_case
                self.client_id = client_id

            def find_counterparty_by_identity(self, **_: object) -> None:
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                edit = self.client.patch(
                    f"/api/crm/clients/{self.client_id}",
                    json={"email": "newer@example.test", "expectedVersion": 1},
                )
                self.test_case.assertEqual(200, edit.status_code, edit.text)
                return {
                    "Ref_Key": "race-created-1",
                    "Description": "Созданная гонка",
                    "НаименованиеПолное": "Созданная гонка",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Созданная гонка", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        claimed = self.service.db.claim_next_crm_sync_job()
        self.assertIsNotNone(claimed)
        fake_onec = EditingDuringCreateOneC(self.client, self, client_id)
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        outcome = self.service._process_crm_create_job(claimed or {})
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT status FROM crm_sync_jobs WHERE crm_client_id = ? ORDER BY id", (client_id,)
            ).fetchall()
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual("completed", outcome)
        self.assertEqual("pending", card["sync_status"])
        self.assertIsNotNone(card["linked_counterparty_id"])
        self.assertEqual(["completed", "pending"], [job["status"] for job in jobs])

    def test_preclaim_edit_keeps_explicit_create_job_before_worker_runs(self) -> None:
        """Editing a queued local lead must not replace its explicit create intent."""
        class SuccessfulOneC:
            def find_counterparty_by_identity(self, **_: object) -> None:
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                return {
                    "Ref_Key": "preclaim-create-1",
                    "Description": "До получения заявки",
                    "НаименованиеПолное": "До получения заявки",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "До получения заявки", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        edit = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"email": "queued-edit@example.test", "expectedVersion": 1},
        )
        self.assertEqual(200, edit.status_code, edit.text)
        self.service.build_user_client = lambda **_: SuccessfulOneC()  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs(limit=1)
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT operation, status FROM crm_sync_jobs WHERE crm_client_id = ? ORDER BY id", (client_id,)
            ).fetchall()
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual({"processed": 1, "completed": 1}, result)
        self.assertIsNotNone(card["linked_counterparty_id"])
        self.assertEqual("pending", card["sync_status"])
        self.assertEqual([("create", "completed"), ("update", "pending")], [(job["operation"], job["status"]) for job in jobs])

    def test_stale_create_completion_keeps_newer_running_job_pending(self) -> None:
        """A newer claimed update is still outstanding when an older create returns."""
        class ClaimingEditDuringCreateOneC:
            def __init__(self, client: TestClient, test_case: unittest.TestCase, client_id: int, service: WebStockSyncService) -> None:
                self.client = client
                self.test_case = test_case
                self.client_id = client_id
                self.service = service

            def find_counterparty_by_identity(self, **_: object) -> None:
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                edit = self.client.patch(
                    f"/api/crm/clients/{self.client_id}",
                    json={"email": "running-edit@example.test", "expectedVersion": 1},
                )
                self.test_case.assertEqual(200, edit.status_code, edit.text)
                self.test_case.assertIsNotNone(self.service.db.claim_next_crm_sync_job())
                return {
                    "Ref_Key": "running-race-1",
                    "Description": "Гонка выполняющейся заявки",
                    "НаименованиеПолное": "Гонка выполняющейся заявки",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Гонка выполняющейся заявки", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        claimed_create = self.service.db.claim_next_crm_sync_job()
        self.assertIsNotNone(claimed_create)
        self.service.build_user_client = lambda **_: ClaimingEditDuringCreateOneC(self.client, self, client_id, self.service)  # type: ignore[method-assign]

        outcome = self.service._process_crm_create_job(claimed_create or {})
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT operation, status FROM crm_sync_jobs WHERE crm_client_id = ? ORDER BY id", (client_id,)
            ).fetchall()
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual("completed", outcome)
        self.assertIsNotNone(card["linked_counterparty_id"])
        self.assertEqual("pending", card["sync_status"])
        self.assertEqual([("create", "completed"), ("update", "running")], [(job["operation"], job["status"]) for job in jobs])

    def test_export_returns_only_the_current_owner_crm_workbook(self) -> None:
        self.client.post("/api/crm/clients", json={"documentName": "Экспорт владельца", "inn": "001234567890"})
        self.as_user(self.other_id)
        self.client.post("/api/crm/clients", json={"documentName": "Чужой экспорт"})
        self.as_user(self.owner_id)

        response = self.client.get("/api/crm/export?scope=all")
        workbook = load_workbook(BytesIO(response.content))

        self.assertEqual(200, response.status_code)
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertEqual(["Клиенты", "Контакты"], workbook.sheetnames)
        self.assertEqual("Экспорт владельца", workbook["Клиенты"]["A2"].value)
        self.assertEqual(2, workbook["Клиенты"].max_row)

    def test_local_crm_card_can_only_be_sent_to_onec_explicitly_after_validation(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Лид без ИНН"}).json()

        response = self.client.post(f"/api/crm/clients/{created['client']['id']}/send-to-onec")

        self.assertEqual(400, response.status_code)
        self.assertIn("ИНН", response.json()["detail"])

    def test_local_legal_entity_requires_a_valid_tax_identity_before_queueing_onec_create(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид с ошибочным ИНН", "inn": "7707", "kpp": "770701001"},
        ).json()

        response = self.client.post(f"/api/crm/clients/{created['client']['id']}/send-to-onec")

        self.assertEqual(400, response.status_code)
        self.assertIn("10 цифр", response.json()["detail"])

    def test_new_lead_stores_initial_contact_and_comment_in_the_owner_crm(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={
                "documentName": "Личный первый контакт",
                "contactPerson": "Анна",
                "email": "anna@example.test",
                "phone": "+7 900 000-00-00",
                "notes": "Перезвонить после выставки",
            },
        )
        client_id = created.json()["client"]["id"]

        contacts = self.client.get(f"/api/crm/clients/{client_id}/contacts")
        events = self.client.get(f"/api/crm/clients/{client_id}/events")
        with self.service.db.connect() as conn:
            card = conn.execute("SELECT contact_person, email, phone, notes FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual(201, created.status_code)
        self.assertEqual("", created.json()["client"]["email"])
        self.assertEqual("", created.json()["client"]["phone"])
        self.assertEqual([{"name": "Анна", "email": "anna@example.test", "phone": "+7 900 000-00-00", "isPrimary": True}], [
            {key: contact[key] for key in ("name", "email", "phone", "isPrimary")}
            for contact in contacts.json()["items"]
        ])
        self.assertEqual(["Перезвонить после выставки"], [event["body"] for event in events.json()["items"]])
        self.assertEqual((None, None, None, None), tuple(card))

    def test_create_lead_rolls_back_all_rows_when_initial_comment_insert_fails(self) -> None:
        with self.service.db.transaction() as conn:
            conn.execute("""
                CREATE TRIGGER fail_initial_crm_comment
                BEFORE INSERT ON crm_events WHEN NEW.kind = 'comment'
                BEGIN SELECT RAISE(ABORT, 'injected initial comment failure'); END
            """)

        response = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Атомарный лид", "contactPerson": "Ирина", "email": "irina@example.test", "phone": "+79990000000", "notes": "Первый комментарий"},
        )
        with self.service.db.connect() as conn:
            counts = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in ("crm_clients", "crm_assignments", "crm_contacts", "crm_events")}

        self.assertEqual(400, response.status_code)
        self.assertEqual({"crm_clients": 0, "crm_assignments": 0, "crm_contacts": 0, "crm_events": 0}, counts)

    def test_explicit_onec_create_persists_a_job_without_calling_onec(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Очередной лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]

        response = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT operation, author_user_id, status FROM crm_sync_jobs WHERE crm_client_id = ?",
                (client_id,),
            ).fetchone()

        self.assertEqual(202, response.status_code, response.text)
        self.assertEqual("queued", response.json()["sync"]["status"])
        self.assertEqual("pending", response.json()["client"]["syncStatus"])
        self.assertIsNotNone(job)
        self.assertEqual("create", job["operation"])
        self.assertEqual(self.owner_id, job["author_user_id"])
        self.assertEqual("pending", job["status"])

    def test_explicit_onec_create_persists_a_stable_idempotency_key(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Идемпотентный лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]

        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        with self.service.db.connect() as conn:
            columns = {row["name"] for row in conn.execute("PRAGMA table_info(crm_sync_jobs)").fetchall()}
            jobs = conn.execute("SELECT id, idempotency_key FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchall()

        self.assertIn("idempotency_key", columns)
        self.assertEqual(1, len(jobs))
        self.assertEqual(f"crm-create-{client_id}", jobs[0]["idempotency_key"])

    def test_archived_local_lead_never_runs_an_already_queued_onec_create(self) -> None:
        """An administrative archive must cancel the pending remote-create intent."""
        class RecordingOneC:
            def __init__(self) -> None:
                self.lookup_calls = 0
                self.create_calls = 0

            def find_counterparty_by_identity(self, **_: object) -> None:
                self.lookup_calls += 1
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                self.create_calls += 1
                return {
                    "Ref_Key": "must-not-be-created",
                    "Description": "Архивный лид",
                    "НаименованиеПолное": "Архивный лид",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Архивный лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        queued = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Отменено", "expectedVersion": queued.json()["client"]["version"]},
        )
        fake_onec = RecordingOneC()
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchone()
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(202, queued.status_code, queued.text)
        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual({}, result)
        self.assertEqual(0, fake_onec.lookup_calls)
        self.assertEqual(0, fake_onec.create_calls)
        self.assertEqual("completed", job["status"])
        self.assertIsNone(card["linked_counterparty_id"])
        self.assertEqual("archived", card["sync_status"])

    def test_archived_local_lead_stays_unlinked_when_a_claimed_create_finishes(self) -> None:
        """A create already sent to 1C must not revive a subsequently archived lead."""
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Гонка архива", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        queued = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")
        claimed = self.service.db.claim_next_crm_sync_job()
        self.assertIsNotNone(claimed)

        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, updated_at) VALUES (?, ?, ?, ?)",
                (870, "onec-870", "Гонка архива", "2026-09-04T00:00:00"),
            )

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Архивировать после POST", "expectedVersion": queued.json()["client"]["version"]},
        )
        self.service.db.complete_crm_create_job(int(claimed["id"]), counterparty_id=870)

        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE id = ?", (claimed["id"],)).fetchone()
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status, is_inactive FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(202, queued.status_code, queued.text)
        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual("completed", job["status"])
        self.assertIsNone(card["linked_counterparty_id"])
        self.assertEqual("archived", card["sync_status"])
        self.assertEqual(1, card["is_inactive"])

    def test_restoring_archived_local_lead_does_not_revive_cancelled_onec_create(self) -> None:
        """Restoring a card must require a new explicit request before 1C creation."""
        class RecordingOneC:
            def __init__(self) -> None:
                self.lookup_calls = 0
                self.create_calls = 0

            def find_counterparty_by_identity(self, **_: object) -> None:
                self.lookup_calls += 1
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                self.create_calls += 1
                return {}

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Восстановленный лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        queued = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Отменено", "expectedVersion": queued.json()["client"]["version"]},
        )
        restored = self.client.post(
            f"/api/crm/clients/{client_id}/local-restore?ownerId={self.owner_id}",
            json={"expectedVersion": archived.json()["version"]},
        )
        fake_onec = RecordingOneC()
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(202, queued.status_code, queued.text)
        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual(200, restored.status_code, restored.text)
        self.assertEqual({}, result)
        self.assertEqual(0, fake_onec.lookup_calls)
        self.assertEqual(0, fake_onec.create_calls)
        self.assertEqual("completed", job["status"])

    def test_archived_local_lead_cannot_queue_another_onec_create(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Закрытый лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Отменено", "expectedVersion": created["client"]["version"]},
        )
        self.as_user(self.owner_id)
        requeue = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchall()
            card = conn.execute(
                "SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual(400, requeue.status_code)
        self.assertIn("архив", requeue.json()["detail"].casefold())
        self.assertEqual([], jobs)
        self.assertEqual("archived", card["sync_status"])

    def test_archived_local_lead_cannot_retry_a_blocked_onec_create(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Архивный повтор", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.service.save_system_settings({"base_url": "https://onec.example.test"})
        queued = self.client.post(f"/api/crm/clients/{client_id}/send-to-onec")
        blocked = self.service.run_due_crm_sync_jobs()

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Отменено", "expectedVersion": queued.json()["client"]["version"]},
        )
        self.as_user(self.owner_id)
        retry = self.client.post(f"/api/crm/clients/{client_id}/retry-onec")
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchone()
            card = conn.execute(
                "SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(202, queued.status_code, queued.text)
        self.assertEqual({"processed": 1, "blocked": 1}, blocked)
        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual(400, retry.status_code)
        self.assertIn("архив", retry.json()["detail"].casefold())
        self.assertEqual("blocked_credentials", job["status"])
        self.assertEqual("archived", card["sync_status"])

    def test_archived_local_lead_cannot_link_an_existing_onec_counterparty(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Архивное связывание", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, inn, kpp, updated_at) "
                "VALUES (904, 'onec-904', 'Контрагент архива', '7707083893', '770701001', '2026-09-04T00:00:00')"
            )

        self.as_user(self.admin_id, "admin")
        archived = self.client.post(
            f"/api/crm/clients/{client_id}/local-archive?ownerId={self.owner_id}",
            json={"reason": "Отменено", "expectedVersion": created["client"]["version"]},
        )
        self.as_user(self.owner_id)
        link = self.client.post(
            f"/api/crm/clients/{client_id}/link-existing",
            json={"counterpartyId": 904, "expectedVersion": archived.json()["version"]},
        )
        with self.service.db.connect() as conn:
            card = conn.execute(
                "SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(200, archived.status_code, archived.text)
        self.assertEqual(400, link.status_code)
        self.assertIn("архив", link.json()["detail"].casefold())
        self.assertIsNone(card["linked_counterparty_id"])
        self.assertEqual("archived", card["sync_status"])

    def test_create_worker_blocks_missing_submitter_onec_credentials_without_fallback(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид без учётных данных", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.service.save_system_settings({"base_url": "https://onec.example.test"})
        self.service.update_user_profile(
            user_id=self.admin_id,
            full_name="Администратор",
            onec_username="admin-onec",
            onec_password="admin-password",
        )
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT author_user_id, status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchone()
            card = conn.execute(
                "SELECT sync_status, sync_error FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual(self.owner_id, job["author_user_id"])
        self.assertEqual("blocked_credentials", job["status"])
        self.assertEqual("blocked_credentials", card["sync_status"])
        self.assertIn("учётные данные", card["sync_error"].casefold())

    def test_submitter_can_explicitly_retry_a_blocked_credential_create_job(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Повторный лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.service.save_system_settings({"base_url": "https://onec.example.test"})
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.assertEqual({"processed": 1, "blocked": 1}, self.service.run_due_crm_sync_jobs())
        self.service.update_user_profile(
            user_id=self.owner_id,
            full_name="Владелец",
            onec_username="owner-onec",
            onec_password="owner-password",
        )

        response = self.client.post(f"/api/crm/clients/{client_id}/retry-onec")
        with self.service.db.connect() as conn:
            job = conn.execute(
                "SELECT author_user_id, status, claimed_at FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchone()
            card = conn.execute(
                "SELECT sync_status, sync_error FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(202, response.status_code, response.text)
        self.assertEqual("queued", response.json()["sync"]["status"])
        self.assertEqual(self.owner_id, job["author_user_id"])
        self.assertEqual("pending", job["status"])
        self.assertIsNone(job["claimed_at"])
        self.assertEqual("pending", card["sync_status"])
        self.assertIsNone(card["sync_error"])

    def test_non_submitter_cannot_retry_a_blocked_credential_create_job(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Чужая заявка", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.service.save_system_settings({"base_url": "https://onec.example.test"})
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.assertEqual({"processed": 1, "blocked": 1}, self.service.run_due_crm_sync_jobs())

        self.as_user(self.admin_id, "admin")
        response = self.client.post(f"/api/crm/clients/{client_id}/retry-onec?ownerId={self.owner_id}")

        self.assertEqual(403, response.status_code)
        self.assertIn("только автор", response.json()["detail"].casefold())

    def test_create_worker_blocks_revoked_submitter_onec_access(self) -> None:
        class RevokedAccessOneC:
            def find_counterparty_by_identity(self, **_: object) -> dict[str, object] | None:
                raise OneCClientError("1С вернула HTTP 401: доступ отозван")

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                raise AssertionError("Создание без действующего доступа 1С запрещено")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид с отозванным доступом", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.service.build_user_client = lambda **_: RevokedAccessOneC()  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()
            card = conn.execute("SELECT sync_status, sync_error FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual("blocked_credentials", job["status"])
        self.assertEqual("blocked_credentials", card["sync_status"])
        self.assertIn("доступ", card["sync_error"].casefold())

    def test_create_worker_recovers_an_unknown_post_without_a_second_create(self) -> None:
        class UnknownPostOneC:
            def __init__(self) -> None:
                self.created_cards: list[dict[str, object]] = []
                self.remote: dict[str, object] | None = None
                self.lookup_count = 0

            def find_counterparty_by_identity(self, **_: object) -> dict[str, object] | None:
                self.lookup_count += 1
                return self.remote if self.lookup_count >= 3 else None

            def create_counterparty(self, card: dict[str, object]) -> dict[str, object]:
                self.created_cards.append(dict(card))
                self.remote = {
                    "Ref_Key": "crm-queue-1",
                    "Description": "Очередной лид",
                    "НаименованиеПолное": "Очередной лид",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }
                raise OneCClientError("Соединение оборвалось после POST")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Очередной лид", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        fake_onec = UnknownPostOneC()
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        first = self.service.run_due_crm_sync_jobs()
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_sync_jobs SET available_at = ? WHERE crm_client_id = ?", ("1970-01-01T00:00:00+00:00", client_id))
        second = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status, attempt_count FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()
            card = conn.execute("SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "retried": 1}, first)
        self.assertEqual({"processed": 1, "completed": 1}, second)
        self.assertEqual(1, len(fake_onec.created_cards))
        self.assertEqual("completed", job["status"])
        self.assertEqual(2, job["attempt_count"])
        self.assertIsNotNone(card["linked_counterparty_id"])
        self.assertEqual("synced", card["sync_status"])

    def test_create_worker_never_repeats_an_unknown_post_while_identity_is_not_visible(self) -> None:
        """Eventual 1C visibility after POST must not turn a timeout into a duplicate."""
        class DelayedUnknownPostOneC:
            def __init__(self) -> None:
                self.created_cards: list[dict[str, object]] = []
                self.lookup_count = 0

            def find_counterparty_by_identity(self, **_: object) -> None:
                self.lookup_count += 1
                return None

            def create_counterparty(self, card: dict[str, object]) -> dict[str, object]:
                self.created_cards.append(dict(card))
                raise OneCClientError("Соединение оборвалось после POST")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Отложенная видимость", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        fake_onec = DelayedUnknownPostOneC()
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        first = self.service.run_due_crm_sync_jobs()
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_sync_jobs SET available_at = ? WHERE crm_client_id = ?", ("1970-01-01T00:00:00+00:00", client_id))
        second = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status, payload FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "retried": 1}, first)
        self.assertEqual({"processed": 1, "retried": 1}, second)
        self.assertEqual(1, len(fake_onec.created_cards))
        self.assertEqual("pending", job["status"])
        self.assertTrue(json.loads(job["payload"])["post_uncertain"])

    def test_create_worker_blocks_a_non_retriable_validation_error(self) -> None:
        """A rejected payload is not an uncertain POST and must not be retried."""
        class ValidationRejectingOneC:
            def find_counterparty_by_identity(self, **_: object) -> None:
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                raise OneCClientError("1С вернула HTTP 400: некорректные реквизиты")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Некорректные реквизиты", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.service.build_user_client = lambda **_: ValidationRejectingOneC()  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status, available_at FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()
            card = conn.execute("SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual("blocked_validation", job["status"])
        self.assertEqual("blocked_validation", card["sync_status"])

    def test_create_worker_blocks_a_metadata_error_without_retry(self) -> None:
        """A publication capability error cannot become safe by retrying the same payload."""
        class MetadataRejectingOneC:
            def find_counterparty_by_identity(self, **_: object) -> None:
                return None

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                raise OneCClientError("В OData metadata Контрагенты не найдено обязательное поле")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Неполная публикация", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        self.service.build_user_client = lambda **_: MetadataRejectingOneC()  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual("blocked_validation", job["status"])

    def test_create_worker_blocks_a_preexisting_identity_for_manual_linking(self) -> None:
        class ExistingOneC:
            def __init__(self) -> None:
                self.create_calls = 0

            def find_counterparty_by_identity(self, **_: object) -> dict[str, object]:
                return {
                    "Ref_Key": "existing-crm-queue-1",
                    "Description": "Уже существует",
                    "НаименованиеПолное": "Уже существует",
                    "ИНН": "7707083893",
                    "КПП": "770701001",
                }

            def create_counterparty(self, _: dict[str, object]) -> dict[str, object]:
                self.create_calls += 1
                raise AssertionError("Создание при найденном совпадении запрещено")

        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид с совпадением", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        self.assertEqual(202, self.client.post(f"/api/crm/clients/{client_id}/send-to-onec").status_code)
        fake_onec = ExistingOneC()
        self.service.build_user_client = lambda **_: fake_onec  # type: ignore[method-assign]

        result = self.service.run_due_crm_sync_jobs()
        with self.service.db.connect() as conn:
            job = conn.execute("SELECT status FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)).fetchone()
            card = conn.execute("SELECT linked_counterparty_id, sync_status FROM crm_clients WHERE id = ?", (client_id,)).fetchone()

        self.assertEqual({"processed": 1, "blocked": 1}, result)
        self.assertEqual(0, fake_onec.create_calls)
        self.assertEqual("blocked_duplicate", job["status"])
        self.assertIsNone(card["linked_counterparty_id"])
        self.assertEqual("blocked_duplicate", card["sync_status"])

    def test_owner_can_confirm_matching_existing_onec_counterparty(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид для связывания", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, inn, kpp, updated_at) "
                "VALUES (901, 'onec-901', 'Найденная компания', '7707083893', '770701001', '2026-09-04T00:00:00')"
            )

        response = self.client.post(
            f"/api/crm/clients/{client_id}/link-existing",
            json={"counterpartyId": 901, "expectedVersion": 1},
        )

        self.assertEqual(200, response.status_code)
        self.assertEqual(901, response.json()["client"]["linkedCounterpartyId"])
        self.assertEqual("synced", response.json()["client"]["syncStatus"])
        self.assertEqual(2, response.json()["client"]["version"])
        audit = self.client.get(f"/api/crm/clients/{client_id}/audit")
        self.assertEqual("link_existing_counterparty", audit.json()["items"][-1]["action"])

    def test_link_candidates_show_only_same_legal_identity(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Лид для проверки", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]
        with self.service.db.transaction() as conn:
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, inn, kpp, updated_at) "
                "VALUES (902, 'onec-902', 'Точное совпадение', '7707083893', '770701001', '2026-09-04T00:00:00')"
            )
            conn.execute(
                "INSERT INTO counterparties(id, onec_key, name, inn, kpp, updated_at) "
                "VALUES (903, 'onec-903', 'Другая КПП', '7707083893', '770799999', '2026-09-04T00:00:00')"
            )

        response = self.client.get(f"/api/crm/clients/{client_id}/link-candidates")

        self.assertEqual(200, response.status_code)
        self.assertEqual(
            [{"id": 902, "onecKey": "onec-902", "name": "Точное совпадение", "inn": "7707083893", "kpp": "770701001"}],
            response.json()["items"],
        )

    def test_link_confirmation_keeps_owner_context_server_guarded(self) -> None:
        created = self.client.post(
            "/api/crm/clients",
            json={"documentName": "Закрытый кандидат", "inn": "7707083893", "kpp": "770701001"},
        ).json()
        client_id = created["client"]["id"]

        denied = self.client.post(
            f"/api/crm/clients/{client_id}/link-existing?ownerId={self.other_id}",
            json={"counterpartyId": 901, "expectedVersion": 1},
        )
        self.as_user(self.admin_id, "admin")
        missing_context = self.client.post(
            f"/api/crm/clients/{client_id}/link-existing",
            json={"counterpartyId": 901, "expectedVersion": 1},
        )

        self.assertEqual(403, denied.status_code)
        self.assertEqual(400, missing_context.status_code)

    def test_local_lead_update_without_identity_does_not_enqueue_onec_sync(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Локальный без ИНН"}).json()
        client_id = created["client"]["id"]

        response = self.client.patch(
            f"/api/crm/clients/{client_id}",
            json={"documentName": "Локальный без ИНН: уточнён", "expectedVersion": 1},
        )
        with self.service.db.connect() as conn:
            jobs = conn.execute(
                "SELECT id FROM crm_sync_jobs WHERE crm_client_id = ?", (client_id,)
            ).fetchall()
            card = conn.execute(
                "SELECT sync_status FROM crm_clients WHERE id = ?", (client_id,)
            ).fetchone()

        self.assertEqual(200, response.status_code)
        self.assertEqual([], jobs)
        self.assertEqual("local", card["sync_status"])

    def test_active_lists_exclude_inactive_card_with_legacy_active_assignment(self) -> None:
        created = self.client.post("/api/crm/clients", json={"documentName": "Неактивный лид"}).json()
        client_id = created["client"]["id"]
        tab_id = created["assignment"]["tabId"]
        with self.service.db.transaction() as conn:
            conn.execute("UPDATE crm_clients SET is_inactive = 1, sync_status = 'archived' WHERE id = ?", (client_id,))

        response = self.client.get(f"/api/crm/clients?tabId={tab_id}")

        self.assertEqual(200, response.status_code)
        self.assertEqual([], response.json()["items"])


if __name__ == "__main__":
    unittest.main()
