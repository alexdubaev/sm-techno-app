from __future__ import annotations

import re
from math import ceil
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet


REQUIRED_FIELDS = ["article", "name", "qty", "price_vat"]
DEFAULT_OFFER_ROW_HEIGHT = 17.1
MAX_OFFER_ROW_HEIGHT = 90.0

FIELD_ALIASES: dict[str, list[str]] = {
    "row_no": ["№", "n", "номер", "номер п/п", "no"],
    "article": ["артикул", "код", "номер детали", "part number", "part no", "partno", "каталожный номер"],
    "name": ["наименование", "название", "описание", "товар", "номенклатура"],
    "brand": ["бренд", "производитель", "марка", "brand"],
    "qty": ["количество", "кол-во", "кол во", "qty", "quantity"],
    "price_vat": ["цена с ндс", "цена", "price", "цена, руб", "цена руб"],
    "amount_vat": ["сумма с ндс", "сумма", "итого", "amount"],
    "delivery_time": ["срок поставки", "срок", "поставка", "delivery", "сроки"],
    "note": ["примечание", "комментарий", "note", "notes"],
}


@dataclass(frozen=True)
class CommercialOfferLineInput:
    row_no: int
    article: str | None
    name: str | None
    brand: str | None
    qty: float | None
    price_vat: float | None
    amount_vat: float | None
    delivery_time: str | None
    note: str | None
    item_id: int | None = None
    warehouse_id: int | None = None
    warehouse_name: str = ""


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("ё", "е")
    text = re.sub(r"[\n\r\t]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" .,:;()[]")


def normalize_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[^0-9,\.\- ]", "", text)
    text = text.replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def _find_header_row(ws: Worksheet, max_scan_rows: int = 30) -> tuple[int, dict[str, int]]:
    aliases = {field: [normalize_header(alias) for alias in names] for field, names in FIELD_ALIASES.items()}

    best_row = 0
    best_mapping: dict[str, int] = {}
    best_score = 0
    for row in range(1, min(ws.max_row, max_scan_rows) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            header = normalize_header(ws.cell(row=row, column=col).value)
            if not header:
                continue
            for field, names in aliases.items():
                if header in names and field not in mapping:
                    mapping[field] = col
        if len(mapping) > best_score:
            best_row = row
            best_mapping = mapping
            best_score = len(mapping)

    missing = [field for field in REQUIRED_FIELDS if field not in best_mapping]
    if missing:
        labels = {
            "article": "Артикул",
            "name": "Наименование",
            "qty": "Количество",
            "price_vat": "Цена с НДС",
        }
        readable = ", ".join(labels.get(field, field) for field in missing)
        raise ValueError(f"Не удалось найти обязательные столбцы во входном файле: {readable}.")

    return best_row, best_mapping


def read_source_offer_lines(source_path: Path | str) -> list[CommercialOfferLineInput]:
    try:
        workbook = load_workbook(source_path, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("Не удалось открыть Excel-файл КП. Проверьте, что файл не повреждён.") from exc
    ws = workbook[workbook.sheetnames[0]]
    header_row, mapping = _find_header_row(ws)

    lines: list[CommercialOfferLineInput] = []
    for excel_row in range(header_row + 1, ws.max_row + 1):
        raw = {field: ws.cell(row=excel_row, column=column).value for field, column in mapping.items()}
        article = str(raw.get("article") or "").strip()
        name = str(raw.get("name") or "").strip()
        if not article and not name:
            continue

        qty = normalize_number(raw.get("qty"))
        price_vat = normalize_number(raw.get("price_vat"))
        amount_vat = normalize_number(raw.get("amount_vat"))
        if amount_vat is None and qty is not None and price_vat is not None:
            amount_vat = qty * price_vat

        lines.append(
            CommercialOfferLineInput(
                row_no=len(lines) + 1,
                article=article or None,
                name=name or None,
                brand=str(raw.get("brand") or "").strip() or None,
                qty=qty,
                price_vat=price_vat,
                amount_vat=amount_vat,
                delivery_time=str(raw.get("delivery_time") or "").strip() or None,
                note=str(raw.get("note") or "").strip() or None,
            )
        )

    if not lines:
        raise ValueError("Во входном файле не найдено ни одной позиции для КП.")
    return lines


def _save_row_style(ws: Worksheet, row: int, max_col: int) -> tuple[list[dict[str, Any]], float | None]:
    styles: list[dict[str, Any]] = []
    for column in range(1, max_col + 1):
        cell = ws.cell(row=row, column=column)
        styles.append(
            {
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
            }
        )
    return styles, ws.row_dimensions[row].height


def _apply_row_style(ws: Worksheet, row: int, styles: list[dict[str, Any]], row_height: float | None) -> None:
    for column, style in enumerate(styles, start=1):
        cell = ws.cell(row=row, column=column)
        cell.font = copy(style["font"])
        cell.fill = copy(style["fill"])
        cell.border = copy(style["border"])
        cell.alignment = copy(style["alignment"])
        cell.number_format = style["number_format"]
        cell.protection = copy(style["protection"])
    ws.row_dimensions[row].height = row_height


def _set_cell_alignment(cell: Any, **updates: Any) -> None:
    alignment = copy(cell.alignment)
    for key, value in updates.items():
        setattr(alignment, key, value)
    cell.alignment = alignment


def _cell(column_letter: str, row: int) -> str:
    return f"{column_letter}{row}"


def _find_total_row(ws: Worksheet, first_row: int) -> int:
    for row in range(first_row, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=column).value
            if isinstance(value, str) and "итого" in normalize_header(value):
                return row
    return first_row + 20


def _pop_merged_ranges(ws: Worksheet) -> list[CellRange]:
    ranges = [CellRange(str(merged_range)) for merged_range in ws.merged_cells.ranges]
    for merged_range in ranges:
        ws.unmerge_cells(str(merged_range))
    return ranges


def _restore_shifted_merged_ranges(
    ws: Worksheet,
    ranges: list[CellRange],
    *,
    first_row: int,
    template_total_row: int,
    row_delta: int,
) -> None:
    for merged_range in ranges:
        if merged_range.max_row < first_row:
            ws.merge_cells(str(merged_range))
            continue

        if merged_range.min_row >= template_total_row:
            shifted_range = CellRange(
                min_col=merged_range.min_col,
                min_row=merged_range.min_row + row_delta,
                max_col=merged_range.max_col,
                max_row=merged_range.max_row + row_delta,
            )
            ws.merge_cells(str(shifted_range))


def _quantity_cell_value(qty: float | None) -> float | int | None:
    if isinstance(qty, (int, float)) and float(qty).is_integer():
        return int(qty)
    return qty


def _quantity_number_format(qty: float | None) -> str:
    if isinstance(qty, (int, float)) and float(qty).is_integer():
        return "0"
    return "0.###"


def _estimate_wrapped_line_count(value: str | None, column_width: float | None) -> int:
    if not value:
        return 1

    chars_per_line = max(int((column_width or 23) * 1.15), 18)
    return max(1, sum(max(1, ceil(len(part) / chars_per_line)) for part in str(value).splitlines()))


def _fit_name_row_height(ws: Worksheet, row: int, name: str | None, base_height: float | None) -> None:
    line_count = _estimate_wrapped_line_count(name, ws.column_dimensions["C"].width)
    if line_count <= 1:
        return

    row_height = base_height or DEFAULT_OFFER_ROW_HEIGHT
    current_height = ws.row_dimensions[row].height or row_height
    ws.row_dimensions[row].height = min(max(current_height, row_height * line_count), MAX_OFFER_ROW_HEIGHT)


def _client_cell_value(client_name: str) -> CellRichText:
    return CellRichText("Покупатель: ", TextBlock(InlineFont(b=True), client_name))


def generate_commercial_offer_workbook(
    *,
    template_path: Path | str,
    output_path: Path | str,
    lines: list[CommercialOfferLineInput],
    offer_number: str,
    client_name: str,
    offer_date: date,
) -> None:
    if not lines:
        raise ValueError("Нельзя сформировать КП без позиций.")

    workbook = load_workbook(template_path)
    if "КП" not in workbook.sheetnames:
        raise ValueError("В шаблоне КП не найден лист 'КП'.")

    ws = workbook["КП"]
    first_row = 13
    max_col = 9
    columns = {
        "row_no": "A",
        "article": "B",
        "name": "C",
        "brand": "D",
        "qty": "E",
        "price_vat": "F",
        "amount_vat": "G",
        "delivery_time": "H",
        "note": "I",
    }

    template_total_row = _find_total_row(ws, first_row)
    default_rows = max(template_total_row - first_row, 1)
    row_delta = len(lines) - default_rows
    merged_ranges = _pop_merged_ranges(ws)
    row_style, row_height = _save_row_style(ws, first_row, max_col)
    ws.delete_rows(first_row, default_rows)
    ws.insert_rows(first_row, len(lines))
    _restore_shifted_merged_ranges(
        ws,
        merged_ranges,
        first_row=first_row,
        template_total_row=template_total_row,
        row_delta=row_delta,
    )

    for offset, line in enumerate(lines):
        row = first_row + offset
        _apply_row_style(ws, row, row_style, row_height)
        ws[_cell(columns["row_no"], row)] = offset + 1
        ws[_cell(columns["article"], row)] = line.article
        ws[_cell(columns["name"], row)] = line.name
        _set_cell_alignment(ws[_cell(columns["name"], row)], wrap_text=True, vertical="top")
        _fit_name_row_height(ws, row, line.name, row_height)
        ws[_cell(columns["brand"], row)] = line.brand
        ws[_cell(columns["qty"], row)] = _quantity_cell_value(line.qty)
        ws[_cell(columns["price_vat"], row)] = line.price_vat
        ws[_cell(columns["amount_vat"], row)] = (
            f'=IF(OR({_cell(columns["qty"], row)}="",{_cell(columns["price_vat"], row)}=""),"",'
            f'{_cell(columns["qty"], row)}*{_cell(columns["price_vat"], row)})'
        )
        ws[_cell(columns["delivery_time"], row)] = line.delivery_time
        ws[_cell(columns["note"], row)] = line.note

    total_row = first_row + len(lines)
    ws[_cell(columns["amount_vat"], total_row)] = (
        f'=SUM({_cell(columns["amount_vat"], first_row)}:{_cell(columns["amount_vat"], total_row - 1)})'
    )
    ws["A7"] = f"Коммерческое предложение № {offer_number}"
    ws["A9"] = _client_cell_value(client_name)
    ws["F9"] = f"Дата: {offer_date.strftime('%d.%m.%Y')}"
    _set_cell_alignment(ws["A9"], vertical="center")
    _set_cell_alignment(ws["F9"], horizontal="right", vertical="center")

    for row in range(first_row, total_row + 1):
        ws[_cell(columns["qty"], row)].number_format = _quantity_number_format(ws[_cell(columns["qty"], row)].value)
        ws[_cell(columns["price_vat"], row)].number_format = "#,##0.00 ₽"
        ws[_cell(columns["amount_vat"], row)].number_format = "#,##0.00 ₽"

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
