"""Read and plan CRM workbooks using only the supplied local CRM snapshot.

No service, database, synchronization or counterparty dependencies belong here.
The repository owns authorization and the single final write transaction.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


CLIENT_FIELDS = {
    "Компания": "document_name", "ИНН": "inn", "КПП": "kpp", "Город": "city",
    "Сайт": "website", "Основной контакт": "contact_person", "Телефон": "phone",
    "Почта": "email", "Telegram": "telegram", "MAX": "max_link", "Комментарий": "notes",
}
CONTACT_FIELDS = {"Контактное лицо": "name", "Телефон": "phone", "Почта": "email"}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = str(value).strip()
    # Undo the exporter's formula-protection prefix, preserving real apostrophes.
    if result.startswith("'") and result[1:].lstrip().startswith(("=", "+", "-", "@")):
        return result[1:]
    return result


def normalized(value: Any) -> str:
    return " ".join(text(value).casefold().split())


def phone(value: Any) -> str:
    return re.sub(r"\D", "", text(value))


def positive_id(value: Any) -> int | None:
    raw = text(value)
    return int(raw) if raw.isascii() and raw.isdigit() and 0 < int(raw) <= 9223372036854775807 else None


@dataclass
class WorkbookRows:
    clients: list[dict[str, Any]]
    contacts: list[dict[str, Any]]
    errors: list[dict[str, Any]]


class ImportValidationError(ValueError):
    def __init__(self, preview: dict[str, Any]):
        super().__init__("Исправьте ошибки строк перед импортом.")
        self.preview = preview


def read_crm_workbook(content: bytes) -> WorkbookRows:
    """Accept legacy and version-1 books; formulas are never executed or imported."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError("Книга Excel слишком велика.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError, TypeError, ValueError, SyntaxError) as exc:
        raise ValueError("Загрузите корректную книгу XLSX CRM.") from exc
    errors: list[dict[str, Any]] = []
    result: dict[str, list[dict[str, Any]]] = {}
    try:
        if "Клиенты" not in book.sheetnames:
            raise ValueError("В книге отсутствует лист «Клиенты».")
        for name in ("Клиенты", "Контакты"):
            result[name] = []
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            if sheet.max_row > 20001 or sheet.max_column > 100:
                raise ValueError("В книге слишком много строк или колонок.")
            iterator = sheet.iter_rows()
            header = [text(cell.value) for cell in next(iterator, ())]
            required = "Компания" if name == "Клиенты" else "Контактное лицо"
            if required not in header or len([h for h in header if h]) != len(set(h for h in header if h)):
                raise ValueError(f"Проверьте заголовки листа «{name}».")
            for number, cells in enumerate(iterator, start=2):
                if not any(text(cell.value) for cell in cells):
                    continue
                row = {key: text(cell.value) for key, cell in zip(header, cells) if key}
                row["_row"] = number
                for key, cell in zip(header, cells):
                    if cell.data_type == "f":
                        errors.append(dict(sheet=name, row=number, field=key, code="formula_not_allowed", message="Замените формулу обычным значением."))
                if name == "Клиенты" and row.get("__export_version") not in (None, "", "1"):
                    errors.append(dict(sheet=name, row=number, field="__export_version", code="unsupported_version", message="Версия экспорта не поддерживается."))
                result[name].append(row)
    except (BadZipFile, KeyError, TypeError, SyntaxError) as exc:
        raise ValueError("Загрузите корректную книгу XLSX CRM.") from exc
    finally:
        book.close()
    return WorkbookRows(result["Клиенты"], result["Контакты"], errors)


def _same_channel(left: dict[str, Any], right: dict[str, Any]) -> bool:
    return bool((phone(left.get("phone")) and phone(left.get("phone")) == phone(right.get("phone"))) or (normalized(left.get("email")) and normalized(left.get("email")) == normalized(right.get("email"))))


def plan_crm_import(book: WorkbookRows, *, owner_id: int, target_tab_id: int | None,
                    new_tab_name: str | None, include_existing_clients: bool,
                    clients: list[dict[str, Any]], contacts: list[dict[str, Any]],
                    assignments: list[dict[str, Any]], colors: frozenset[str],
                    archived_clients: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    preview = dict(ownerId=owner_id, target=dict(tabId=target_tab_id, newTabName=new_tab_name),
                   clientsToCreate=0, clientsToUpdate=0, unchangedClients=0, clientsToAssign=0,
                   contactsToCreate=0, contactsToUpdate=0, duplicateConflicts=0,
                   skippedOneCLinked=0, skippedArchived=0, errors=list(book.errors))
    client_actions: list[dict[str, Any]] = []
    contact_actions: list[dict[str, Any]] = []
    accessible = [c for c in clients if c.get("linked_counterparty_id") is not None or c.get("crm_owner_user_id") == owner_id]
    archived_accessible = [
        c for c in (archived_clients or [])
        if c.get("linked_counterparty_id") is not None or c.get("crm_owner_user_id") == owner_id
    ]
    staged = list(accessible)
    resolved: dict[int, dict[str, Any]] = {}
    source_ids: dict[int, int] = {}
    aliases: dict[str, set[int]] = {}
    seen: set[int] = set()

    def error(sheet, row, code, message, field=None):
        item = dict(sheet=sheet, row=row["_row"], code=code, message=message)
        if field:
            item["field"] = field
        preview["errors"].append(item)
        if code.startswith("ambiguous_") or code.startswith("duplicate_"):
            preview["duplicateConflicts"] += 1

    def matches_for(values: dict[str, Any], client_id: int | None,
                    candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
        matches = [client for client in candidates if client["id"] == client_id] if client_id else []
        if not matches and values.get("inn"):
            matches = [
                client for client in candidates
                if text(client.get("inn")) == values["inn"]
                and text(client.get("kpp")) == text(values.get("kpp"))
            ]
        if not matches and values.get("document_name") and (phone(values.get("phone")) or normalized(values.get("email"))):
            matches = [
                client for client in candidates
                if normalized(client.get("document_name")) == normalized(values["document_name"])
                and _same_channel(values, client)
            ]
        return matches

    for row in book.clients:
        values = {field: row[header] for header, field in CLIENT_FIELDS.items() if row.get(header)}
        raw_id = row.get("__crm_client_id")
        client_id = positive_id(raw_id)
        if raw_id and not client_id:
            error("Клиенты", row, "invalid_client_id", "Некорректный идентификатор клиента.", "__crm_client_id")
            continue
        matches = [c for c in clients if c["id"] == client_id] if client_id else []
        if matches and matches[0] not in accessible:
            error("Клиенты", row, "inaccessible_client", "Клиент недоступен в выбранной CRM.")
            continue
        archived_matches = [c for c in archived_accessible if c["id"] == client_id] if client_id else []
        if not matches and not archived_matches:
            archived_matches = matches_for(values, None, archived_accessible)
        if archived_matches:
            existing = archived_matches[0]
            key = existing["id"]
            action = dict(
                key=key,
                existing=existing,
                values=values,
                changes={},
                skipped=True,
                assign=False,
                assignment=next((a for a in assignments if a["crm_client_id"] == key), None),
                color=None,
            )
            resolved[key] = action
            if client_id:
                source_ids[client_id] = key
            for company in (row.get("Компания"), existing.get("document_name")):
                if normalized(company):
                    aliases.setdefault(normalized(company), set()).add(key)
            client_actions.append(action)
            preview["skippedArchived"] += 1
            continue
        if not matches:
            matches = matches_for(values, None, staged)
        if not matches and values.get("inn") and values.get("kpp"):
            same_inn = [
                client
                for client in staged
                if text(client.get("inn")) == values["inn"]
                and text(client.get("kpp")) != text(values["kpp"])
            ]
            if same_inn:
                error(
                    "Клиенты",
                    row,
                    "inn_kpp_conflict",
                    "Найден клиент с тем же ИНН, но другим КПП. Обновление или создание дубликата запрещено.",
                    "КПП",
                )
                continue
        if len(matches) > 1:
            error("Клиенты", row, "ambiguous_client", "Найдено несколько клиентов с такими реквизитами.")
            continue
        existing = matches[0] if matches else None
        if not existing and not (values.get("inn") or (values.get("document_name") and (phone(values.get("phone")) or normalized(values.get("email"))))):
            error("Клиенты", row, "insufficient_identity", "Укажите ИНН либо компанию и телефон или почту.")
            continue
        key = existing["id"] if existing else -row["_row"]
        if client_id in source_ids and source_ids[client_id] != key:
            error("Клиенты", row, "duplicate_source_client_id", "Один исходный идентификатор указан для разных клиентов.", "__crm_client_id")
            continue
        if key in seen:
            error("Клиенты", row, "duplicate_client", "Клиент повторяется в книге.")
            continue
        seen.add(key)
        changes = {field: value for field, value in values.items() if not existing or text(existing.get(field)) != value}
        skipped = bool(existing and existing.get("linked_counterparty_id") is not None)
        assignment = next((a for a in assignments if a["crm_client_id"] == key), None)
        assign = not existing or include_existing_clients
        if not skipped and assign and assignment and assignment.get("archived_at"):
            error("Клиенты", row, "archived_assignment", "Архивное назначение сначала должен восстановить администратор.")
        action = dict(key=key, existing=existing, values=values, changes=changes, skipped=skipped,
                      assign=assign, assignment=assignment, color=row.get("__color_key") if row.get("__color_key") in colors else None)
        resolved[key] = action
        if client_id:
            source_ids[client_id] = key
        for company in (row.get("Компания"), (existing or {}).get("document_name")):
            if normalized(company):
                aliases.setdefault(normalized(company), set()).add(key)
        client_actions.append(action)
        if skipped:
            preview["skippedOneCLinked"] += 1
            continue
        if existing:
            preview["clientsToUpdate" if changes else "unchangedClients"] += 1
            staged = [dict(c, **changes) if c["id"] == key else c for c in staged]
        else:
            preview["clientsToCreate"] += 1
            staged.append(dict(values, id=key, crm_owner_user_id=owner_id))
        if assign and (not assignment or assignment["tab_id"] != target_tab_id):
            preview["clientsToAssign"] += 1

    staged_contacts = list(contacts)
    seen_contacts: set[int] = set()
    for row in book.contacts:
        raw_client_id = row.get("__crm_client_id")
        key = positive_id(raw_client_id)
        if raw_client_id and not key:
            error("Контакты", row, "invalid_client_id", "Некорректный идентификатор клиента.")
            continue
        if key is not None:
            key = source_ids.get(key, key)
        if key is None:
            keys = aliases.get(normalized(row.get("Компания")), set())
            if len(keys) > 1:
                error("Контакты", row, "ambiguous_client", "Компания контакта неоднозначна.")
                continue
            key = next(iter(keys), None)
        action = resolved.get(key)
        if not action:
            error("Контакты", row, "unresolved_client", "Клиент контакта отсутствует среди распознанных строк книги.")
            continue
        if action["skipped"]:
            continue
        values = {field: row[header] for header, field in CONTACT_FIELDS.items() if row.get(header)}
        if row.get("Основной контакт"):
            primary = normalized(row["Основной контакт"])
            if primary not in {"да", "нет", "true", "false", "1", "0"}:
                error("Контакты", row, "invalid_primary", "Основной контакт: укажите Да или Нет.")
                continue
            values["is_primary"] = int(primary in {"да", "true", "1"})
        raw_id = row.get("__crm_contact_id")
        contact_id = positive_id(raw_id)
        if raw_id and not contact_id:
            error("Контакты", row, "invalid_contact_id", "Некорректный идентификатор контакта.")
            continue
        matches = [c for c in contacts if c["id"] == contact_id] if contact_id else []
        if matches and (matches[0]["crm_client_id"] != key or matches[0]["owner_user_id"] != owner_id):
            error("Контакты", row, "inaccessible_contact", "Контакт не принадлежит выбранному клиенту и владельцу.")
            continue
        if not matches:
            matches = [c for c in staged_contacts if c["crm_client_id"] == key and c["owner_user_id"] == owner_id and normalized(c.get("name")) == normalized(values.get("name")) and _same_channel(values, c)]
        if len(matches) > 1:
            error("Контакты", row, "ambiguous_contact", "Найдено несколько одинаковых контактов.")
            continue
        existing = matches[0] if matches else None
        if not existing and not (values.get("name") and (phone(values.get("phone")) or normalized(values.get("email")))):
            error("Контакты", row, "insufficient_identity", "Укажите имя контакта и телефон или почту.")
            continue
        contact_key = existing["id"] if existing else -row["_row"]
        if contact_key in seen_contacts:
            error("Контакты", row, "duplicate_contact", "Контакт повторяется в книге.")
            continue
        seen_contacts.add(contact_key)
        changes = {field: value for field, value in values.items() if not existing or text(existing.get(field)) != text(value)}
        contact_actions.append(dict(key=key, existing=existing, changes=changes, values=values))
        if not existing:
            preview["contactsToCreate"] += 1
            staged_contacts.append(dict(values, id=contact_key, crm_client_id=key, owner_user_id=owner_id))
        else:
            staged_contacts = [dict(c, **changes) if c["id"] == contact_key else c for c in staged_contacts]

    # Resolve explicit primary choices as final state before computing deltas.
    # The last explicitly primary row wins, independently of the stored primary.
    last_primary = {action["key"]: index for index, action in enumerate(contact_actions)
                    if action["values"].get("is_primary") == 1}
    for index, action in enumerate(contact_actions):
        values, existing = action["values"], action["existing"]
        if values.get("is_primary") == 1:
            values["is_primary"] = int(last_primary[action["key"]] == index)
        action["changes"] = {field: value for field, value in values.items()
                             if not existing or text(existing.get(field)) != text(value)}
        if existing and action["changes"]:
            preview["contactsToUpdate"] += 1
    return dict(preview=preview, clients=client_actions, contacts=contact_actions)
