"""CRM-specific XLSX export construction.

The API layer is responsible for authorizing the requested export scope and
collecting *all* active rows (not just the current page).  It then calls::

    build_crm_export_xlsx(client_rows=rows, contact_rows=contacts)

``client_rows`` and ``contact_rows`` are mappings from the CRM list/detail
queries.  Both snake_case and camelCase field names are accepted so the module
can be used directly while the API contract is being integrated.  This module
has no database or FastAPI dependency and never mutates its inputs.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MOSCOW_TIMEZONE = ZoneInfo("Europe/Moscow")

CLIENT_HEADERS = (
    "Компания",
    "ИНН",
    "КПП",
    "Город",
    "Сайт",
    "Основной контакт",
    "Телефон",
    "Почта",
    "Комментарий",
    "Дата звонка",
    "Напоминание",
    "Личная вкладка",
    "Связь с 1С",
)
CONTACT_HEADERS = ("Компания", "Контактное лицо", "Телефон", "Почта", "Основной контакт")
_CLIENT_ROUND_TRIP_HEADERS = ("__crm_client_id", "__color_key", "__export_version")
_CONTACT_ROUND_TRIP_HEADERS = ("__crm_client_id", "__crm_contact_id")
_CLIENT_COLUMN_WIDTHS = (30, 15, 14, 18, 28, 24, 20, 28, 42, 20, 20, 22, 24)
_CONTACT_COLUMN_WIDTHS = (30, 24, 20, 28, 18)

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TEXT_FORMAT = "@"
_DATE_FORMAT = "DD.MM.YYYY HH:MM"
_DATE_ONLY_FORMAT = "DD.MM.YYYY"
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_ROW_COLOR_HEX = {
    "red": "FCE4D6",
    "orange": "FCE4D6",
    "amber": "FFF2CC",
    "yellow": "FFF2CC",
    "lime": "E2F0D9",
    "green": "E2F0D9",
    "teal": "DDEBF7",
    "cyan": "DDEBF7",
    "blue": "DDEBF7",
    "indigo": "D9EAD3",
    "purple": "EADCF8",
    "pink": "FCE4EC",
    "gray": "D9EAD3",
}


def build_crm_export_xlsx(
    *,
    client_rows: Iterable[Mapping[str, Any]],
    contact_rows: Iterable[Mapping[str, Any]],
) -> bytes:
    """Return a safe CRM workbook with ``Клиенты`` and ``Контакты`` sheets.

    Required client information is represented by readable Russian headers.
    In addition to the documented CRM keys, compatible camelCase aliases are
    accepted.  The caller should provide contacts already filtered to the same
    authorized companies; a contact's company can be supplied as
    ``company_name``/``companyName`` or resolved from ``client_id``/``clientId``.

    Every user-controlled string is written as text.  Values whose first
    non-space character could start an Excel formula receive a leading
    apostrophe; INN, KPP and telephone columns additionally use Excel's Text
    format to retain leading zeroes.
    """
    materialized_clients = list(client_rows)
    materialized_contacts = list(contact_rows)
    client_names = {
        _value(row, "id", "client_id", "clientId"): _safe_text(
            _value(row, "document_name", "documentName", "full_name", "fullName", "company_name", "companyName", "name")
        )
        for row in materialized_clients
        if _value(row, "id", "client_id", "clientId") is not None
    }

    workbook = Workbook()
    clients_sheet = workbook.active
    clients_sheet.title = "Клиенты"
    contacts_sheet = workbook.create_sheet("Контакты")

    _prepare_sheet(
        clients_sheet,
        (*CLIENT_HEADERS, *_CLIENT_ROUND_TRIP_HEADERS),
        column_widths=_CLIENT_COLUMN_WIDTHS,
    )
    _prepare_sheet(
        contacts_sheet,
        (*CONTACT_HEADERS, *_CONTACT_ROUND_TRIP_HEADERS),
        column_widths=_CONTACT_COLUMN_WIDTHS,
    )
    _hide_columns(clients_sheet, first_column=len(CLIENT_HEADERS) + 1, count=len(_CLIENT_ROUND_TRIP_HEADERS))
    _hide_columns(contacts_sheet, first_column=len(CONTACT_HEADERS) + 1, count=len(_CONTACT_ROUND_TRIP_HEADERS))

    for row in materialized_clients:
        values = [
            _safe_text(_value(row, "document_name", "documentName", "full_name", "fullName", "company_name", "companyName", "name")),
            _safe_text(_value(row, "inn", "INN")),
            _safe_text(_value(row, "kpp", "KPP")),
            _safe_text(_value(row, "city")),
            _safe_text(_value(row, "website", "site")),
            _safe_text(_value(row, "contact_name", "contactName", "primary_contact_name", "primaryContactName")),
            _safe_text(_value(row, "phone", "telephone")),
            _safe_text(_value(row, "email")),
            _safe_text(_value(row, "last_comment", "lastComment", "comment", "notes")),
            _as_excel_date(_value(row, "last_call_at", "lastCallAt", "call_at", "callAt")),
            _as_excel_date(_value(row, "reminder_at", "reminderAt", "next_reminder_at", "nextReminderAt")),
            _safe_text(_value(row, "tab_name", "tabName") or "Без вкладки"),
            _sync_status_label(_value(row, "sync_status", "syncStatus", "onec_status", "onecStatus")),
            _value(row, "id", "client_id", "clientId"),
            _value(row, "color_key", "colorKey"),
            1,
        ]
        clients_sheet.append(values)
        current_row = clients_sheet.max_row
        _format_client_row(clients_sheet, current_row)
        _apply_row_color(clients_sheet, current_row, _value(row, "color_key", "colorKey"))

    for row in materialized_contacts:
        company = _value(row, "company_name", "companyName", "company", "client_name", "clientName")
        if company is None:
            company = client_names.get(_value(row, "client_id", "clientId"), "")
        contacts_sheet.append(
            [
                _safe_text(company),
                _safe_text(_value(row, "name", "contact_name", "contactName")),
                _safe_text(_value(row, "phone", "telephone")),
                _safe_text(_value(row, "email")),
                "Да" if _value(row, "is_primary", "isPrimary") else "Нет",
                _value(row, "client_id", "clientId"),
                _value(row, "id", "contact_id", "contactId"),
            ]
        )
        _format_contact_row(contacts_sheet, contacts_sheet.max_row)

    _finish_sheet(clients_sheet, len(CLIENT_HEADERS))
    _finish_sheet(contacts_sheet, len(CONTACT_HEADERS))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _value(row: Mapping[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and row[key] is not None:
            return row[key]
    return None


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _as_excel_date(value: Any) -> date | datetime | str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.astimezone(MOSCOW_TIMEZONE).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        normalized = value.replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(normalized)
        except ValueError:
            return _safe_text(value)
        return parsed.astimezone(MOSCOW_TIMEZONE).replace(tzinfo=None) if parsed.tzinfo else parsed
    return _safe_text(value)


def _sync_status_label(value: Any) -> str:
    labels = {
        "linked": "Связан с 1С",
        "synced": "Связан с 1С",
        "local": "Локальный",
        "pending": "Ожидает синхронизации",
        "error": "Ошибка синхронизации",
    }
    text = _safe_text(value)
    return labels.get(text.lower(), text)


def _prepare_sheet(sheet: Any, headers: tuple[str, ...], *, column_widths: tuple[int, ...]) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _format_client_row(sheet: Any, row_number: int) -> None:
    for column in (2, 3, 7):
        sheet.cell(row_number, column).number_format = _TEXT_FORMAT
    for column in (10, 11):
        cell = sheet.cell(row_number, column)
        if isinstance(cell.value, datetime):
            cell.number_format = _DATE_FORMAT
        elif isinstance(cell.value, date):
            cell.number_format = _DATE_ONLY_FORMAT
    sheet.cell(row_number, 9).alignment = Alignment(vertical="top", wrap_text=True)


def _format_contact_row(sheet: Any, row_number: int) -> None:
    sheet.cell(row_number, 3).number_format = _TEXT_FORMAT


def _hide_columns(sheet: Any, *, first_column: int, count: int) -> None:
    for column_index in range(first_column, first_column + count):
        sheet.column_dimensions[get_column_letter(column_index)].hidden = True


def _apply_row_color(sheet: Any, row_number: int, color_key: Any) -> None:
    color = _ROW_COLOR_HEX.get(str(color_key or "").lower())
    if color is None:
        return
    fill = PatternFill("solid", fgColor=color)
    for cell in sheet[row_number]:
        cell.fill = fill


def _finish_sheet(sheet: Any, column_count: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{sheet.max_row}"
