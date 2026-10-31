from __future__ import annotations

import base64
from datetime import datetime, timezone
from io import BytesIO
from decimal import Decimal
from pathlib import Path
import tempfile
import os
import re
import shutil
import time
import uuid
from contextlib import contextmanager
from collections import Counter
from copy import copy

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse
from fastapi.routing import APIRoute
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field, StrictInt, StrictStr
from starlette.concurrency import run_in_threadpool

from stock_sync_web.assistant_extraction import (CodexDocumentExtractor, DocumentSourceReader,
    ExtractionError, ExtractionRequest, FORMATS, validate_payload, create_document_extractor)
from stock_sync_web.assistant_service import _extraction_slots, require_document_reader
from stock_sync_web.assistant_storage import UploadLimits
from stock_sync_web.service import CLIENT_PRICE_TEMPLATE_PATH


def _reparse(path):
    stat = path.lstat()
    return path.is_symlink() or bool(getattr(stat, 'st_file_attributes', 0) & 0x400)


def _lock(file, blocking):
    if os.name == 'nt':
        import msvcrt
        file.seek(0)
        try: msvcrt.locking(file.fileno(), msvcrt.LK_LOCK if blocking else msvcrt.LK_NBLCK, 1)
        except OSError: return False
    else:
        import fcntl
        try: fcntl.flock(file.fileno(), fcntl.LOCK_EX | (0 if blocking else fcntl.LOCK_NB))
        except OSError: return False
    return True


def _unlock(file):
    if os.name == 'nt':
        import msvcrt
        file.seek(0); msvcrt.locking(file.fileno(), msvcrt.LK_UNLCK, 1)
    else:
        import fcntl
        fcntl.flock(file.fileno(), fcntl.LOCK_UN)


def _safe_tree(path):
    # Do not traverse or delete a tree containing a link or Windows junction.
    for root, dirs, files in os.walk(path, followlinks=False):
        for name in dirs + files:
            if _reparse(Path(root) / name): return False
    return True


def _workspace_root():
    root = Path(tempfile.gettempdir()) / 'sm-techno-stock-valuation-v1'
    root.mkdir(mode=0o700, exist_ok=True)
    if _reparse(root) or not root.is_dir(): raise RuntimeError('Unsafe temporary workspace')
    if os.name != 'nt':
        stat = root.stat()
        if stat.st_uid != os.getuid() or stat.st_mode & 0o077: raise RuntimeError('Unsafe temporary workspace')
    return root


def _scavenge(root):
    for path in root.iterdir():
        if not re.fullmatch(r'stock-[0-9a-f]{32}', path.name): continue
        try:
            if _reparse(path) or not path.is_dir() or time.time() - path.stat().st_mtime < 600: continue
            lease = path / '.lease'
            if _reparse(lease) or not lease.is_file(): continue
            with lease.open('r+b') as lock:
                if not _lock(lock, False): continue
                try:
                    lock.seek(0)
                    if lock.read() != b'sm-techno-stock-valuation-v1': continue
                    if not _safe_tree(path): continue
                    os.utime(path, None)
                finally: _unlock(lock)
            shutil.rmtree(path)
        except (OSError, ValueError):
            continue


@contextmanager
def _stock_workspace():
    root = _workspace_root()
    _scavenge(root)
    path = root / ('stock-' + uuid.uuid4().hex)
    path.mkdir(mode=0o700)
    try:
        with (path / '.lease').open('x+b') as lease:
            lease.write(b'sm-techno-stock-valuation-v1'); lease.flush()
            if not _lock(lease, True): raise RuntimeError('Cannot acquire workspace lease')
            try: yield path
            finally:
                os.utime(path, None)
                _unlock(lease)
    finally:
        if path.exists() and not _reparse(path) and _safe_tree(path): shutil.rmtree(path)


class PrivateRoute(APIRoute):
    def get_route_handler(self):
        handler = super().get_route_handler()
        async def private_handler(request):
            try:
                response = await handler(request)
            except RequestValidationError:
                response = JSONResponse({'detail': 'Проверьте данные запроса.'}, status_code=422)
            except HTTPException as exc:
                response = JSONResponse({'detail': exc.detail}, status_code=exc.status_code, headers=exc.headers)
            except Exception:
                response = JSONResponse({'detail': 'Не удалось выполнить запрос. Повторите попытку.'}, status_code=500)
            response.headers['Cache-Control'] = 'no-store'
            return response
        return private_handler


class Row(BaseModel):
    model_config = ConfigDict(extra='forbid')
    article: StrictStr = Field(min_length=1, max_length=5000)
    requestedQty: StrictStr | None = None
    sourceRow: StrictInt = Field(ge=1)

class ResolveRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    rows: list[Row] = Field(min_length=1, max_length=1000)


def _authorize(actor):
    try: require_document_reader(actor)
    except Exception: raise HTTPException(403, 'Доступ запрещён.') from None


def _payload(rows):
    payload = {'rows': [{'article': r.article, 'requested_qty': r.requestedQty, 'source_row': r.sourceRow} for r in rows]}
    validate_payload('articles', payload)
    if any(not r.article.strip() for r in rows): raise ValueError('Укажите артикул каждой строки.')
    return payload


def _workbook(rows):
    book = load_workbook(CLIENT_PRICE_TEMPLATE_PATH)
    # Retain the actual template's column sizing/styles, but remove every sample,
    # formula, customer field, hidden sheet, relationship and embedded drawing.
    source = book.active
    widths = [source.column_dimensions[chr(65+i)].width for i in range(7)]
    header_styles = [(copy(source.cell(1, i).font), copy(source.cell(1, i).fill), copy(source.cell(1, i).border), copy(source.cell(1, i).alignment)) for i in range(1, 8)]
    book.close()
    book = Workbook()
    ws = book.active
    ws.title = 'Расценка'
    for i, width in enumerate(widths): ws.column_dimensions[chr(65+i)].width = width or 18
    ws.append(['Артикул', 'Количество запроса', 'Статус', 'Остаток', 'Цена', 'Склад', 'Строка источника'])
    for cell, (font, fill, border, alignment) in zip(ws[1], header_styles):
        cell.font, cell.fill, cell.border, cell.alignment = font, fill, border, alignment
    for row in rows:
        for variant in row['variants'] or [None]:
            values = [row['article'], row['requestedQty'], row['status'], variant['physicalStock'] if variant else None,
                      variant['price'] if variant else None, variant['warehouseName'] if variant else None, row['sourceRow']]
            ws.append(values)
            # Force text cells, including leading-zero and formula-like strings.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str): cell.data_type = 's'
    output = BytesIO(); book.save(output); book.close()
    return base64.b64encode(output.getvalue()).decode('ascii')


def create_stock_valuation_router(current_user, get_service) -> APIRouter:
    router = APIRouter(prefix='/api/stock/valuation', route_class=PrivateRoute)

    @router.post('/intake')
    async def intake(response: Response, file: UploadFile = File(...), actor=Depends(current_user)):
        response.headers['Cache-Control'] = 'no-store'
        _authorize(actor)
        suffix = Path(file.filename or '').suffix.lower()
        if suffix not in FORMATS: raise HTTPException(400, 'Этот формат не поддерживается.')
        limits = UploadLimits()
        try:
            with _stock_workspace() as workspace:
                source = workspace / ('source' + suffix)
                size = 0
                with source.open('wb') as target:
                    while chunk := await file.read(1024 * 1024):
                        size += len(chunk)
                        if size > limits.max_bytes: raise HTTPException(413, 'Превышен допустимый размер файла.')
                        target.write(chunk)
                def extract_source():
                    if not _extraction_slots.acquire(blocking=False):
                        raise HTTPException(429, 'Распознавание занято. Повторите попытку позже.')
                    try:
                        texts, images = DocumentSourceReader(limits).read(source, workspace)
                        return create_document_extractor(codex_factory=CodexDocumentExtractor).extract(ExtractionRequest('articles', texts, images, []), temp_root=workspace)
                    finally:
                        _extraction_slots.release()
                result = await run_in_threadpool(extract_source)
                validate_payload('articles', result.payload)
                rows = [Row(article=r['article'], requestedQty=r['requested_qty'], sourceRow=r['source_row']) for r in result.payload['rows']]
                _payload(rows)
                if not rows: raise ValueError('В файле не найдено артикулов.')
                return {'rows': [r.model_dump() for r in rows], 'warnings': result.warnings}
        except HTTPException: raise
        except ExtractionError: raise HTTPException(400, 'Не удалось распознать файл. Проверьте формат или повторите попытку.') from None
        except Exception: raise HTTPException(400, 'Не удалось распознать файл. Проверьте формат или повторите попытку.') from None
        finally: await file.close()

    @router.post('/resolve')
    def resolve(body: ResolveRequest, response: Response, actor=Depends(current_user)):
        response.headers['Cache-Control'] = 'no-store'; _authorize(actor)
        try: _payload(body.rows)
        except ValueError: raise HTTPException(400, 'Проверьте артикулы и количество.') from None
        db = get_service().db
        keys = [db._normalize_stock_sku(r.article) for r in body.rows]
        counts = Counter(keys)
        quantities = {}
        for key, row in zip(keys, body.rows): quantities.setdefault(key, set()).add(Decimal(row.requestedQty) if row.requestedQty else None)
        snapshot = db.lookup_stock_articles([r.article for r in body.rows])
        seen = set(); resolved = []
        for key, row in zip(keys, body.rows):
            matches = snapshot.get(key, [])
            ids = {item['item_id'] for item in matches}
            variants = [{'itemId': x['item_id'], 'warehouseId': x['warehouse_id'], 'warehouseName': x['warehouse_name'],
                         'physicalStock': x['quantity'], 'price': x['price']} for x in matches if x['warehouse_id'] is not None and x['quantity'] > 0]
            if counts[key] > 1 and len(quantities[key]) > 1: status = 'ambiguous'; variants = []
            elif key in seen: status = 'duplicate'; variants = []
            elif not matches: status = 'not_found'; variants = []
            elif len(ids) != 1: status = 'ambiguous'; variants = []
            elif not variants: status = 'zero_stock'
            else: status = 'matched'
            seen.add(key)
            resolved.append({**row.model_dump(), 'status': status, 'variants': variants})
        return {'resolvedAt': datetime.now(timezone.utc).isoformat(), 'articleResolution': resolved,
                'workbook': {'filename': 'stock-valuation.xlsx', 'mediaType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'base64': _workbook(resolved)}}
    return router
