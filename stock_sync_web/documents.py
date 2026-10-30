from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate


@dataclass(frozen=True)
class DocumentLineInput:
    row_no: int
    article: str
    name: str
    qty: float
    price: float


CLIENT_TEMPLATE_FIELDS = [
    "document_name",
    "full_name",
    "inn",
    "kpp",
    "ogrn",
    "legal_address",
    "actual_address",
    "bank_name",
    "bank_bik",
    "bank_account",
    "correspondent_account",
    "signer_position",
    "signer_position_genitive",
    "signer_name",
    "signer_name_genitive",
    "signer_short_name",
    "signer_basis",
    "contact_person",
    "email",
    "phone",
]


def build_document_context(
    *,
    client: dict[str, Any],
    document_number: str,
    document_date: str | date | datetime,
    lines: list[DocumentLineInput] | None = None,
) -> dict[str, Any]:
    prepared_lines = [_line_to_context(line) for line in (lines or [])]
    total_amount = sum(float(line.qty or 0) * float(line.price or 0) for line in (lines or []))

    return {
        "client": _client_to_context(client),
        "document": {
            "number": str(document_number or "").strip(),
            "date": format_document_date(document_date),
        },
        "spec": {
            "lines": prepared_lines,
            "total_amount": format_money(total_amount),
        },
    }


def find_missing_client_fields(context: dict[str, Any], required_fields: list[str]) -> list[str]:
    client = context.get("client") if isinstance(context, dict) else {}
    if not isinstance(client, dict):
        return required_fields
    return [field for field in required_fields if not str(client.get(field) or "").strip()]


def generate_document_docx(
    *,
    template_path: Path | str,
    output_path: Path | str,
    context: dict[str, Any],
) -> None:
    doc = DocxTemplate(template_path)
    doc.render(context, autoescape=True)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    polish_generated_document_layout(output)


def format_document_date(value: str | date | datetime) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")

    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return date.fromisoformat(text[:10]).strftime("%d.%m.%Y")
    except ValueError:
        return text


def _client_to_context(client: dict[str, Any]) -> dict[str, str]:
    bank_name_or_bik = _get_value(client, "bank_name_or_bik", "bankNameOrBik")
    signer_position = _get_value(client, "signer_position", "signerPosition")
    signer_name = _get_value(client, "signer_name", "signerName")
    document_name = _get_value(client, "document_name", "documentName")
    full_name = _get_value(client, "full_name", "fullName")
    display_name = document_name or full_name
    values = {
        "document_name": display_name,
        "full_name": full_name or document_name,
        "inn": _get_value(client, "inn"),
        "kpp": _get_value(client, "kpp"),
        "ogrn": _get_value(client, "ogrn"),
        "legal_address": _get_value(client, "legal_address", "legalAddress"),
        "actual_address": _get_value(client, "actual_address", "actualAddress"),
        "bank_name": _get_value(client, "bank_name", "bankName") or bank_name_or_bik,
        "bank_bik": _get_value(client, "bank_bik", "bankBik"),
        "bank_account": _get_value(client, "bank_account", "bankAccount"),
        "correspondent_account": _get_value(client, "correspondent_account", "correspondentAccount"),
        "signer_position": signer_position,
        "signer_position_genitive": decline_signer_position_genitive(signer_position),
        "signer_name": signer_name,
        "signer_name_genitive": decline_person_name_genitive(signer_name),
        "signer_short_name": format_person_short_name(signer_name),
        "signer_basis": _get_value(client, "signer_basis", "signerBasis"),
        "contact_person": _get_value(client, "contact_person", "contactPerson"),
        "email": _get_value(client, "email"),
        "phone": _get_value(client, "phone"),
    }
    return {field: str(values.get(field) or "").strip() for field in CLIENT_TEMPLATE_FIELDS}


def _line_to_context(line: DocumentLineInput) -> dict[str, str]:
    amount = float(line.qty or 0) * float(line.price or 0)
    return {
        "row_no": str(line.row_no),
        "article": str(line.article or ""),
        "name": str(line.name or ""),
        "qty": format_quantity(line.qty),
        "price": format_money(line.price),
        "amount": format_money(amount),
    }


def _get_value(source: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = source.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def polish_generated_document_layout(path: Path | str) -> None:
    document = WordDocument(path)
    changed = _remove_empty_paragraphs_between_numbered_blocks(document)
    changed = _compact_requisites_section(document) or changed
    if changed:
        document.save(path)


def _remove_empty_paragraphs_between_numbered_blocks(document) -> bool:
    changed = False
    for child in list(document.element.body.iterchildren()):
        if not _is_empty_paragraph_element(child, document):
            continue
        previous = child.getprevious()
        next_element = child.getnext()
        previous_text = _paragraph_text(previous, document) if _is_paragraph_element(previous) else ""
        next_text = _paragraph_text(next_element, document) if _is_paragraph_element(next_element) else ""
        if _starts_with_clause_number(previous_text) or _starts_with_clause_number(next_text):
            child.getparent().remove(child)
            changed = True
    return changed


def _compact_requisites_section(document) -> bool:
    heading_element = _find_requisites_heading_element(document)
    if heading_element is None:
        return False

    changed = False
    heading = Paragraph(heading_element, document)
    if heading.paragraph_format.keep_with_next is not True:
        heading.paragraph_format.keep_with_next = True
        changed = True

    while _is_empty_paragraph_element(heading_element.getprevious(), document):
        heading_element.getparent().remove(heading_element.getprevious())
        changed = True
    while _is_empty_paragraph_element(heading_element.getnext(), document):
        heading_element.getparent().remove(heading_element.getnext())
        changed = True

    table_element = heading_element.getnext()
    if table_element is None or table_element.tag != qn("w:tbl"):
        return changed

    table = Table(table_element, document)
    changed = _remove_empty_trailing_rows(table) or changed
    changed = _compact_table_cell_blank_paragraphs(table) or changed
    changed = _keep_table_together(table) or changed
    while _is_empty_paragraph_element(table_element.getnext(), document):
        table_element.getparent().remove(table_element.getnext())
        changed = True
    return changed


def _find_requisites_heading_element(document):
    for child in document.element.body.iterchildren():
        if _is_paragraph_element(child) and _paragraph_text(child, document).strip().startswith("13."):
            return child
    return None


def _remove_empty_trailing_rows(table: Table) -> bool:
    changed = False
    for row in list(table.rows)[::-1]:
        if any(cell.text.strip() for cell in row.cells):
            break
        table._tbl.remove(row._tr)
        changed = True
    return changed


def _keep_table_together(table: Table) -> bool:
    changed = False
    cell_paragraphs: list[Paragraph] = []
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
            changed = True
        for cell in row.cells:
            cell_paragraphs.extend(cell.paragraphs)

    for paragraph in cell_paragraphs[:-1]:
        if paragraph.paragraph_format.keep_with_next is not True:
            paragraph.paragraph_format.keep_with_next = True
            changed = True
    return changed


def _compact_table_cell_blank_paragraphs(table: Table, max_blank: int = 2) -> bool:
    changed = False
    for row in table.rows:
        for cell in row.cells:
            blank_count = 0
            for paragraph in list(cell.paragraphs):
                if paragraph.text.strip():
                    blank_count = 0
                    continue
                blank_count += 1
                if blank_count <= max_blank or len(cell.paragraphs) <= 1:
                    continue
                paragraph._element.getparent().remove(paragraph._element)
                changed = True
    return changed


def _is_paragraph_element(element) -> bool:
    return element is not None and element.tag == qn("w:p")


def _is_empty_paragraph_element(element, document) -> bool:
    return _is_paragraph_element(element) and not _paragraph_text(element, document).strip()


def _paragraph_text(element, document) -> str:
    return Paragraph(element, document).text if element is not None else ""


def _starts_with_clause_number(value: str) -> bool:
    text = value.strip()
    return bool(text) and text[0].isdigit() and "." in text[:4]


def decline_signer_position_genitive(value: str) -> str:
    text = _normalize_spaces(value)
    if not text:
        return ""
    if not _has_cyrillic(text):
        return text
    normalized = text.lower()
    position_map = {
        "генеральный директор": "генерального директора",
        "директор": "директора",
        "исполнительный директор": "исполнительного директора",
        "коммерческий директор": "коммерческого директора",
        "заместитель директора": "заместителя директора",
        "главный бухгалтер": "главного бухгалтера",
        "индивидуальный предприниматель": "индивидуального предпринимателя",
        "ип": "индивидуального предпринимателя",
    }
    if normalized in position_map:
        return position_map[normalized]

    word_map = {
        "генеральный": "генерального",
        "исполнительный": "исполнительного",
        "коммерческий": "коммерческого",
        "главный": "главного",
        "директор": "директора",
        "заместитель": "заместителя",
        "бухгалтер": "бухгалтера",
        "предприниматель": "предпринимателя",
    }
    return " ".join(word_map.get(word, word) for word in normalized.split())


def decline_person_name_genitive(value: str) -> str:
    text = _normalize_spaces(value)
    if not text or not _has_cyrillic(text) or "." in text:
        return text

    parts = text.split()
    if len(parts) < 2:
        return text

    gender = _detect_person_gender(parts)
    declined: list[str] = []
    for index, part in enumerate(parts):
        if index == 0:
            declined.append(_decline_hyphenated(part, lambda item: _decline_surname(item, gender)))
        elif index == 1:
            declined.append(_decline_hyphenated(part, lambda item: _decline_given_name(item, gender)))
        elif index == 2:
            declined.append(_decline_hyphenated(part, lambda item: _decline_patronymic(item, gender)))
        else:
            declined.append(part)
    return " ".join(declined)


def format_person_short_name(value: str) -> str:
    text = _normalize_spaces(value)
    if not text or not _has_cyrillic(text) or "." in text:
        return text

    parts = text.split()
    if len(parts) < 2:
        return text

    initials = "".join(f"{part[0].upper()}." for part in parts[1:3] if part)
    return f"{parts[0]} {initials}".strip()


def _normalize_spaces(value: str) -> str:
    return " ".join(str(value or "").strip().split())


def _has_cyrillic(value: str) -> bool:
    return any("а" <= char.lower() <= "я" or char.lower() == "ё" for char in value)


def _detect_person_gender(parts: list[str]) -> str:
    if len(parts) >= 3:
        patronymic = parts[2].lower()
        if patronymic.endswith("ич"):
            return "male"
        if patronymic.endswith(("на", "кызы")):
            return "female"
    first_name = parts[1].lower() if len(parts) > 1 else ""
    if first_name.endswith(("а", "я")) and first_name not in {"илья", "никита", "лука", "кузьма", "фома"}:
        return "female"
    return "male"


def _decline_hyphenated(value: str, decline_part) -> str:
    return "-".join(decline_part(part) for part in value.split("-"))


def _decline_surname(value: str, gender: str) -> str:
    lower = value.lower()
    if gender == "female":
        if lower.endswith(("ова", "ева", "ёва", "ина", "ына")):
            return _match_case(value, lower[:-1] + "ой")
        if lower.endswith("ая"):
            return _match_case(value, lower[:-2] + "ой")
        if lower.endswith("яя"):
            return _match_case(value, lower[:-2] + "ей")
        if lower.endswith("а"):
            return _match_case(value, lower[:-1] + _genitive_a_ending(lower[:-1]))
        if lower.endswith("я"):
            return _match_case(value, lower[:-1] + "и")
        return value

    if lower.endswith(("ов", "ев", "ёв", "ин", "ын")):
        return _match_case(value, lower + "а")
    if lower.endswith(("ский", "цкий", "кий", "ой", "ый")):
        return _match_case(value, lower[:-2] + "ого")
    if lower.endswith("ий"):
        return _match_case(value, lower[:-2] + "ия")
    if _ends_with_consonant(lower):
        return _match_case(value, lower + "а")
    return value


def _decline_given_name(value: str, gender: str) -> str:
    lower = value.lower()
    special = {
        "лев": "льва",
        "павел": "павла",
        "пётр": "петра",
        "петр": "петра",
        "илья": "ильи",
        "никита": "никиты",
        "лука": "луки",
        "кузьма": "кузьмы",
        "фома": "фомы",
        "любовь": "любови",
    }
    if lower in special:
        return _match_case(value, special[lower])

    if gender == "female":
        if lower.endswith("ия"):
            return _match_case(value, lower[:-1] + "и")
        if lower.endswith("а"):
            return _match_case(value, lower[:-1] + _genitive_a_ending(lower[:-1]))
        if lower.endswith("я"):
            return _match_case(value, lower[:-1] + "и")
        if lower.endswith("ь"):
            return _match_case(value, lower[:-1] + "и")
        return value

    if lower.endswith("ий"):
        return _match_case(value, lower[:-2] + "ия")
    if lower.endswith("й"):
        return _match_case(value, lower[:-1] + "я")
    if lower.endswith("ь"):
        return _match_case(value, lower[:-1] + "я")
    if lower.endswith("а"):
        return _match_case(value, lower[:-1] + _genitive_a_ending(lower[:-1]))
    if _ends_with_consonant(lower):
        return _match_case(value, lower + "а")
    return value


def _decline_patronymic(value: str, gender: str) -> str:
    lower = value.lower()
    if lower.endswith("ич"):
        return _match_case(value, lower + "а")
    if gender == "female" and lower.endswith("на"):
        return _match_case(value, lower[:-1] + _genitive_a_ending(lower[:-1]))
    return _decline_given_name(value, gender)


def _genitive_a_ending(stem: str) -> str:
    return "и" if stem.endswith(("г", "к", "х", "ж", "ч", "ш", "щ")) else "ы"


def _ends_with_consonant(value: str) -> bool:
    return bool(value) and value[-1] in "бвгджзклмнпрстфхцчшщ"


def _match_case(source: str, lower_value: str) -> str:
    if source.isupper():
        return lower_value.upper()
    if source[:1].isupper():
        return lower_value[:1].upper() + lower_value[1:]
    return lower_value


def format_quantity(value: float | int | None) -> str:
    number = float(value or 0)
    if number.is_integer():
        return str(int(number))
    return f"{number:,.3f}".rstrip("0").rstrip(".").replace(",", " ")


def format_money(value: float | int | None) -> str:
    return f"{float(value or 0):,.2f}".replace(",", " ")


# Invoice specification uses its own source-preserving template. Legacy generators stay unchanged.
def read_docx_xml(content):
    from io import BytesIO
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(BytesIO(content)) as archive:
        return etree.fromstring(archive.read('word/document.xml'),etree.XMLParser(resolve_entities=False,no_network=True))


_WORD_NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def _xml_text(node):
    return ''.join(node.xpath('.//w:t/text()',namespaces=_WORD_NS))


def docx_source_text(content):
    root=read_docx_xml(content)
    def blocks(node):
        for child in node:
            if child.tag==qn('w:p'):yield _xml_text(child)
            elif child.tag==qn('w:tbl'):
                for row in child.findall('w:tr',_WORD_NS):
                    yield '\t'.join('\n'.join(blocks(cell)) for cell in row.findall('w:tc',_WORD_NS))
    rows=list(blocks(root.find('w:body',_WORD_NS)))
    # Header/footer requisites remain readable without dereferencing media or external relationships.
    from io import BytesIO
    from zipfile import ZipFile
    from lxml import etree
    import re
    with ZipFile(BytesIO(content)) as archive:
        for kind in ('header','footer'):
            for name in archive.namelist():
                if re.fullmatch(r'word/'+kind+r'\d+\.xml',name):
                    part=etree.fromstring(archive.read(name),etree.XMLParser(resolve_entities=False,no_network=True))
                    rows.extend(value for value in blocks(part) if value.strip())
    return '\n'.join(rows)


def _normal_text(value):
    import re
    return re.sub(r'\s+',' ',str(value or '').replace('«','"').replace('»','"').replace('ё','е')).strip().casefold()


def _invoice_decimal(value, *, money=False):
    from decimal import Decimal, InvalidOperation
    import re
    text=re.sub(r'[\s\u00a0\u202f]','',str(value or '')).replace(',','.')
    if not re.fullmatch(r'\d+(?:\.\d{1,6})?',text): raise ValueError('В счёте отсутствует корректное число.')
    try:number=Decimal(text)
    except InvalidOperation: raise ValueError('Некорректное число в счёте.') from None
    if number>Decimal('1000000000000') or money and number.as_tuple().exponent < -2:
        raise ValueError('Проверьте денежную точность и пределы счёта.')
    return number


def rubles_in_words(value):
    """Deterministic RUB boundary, 0..999 billion rubles; no tax-rate assumptions."""
    from decimal import Decimal, ROUND_HALF_UP
    amount=Decimal(value).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
    if amount<0 or amount>=Decimal('1000000000000'): raise ValueError('Сумма вне диапазона прописи.')
    integer=int(amount);kopecks=int((amount-integer)*100)
    units=['','один','два','три','четыре','пять','шесть','семь','восемь','девять']
    teens=['десять','одиннадцать','двенадцать','тринадцать','четырнадцать','пятнадцать','шестнадцать','семнадцать','восемнадцать','девятнадцать']
    tens=['','','двадцать','тридцать','сорок','пятьдесят','шестьдесят','семьдесят','восемьдесят','девяносто']
    hundreds=['','сто','двести','триста','четыреста','пятьсот','шестьсот','семьсот','восемьсот','девятьсот']
    def form(n,forms):return forms[2] if 11<=n%100<=14 else forms[0] if n%10==1 else forms[1] if 2<=n%10<=4 else forms[2]
    words=[]
    for scale,forms,female in [(10**9,('миллиард','миллиарда','миллиардов'),False),(10**6,('миллион','миллиона','миллионов'),False),(1000,('тысяча','тысячи','тысяч'),True),(1,('рубль','рубля','рублей'),False)]:
        group=(integer//scale)%1000
        if not group and scale!=1:continue
        if group:
            if group//100:words.append(hundreds[group//100])
            rest=group%100
            if 10<=rest<20:words.append(teens[rest-10])
            else:
                if rest//10:words.append(tens[rest//10])
                if rest%10:words.append(('одна' if rest%10==1 else 'две') if female and rest%10 in (1,2) else units[rest%10])
        elif integer==0:words.append('ноль')
        words.append(form(group,forms))
    words.extend([f'{kopecks:02d}',form(kopecks,('копейка','копейки','копеек'))])
    return ' '.join(words).capitalize()


def validate_invoice_source(content,payload):
    """Reconcile every goods row with immutable DOCX body XML, independently of model arithmetic."""
    from decimal import Decimal,ROUND_HALF_UP
    import re
    root=read_docx_xml(content);rows=[];summaries=[];found=False
    for table in root.xpath('.//w:tbl',namespaces=_WORD_NS):
        header=False
        for row in table.findall('w:tr',_WORD_NS):
            cells=[_xml_text(c).strip() for c in row.findall('w:tc',_WORD_NS)]
            normalized=[_normal_text(c) for c in cells]
            if '№' in normalized and any('цен' in c and 'ндс' in c for c in normalized) and any('товар' in c or 'наименование' in c for c in normalized):
                start=normalized.index('№');header=True;found=True;continue
            if not header or not any(cells):continue
            values=cells[start:start+8]
            if values==list(map(str,range(1,9))):continue
            if len(values)<8:
                numeric=[v for v in values if v.strip()]
                if rows and values and not values[0] and len(numeric)==2 and all(re.fullmatch(r'[\d\s\u00a0\u202f.,]+',v) for v in numeric) and row.xpath('./w:tc/w:tcPr/w:gridSpan',namespaces=_WORD_NS):
                    summaries.append([_invoice_decimal(v) for v in numeric]);continue
                if any(re.search(r'итого|всего|ндс|наименований',v,re.I) for v in values):continue
                if any(values):raise ValueError('Неполная товарная строка исходного счёта.')
                continue
            if not values[0]:
                if any(values[1:]) and not any(re.search(r'итого|всего|ндс',v,re.I) for v in values):raise ValueError('Строка счёта без номера требует исправления.')
                continue
            if not re.fullmatch(r'[1-9]\d*',values[0]):raise ValueError('Проверьте товарные строки исходного счёта.')
            rows.append(values)
    lines=payload.get('lines') or []
    if not found or not rows or len(lines)!=len(rows):raise ValueError('Количество товарных строк не совпадает с исходным счётом.')
    if len(rows)>1000:raise ValueError('Слишком много товарных строк.')
    total=Decimal(0);vat=Decimal(0);seen=set()
    def cents(n):return n.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
    canonical_lines=[]
    for raw,line in zip(rows,lines):
        if int(raw[0]) in seen or line.get('row_no')!=int(raw[0]):raise ValueError('Номера товарных строк не совпадают с источником.')
        seen.add(int(raw[0]))
        if _normal_text(raw[1]) not in {_normal_text(line.get('name')),_normal_text(str(line.get('article') or '')+' '+str(line.get('name') or ''))} or _normal_text(line.get('unit'))!=_normal_text(raw[3]):
            raise ValueError('Наименование или единица товара не совпадают с исходным счётом.')
        article=str(line.get('article') or '').strip()
        if article and not re.search(r'(?<![\w-])'+re.escape(article)+r'(?![\w-])',raw[1],re.I):raise ValueError('Артикул не найден в строке счёта.')
        numbers=[]
        for key,index in [('qty',2),('unit_price_incl_vat',4),('unit_price_excl_vat',5),('vat_amount',6),('total_amount',7)]:
            expected=_invoice_decimal(raw[index],money=key!='qty');actual=_invoice_decimal(line.get(key),money=key!='qty')
            if expected!=actual:raise ValueError('Числа строки не совпадают с исходным счётом: '+key)
            numbers.append(actual)
        qty,incl,excl,tax,amount=numbers
        if qty<=0 or excl>incl or tax>amount:raise ValueError('Недопустимые количество, цена или НДС.')
        if cents(qty*incl)!=amount:raise ValueError('Количество × цена с НДС не совпадает с суммой строки.')
        # Source may round a unit price to kopecks before multiplying; tax is already a line amount.
        tolerance=max(Decimal('.01'),qty*Decimal('.005'))
        if abs(amount-tax-qty*excl)>tolerance:raise ValueError('Цена без НДС и сумма НДС не сходятся с суммой строки.')
        canonical_lines.append({**line,'name':raw[1]})
        total+=amount;vat+=tax
    if total!=_invoice_decimal(payload.get('total_amount'),money=True) or vat!=_invoice_decimal(payload.get('vat_amount'),money=True):
        raise ValueError('Итоги товаров и НДС не сходятся со счётом.')
    if any(summary!=[sum(_invoice_decimal(line['qty']) for line in lines),total] for summary in summaries):
        raise ValueError('Служебный итог количества и суммы не совпадает с товарами счёта.')
    texts=[_xml_text(p) for p in root.xpath('.//w:p',namespaces=_WORD_NS)]
    # Includes table rows because totals are often split between label and value cells.
    texts += [' '.join(_xml_text(c) for c in r.findall('w:tc',_WORD_NS)) for r in root.xpath('.//w:tr',namespaces=_WORD_NS)]
    text='\n'.join(texts)
    for label,expected in [(r'Итого',total),(r'В том числе НДС',vat),(r'Всего к оплате',total)]:
        matches=re.findall(label+r'\s*:?\s*([\d][\d \u00a0\u202f]*(?:[,.]\d{1,2})?)',text,re.I)
        if not matches or any(_invoice_decimal(m,money=True)!=expected for m in matches):raise ValueError('Итог не подтверждён исходным счётом: '+label)
    count=re.search(r'Всего наименований\s+(\d+)',text,re.I)
    if count and int(count.group(1))!=len(rows):raise ValueError('Число наименований счёта не совпадает с таблицей.')
    words=str(payload.get('amount_in_words') or '').strip()
    if _normal_text(words)!=_normal_text(rubles_in_words(total)) or _normal_text(words) not in _normal_text(text):
        raise ValueError('Сумма прописью не совпадает с проверенной суммой и исходным счётом.')
    if payload.get('currency') not in ('RUB','руб','руб.','рубль','рубли','рублей'):raise ValueError('Для спецификации требуется счёт в рублях.')
    buyer=re.search(r'(?:Покупатель|Заказчик)\s*:?([^\n]*(?:\n[^\n]*)?)',text,re.I)
    if not buyer:raise ValueError('Не найден блок покупателя исходного счёта.')
    buyer_text=_normal_text(buyer.group(1))
    for key in ('buyer_document_name','buyer_inn'):
        if not payload.get(key) or _normal_text(payload[key]) not in buyer_text:raise ValueError('Реквизиты покупателя не подтверждены исходным счётом.')
    if payload.get('buyer_kpp') and str(payload['buyer_kpp']) not in buyer_text:raise ValueError('КПП покупателя не подтверждён счётом.')
    if not payload.get('invoice_number') or not re.search(r'сч[её]т(?: на оплату)?\s*№\s*'+re.escape(payload['invoice_number'])+r'(?![\w-])',text,re.I):
        raise ValueError('Номер исходного счёта не подтверждён.')
    try:invoice_date=date.fromisoformat(payload.get('date'))
    except (TypeError,ValueError):raise ValueError('Проверьте дату исходного счёта.') from None
    months=('января','февраля','марта','апреля','мая','июня','июля','августа','сентября','октября','ноября','декабря')
    title=re.search(r'сч[её]т(?: на оплату)?\s*№\s*'+re.escape(payload['invoice_number'])+r'\s+от\s+([^\n]+)',text,re.I)
    if not title or not (invoice_date.strftime('%d.%m.%Y') in title.group(1) or re.search(
        r'[«"“]?0?'+str(invoice_date.day)+r'[»"”]?\s+'+months[invoice_date.month-1]+r'\s+'+str(invoice_date.year)+r'(?!\d)',title.group(1),re.I)):
        raise ValueError('Дата исходного счёта не совпадает с файлом.')
    return {'total':total,'vat':vat,'lines':canonical_lines,'amount_in_words':words}


def generate_invoice_specification(*,template_path,output_path,number,document_date,contract,invoice):
    """Fill only named source slots, cloning the source row; preserve all other ZIP members."""
    from copy import deepcopy
    from decimal import Decimal
    from zipfile import ZipFile
    from lxml import etree
    root=read_docx_xml(Path(template_path).read_bytes());body=root.find('w:body',_WORD_NS)
    def fill(node,value):
        texts=node.xpath('.//w:t',namespaces=_WORD_NS)
        if not texts:
            p=node if node.tag==qn('w:p') else node.find('w:p',_WORD_NS)
            run=etree.SubElement(p,qn('w:r'));texts=[etree.SubElement(run,qn('w:t'))]
        texts[0].text=str(value)
        for text in texts[1:]:text.text=''
    def money(value):return f'{Decimal(value):,.2f}'.replace(',',' ').replace('.',',')
    legal=contract['legal_fields']
    fill(body[0],'СПЕЦИФИКАЦИЯ № '+str(number))
    fill(body[1],f"к Договору № {contract['number']} от {format_document_date(contract['date'])} г.")
    fill(body[2].findall('w:tr/w:tc',_WORD_NS)[1],format_document_date(document_date)+' г.')
    buyer_texts=body[5].xpath('.//w:t',namespaces=_WORD_NS)
    buyer_texts[-1].text=legal['document_name'] or legal['full_name']
    table=body[7];source_row=table.findall('w:tr',_WORD_NS)[1];table.remove(source_row)
    for line in invoice['lines']:
        row=deepcopy(source_row)
        values=[line['row_no'],line['name'],format(Decimal(line['qty']).normalize(),'f').replace('.',','),line['unit'],money(line['unit_price_incl_vat']),money(line['unit_price_excl_vat']),money(line['vat_amount']),money(line['total_amount'])]
        for cell,value in zip(row.findall('w:tc',_WORD_NS),values):fill(cell,value)
        table.append(row)
    for row,value in zip(body[9].findall('w:tr',_WORD_NS),[invoice['total'],invoice['vat'],invoice['total']]):fill(row.findall('w:tc',_WORD_NS)[1],money(value))
    fill(body[10],'Всего к оплате прописью: '+invoice['amount_in_words'])
    signature=body[17].findall('w:tr',_WORD_NS)
    fill(signature[1].findall('w:tc',_WORD_NS)[1],legal['document_name'] or legal['full_name'])
    fill(signature[2].findall('w:tc',_WORD_NS)[1],legal['signer_position'])
    cell=signature[3].findall('w:tc',_WORD_NS)[1]
    paragraphs=cell.findall('w:p',_WORD_NS)
    fill(paragraphs[0],'______________ / '+format_person_short_name(legal['signer_name'])+' /')
    with ZipFile(template_path) as source,ZipFile(output_path,'w') as target:
        for entry in source.infolist():
            target.writestr(entry,etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) if entry.filename=='word/document.xml' else source.read(entry.filename))


def generated_contract_fields_match(content,document,legal):
    """A stored snapshot is usable only when current title, intro and buyer cells render its actual fields."""
    import re
    root=read_docx_xml(content);body=root.find('w:body',_WORD_NS)
    if len(body)<4:return False
    title=_normal_text(_xml_text(body[0]));intro=_normal_text(_xml_text(body[3]))
    if not re.search(r'договор поставки\s*№\s*'+re.escape(_normal_text(document['number']))+r'(?=\s|$)',title):return False
    if _normal_text(format_document_date(document['document_date'])) not in _normal_text(_xml_text(body[1])):return False
    if any(not legal.get(k) or _normal_text(legal[k]) not in intro for k in ('document_name','signer_position_genitive','signer_name_genitive')):return False
    tables=root.xpath('.//w:tbl',namespaces=_WORD_NS)
    buyers=[]
    for table in tables:
        rows=table.findall('w:tr',_WORD_NS)
        if not rows:continue
        cells=rows[0].findall('w:tc',_WORD_NS)
        if len(cells)==2 and 'покупатель' in _normal_text(_xml_text(cells[1])):
            buyers.append(_normal_text(' '.join(_xml_text(r.findall('w:tc',_WORD_NS)[1]) for r in rows if len(r.findall('w:tc',_WORD_NS))==2)))
    return any(all(legal.get(k) and _normal_text(legal[k]) in buyer for k in ('document_name','inn','signer_position','signer_short_name'))
        and (not legal.get('kpp') or _normal_text(legal['kpp']) in buyer) for buyer in buyers)
