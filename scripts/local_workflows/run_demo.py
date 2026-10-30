"""Task-owned synthetic local demo. No working config, database, storage or 1C."""
from pathlib import Path
import json
import os
import signal
import socket
import subprocess
import sys

SOURCE = Path(__file__).resolve().parents[2]
TASK = SOURCE / ".local-demo"

DATA = TASK / "candidate-data"
FRONTEND_PORT = 15100
BACKEND_PORT = 19100
sys.path.insert(0, str(SOURCE))
from tests.support.isolated_environment import create_test_environment


def environment():
    clean = create_test_environment(os.environ, DATA)
    # Keep extractor workspaces outside the source and current working directory.
    temp_root = Path("D:/codex-local-workflows-test-temp-20261030").resolve()
    if temp_root.is_relative_to(SOURCE.resolve()) or temp_root.is_relative_to(Path.cwd().resolve()):
        raise RuntimeError("The demo temporary directory must be outside source and cwd.")
    temp_root.mkdir(parents=True, exist_ok=True)
    clean.update(TEMP=str(temp_root), TMP=str(temp_root), TMPDIR=str(temp_root))
    path_value = os.environ.get("PATH", os.defpath)
    clean = {key: value for key, value in clean.items() if key.upper() != "PATH"}
    clean["PATH"] = path_value
    clean["SM_TECHNO_INITIAL_ADMIN_PASSWORD"] = "demo-admin-password"
    clean["SM_TECHNO_ALLOWED_ORIGINS"] = f"http://127.0.0.1:{FRONTEND_PORT},http://localhost:{FRONTEND_PORT}"
    # Reuse only a task-created key so restarting can read its own synthetic data.
    DATA.mkdir(exist_ok=True)
    key = DATA / "synthetic-key.txt"
    if not key.exists():
        key.write_text(clean["SM_TECHNO_CREDENTIAL_KEY"], encoding="ascii")
    clean["SM_TECHNO_CREDENTIAL_KEY"] = key.read_text(encoding="ascii")
    converter = TASK / "converter-image-id.txt"
    if converter.exists():
        clean["DOCUMENT_ASSISTANT_CONVERTER_IMAGE"] = converter.read_text().strip()
    return clean


def backend():
    os.environ.clear()
    os.environ.update(environment_from_parent)
    import stock_sync_api
    from tests.support.client_documents_e2e import seed_document_demo
    import uvicorn
    service = stock_sync_api.SERVICE
    user = service.db.get_user_by_username("demo")
    if user is None:
        uid = service.db.create_user(username="demo", password="demo-manager-password", role="user", full_name="Демонстрационный менеджер")
        client_id = seed_document_demo(service, uid)
        from stock_sync_web.crm_repository import CrmRepository
        repo = CrmRepository(service.db)
        tab = repo.ensure_work_tab_for_actor(actor_id=uid, owner_id=uid)
        repo.assign_client_for_actor(actor_id=uid, owner_id=uid, client_id=client_id, tab_id=tab["id"])
        (DATA / "records.json").write_text(json.dumps({"clientId": client_id, "userId": uid}), encoding="utf-8")
    uvicorn.run(stock_sync_api.app, host="127.0.0.1", port=BACKEND_PORT, log_level="warning")


environment_from_parent = environment()
if "--backend" in sys.argv:
    backend()
else:
    for port in (FRONTEND_PORT, BACKEND_PORT):
        with socket.socket() as probe:
            if probe.connect_ex(("127.0.0.1", port)) == 0:
                raise SystemExit(f"Port {port} is occupied; no existing process will be stopped.")
    fixtures = TASK / "demo05-fixtures"
    fixtures.mkdir(exist_ok=True)
    (fixtures / "articles.txt").write_text("00123 4\n", encoding="utf-8")
    (fixtures / "articles-partial.txt").write_text("00123 4\nUNKNOWN-ONLY-DEMO 1\n", encoding="utf-8")
    (fixtures / "requisites.txt").write_text(
        "СИНТЕТИЧЕСКИЕ ДАННЫЕ. НЕ РЕАЛЬНЫЙ КЛИЕНТ.\n"
        "Наименование для документов: ООО Синтетический клиент документов\n"
        "Полное наименование: ООО Синтетический клиент документов\n"
        "ИНН: 7701234560\nКПП: 770101001\nОГРН: 1027700000001\n"
        "Юридический адрес: г. Москва, ул. Синтетическая, д. 1\nФактический адрес: г. Москва, ул. Синтетическая, д. 1\n"
        "Банк: Синтетический банк\nБИК: 044525225\nРасчётный счёт: 40702810900000000001\nКорреспондентский счёт: 30101810400000000225\n"
        "Должность подписанта: Директор\nДолжность в родительном падеже: Директора\n"
        "ФИО подписанта: Иванов Иван Иванович\nФИО в родительном падеже: Иванова Ивана Ивановича\nИнициалы подписанта: И.И. Иванов\nОснование полномочий: Устава\n",
        encoding="utf-8",
    )
    from docx import Document
    invoice = Document()
    invoice.add_paragraph("Счет на оплату № СИНТ-1 от 30.10.2026")
    invoice.add_paragraph("Поставщик: ООО Синтетический продавец, ИНН 7804702073, КПП 780401001")
    invoice.add_paragraph("Покупатель: ООО Синтетический клиент документов, ИНН 7701234560, КПП 770101001")
    table = invoice.add_table(rows=0, cols=8)
    for values in [
        ["№", "Товар (Услуга)", "Кол-во", "Ед.", "Цена с НДС", "Цена без НДС", "Сумма НДС", "Сумма"],
        ["1", "00123 Синтетическая позиция", "1", "шт", "120.00", "100.00", "20.00", "120.00"],
    ]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for text in ["Итого: 120.00", "В том числе НДС: 20.00", "Всего к оплате: 120.00", "Всего наименований 1, на сумму 120.00 руб.", "Сто двадцать рублей 00 копеек"]:
        invoice.add_paragraph(text)
    invoice.save(fixtures / "invoice.docx")
    frontend_environment = {key: value for key, value in environment_from_parent.items() if not key.startswith(("SM_TECHNO_", "DOCUMENT_"))}
    frontend_environment["BACKEND_API_BASE_URL"] = f"http://127.0.0.1:{BACKEND_PORT}"
    logs = [(TASK / name).open("a", encoding="utf-8") for name in ("demo05-backend.log", "demo05-frontend.log")]
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    children = []
    try:
        children.append(subprocess.Popen([sys.executable, str(Path(__file__).resolve()), "--backend"], cwd=SOURCE, env=environment_from_parent, stdout=logs[0], stderr=subprocess.STDOUT, creationflags=flags))
        children.append(subprocess.Popen(["node", "node_modules/next/dist/bin/next", "dev", "--hostname", "127.0.0.1", "--port", str(FRONTEND_PORT)], cwd=SOURCE / "sm-techno-web", env=frontend_environment, stdout=logs[1], stderr=subprocess.STDOUT, creationflags=flags))
        (TASK / "demo05-processes.json").write_text(json.dumps({"launcherPid": os.getpid(), "backendPid": children[0].pid, "frontendPid": children[1].pid, "data": str(DATA), "url": f"http://127.0.0.1:{FRONTEND_PORT}/crm?documents=1"}, indent=2), encoding="utf-8")
        print(f"Synthetic local demo http://127.0.0.1:{FRONTEND_PORT}/crm?documents=1; login demo / demo-manager-password", flush=True)
        while all(child.poll() is None for child in children):
            try:
                children[0].wait(timeout=2)
            except subprocess.TimeoutExpired:
                pass
        raise SystemExit("A demo child exited; see task-owned logs.")
    finally:
        for child in children:
            if child.poll() is None:
                child.terminate()
        for log in logs:
            log.close()
