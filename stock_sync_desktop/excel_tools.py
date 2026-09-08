from __future__ import annotations

import math
import re
import unicodedata
from copy import copy
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, Side


TEMPLATE_COLUMNS = [
    "sku",
    "name",
    "print_name",
    "category_name",
    "group_name",
    "price",
    "warehouse_name",
    "rack",
    "cell",
    "quantity",
    "onec_key",
]

DISPLAY_COLUMNS = {
    "sku": "Артикул",
    "name": "Наименование",
    "print_name": "Наименование для печати",
    "category_name": "Категория",
    "group_name": "Группа",
    "price": "Цена",
    "warehouse_name": "Склад",
    "rack": "Стеллаж",
    "cell": "Ячейка",
    "quantity": "Остаток",
    "onec_key": "Ключ номенклатуры 1С",
    "unit_key": "Ключ единицы измерения 1С",
    "unit_name": "Единица измерения",
}

COLUMN_ALIASES = {
    "sku": {"sku", "артикул", "код", "article"},
    "name": {"name", "наименование", "название"},
    "print_name": {
        "print_name",
        "наименованиедляпечати",
        "печатноенаименование",
        "наимдляпечати",
    },
    "category_name": {"category_name", "категория", "категорияноменклатуры", "category"},
    "group_name": {"group_name", "группа", "входитвгруппу", "группаноменклатуры", "group"},
    "warehouse_name": {"warehouse_name", "склад", "warehouse", "storage"},
    "rack": {"rack", "shelf", "стеллаж", "стелаж"},
    "cell": {"cell", "ячейка", "ячейкахранения"},
    "onec_key": {
        "onec_key",
        "1c_key",
        "ключ1с",
        "ключноменклатуры1с",
        "ключноменклатуры",
        "ref_key",
    },
    "unit_key": {
        "unit_key",
        "единицаизмерения_key",
        "едизмключ",
        "ключединицыизмерения1с",
        "ключединицыизмерения",
    },
    "unit_name": {"unit_name", "единица", "едизм", "единицаизмерения", "unit"},
    "price": {"price", "цена"},
    "quantity": {"quantity", "остаток", "количество", "stock"},
}

CLIENT_PRICE_GRID_BORDER = Border(
    left=Side(style="thin", color="D7DFEA"),
    right=Side(style="thin", color="D7DFEA"),
    top=Side(style="thin", color="D7DFEA"),
    bottom=Side(style="thin", color="D7DFEA"),
)


def create_import_template(path: str | Path) -> Path:
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "stocks"
    sheet.append([DISPLAY_COLUMNS[column] for column in TEMPLATE_COLUMNS])
    sheet.append(
        [
            "SKU-001",
            "Пример товара",
            "Пример товара",
            "CATERPILLAR",
            "CAT",
            1500,
            "Основной склад",
            "Стеллаж 1",
            "A1",
            10,
            "",
        ]
    )

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    instructions = workbook.create_sheet("instructions")
    instructions["A1"] = "Как заполнять шаблон"
    instructions["A1"].font = Font(bold=True)
    rows = [
        "1. Артикул: внутренний артикул товара в приложении и то, что уйдет в поле 'Артикул' в 1С.",
        "2. Наименование: основное имя товара.",
        "3. Наименование для печати: если оставить пустым, приложение возьмет такое же значение, как в 'Наименование'.",
        "4. Категория: имя категории номенклатуры в 1С. Нужно для автосоздания новой карточки товара.",
        "5. Группа: папка номенклатуры в 1С. Если такой группы нет, приложение попробует создать ее автоматически.",
        "6. Цена: цена за единицу товара.",
        "7. Склад: одна строка = один товар на одном складе. Если один товар лежит на трех складах, у него должно быть три строки.",
        "8. Стеллаж и Ячейка: место хранения товара на указанном складе. Можно оставить пустыми.",
        "9. Остаток: остаток именно на указанном складе.",
        "10. Ключ номенклатуры 1С: Ref_Key товара из 1С, если товар уже связан с базой. Если не знаешь ключ, оставь ячейку пустой.",
        "11. Если колонки 'Склад' нет, приложение автоматически загрузит остаток в 'Основной склад'.",
        "12. При повторном импорте остаток по паре товар + склад считается итоговым значением, а не добавкой.",
        "13. Для загрузки только стеллажей и ячеек можно использовать лист 'Сопоставление' с колонками: Артикул, Стеллаж, Ячейка.",
    ]
    for index, value in enumerate(rows, start=2):
        instructions[f"A{index}"] = value

    workbook.save(target)
    return target


def export_stock_snapshot(path: str | Path, rows: list[dict[str, Any]]) -> Path:
    target = Path(path)
    frame = pd.DataFrame(rows)
    ordered = frame[[column for column in TEMPLATE_COLUMNS if column in frame.columns]]
    ordered = ordered.rename(columns=DISPLAY_COLUMNS)
    ordered.to_excel(target, index=False)
    return target


def export_client_price(
    path: str | Path,
    rows: list[dict[str, Any]],
    *,
    template_path: str | Path | None = None,
) -> Path:
    target = Path(path)
    if template_path is None:
        frame = pd.DataFrame(
            [
                {
                    "sku": row.get("sku") or "",
                    "brand": row.get("category_name") or "",
                    "name": row.get("print_name") or row.get("name") or "",
                    "quantity": int(round(_resolve_client_price_quantity(row))),
                    "price": round(float(row.get("price") or 0), 2),
                    "warehouse": _resolve_client_price_warehouse(row),
                }
                for row in rows
            ]
        )
        if frame.empty:
            frame = pd.DataFrame(
                columns=["sku", "brand", "name", "quantity", "price", "warehouse"]
            )
        ordered = frame[["sku", "brand", "name", "quantity", "price", "warehouse"]]
        ordered = ordered.rename(
            columns={
                "sku": "Артикул",
                "brand": "Бренд",
                "name": "Наименование",
                "quantity": "Наличие, шт.",
                "price": "Цена, ₽",
                "warehouse": "Склад",
            }
        )
        ordered.to_excel(target, index=False)
        return target

    workbook = load_workbook(template_path)
    sheet = workbook.active
    sheet.cell(row=1, column=6, value="Склад")
    existing_max_row = sheet.max_row
    data_start_row = 2
    last_data_row = max(data_start_row, len(rows) + 1)

    for row_idx in range(data_start_row, existing_max_row + 1):
        for column_idx in range(1, 7):
            sheet.cell(row=row_idx, column=column_idx).value = None

    style_source_cells = [sheet.cell(row=2, column=column_idx) for column_idx in range(1, 7)]
    style_source_height = sheet.row_dimensions[2].height

    for row_idx, row in enumerate(rows, start=data_start_row):
        if row_idx > existing_max_row:
            _apply_client_template_row_style(
                sheet,
                row_idx=row_idx,
                source_cells=style_source_cells,
                row_height=style_source_height,
            )

        sheet.cell(row=row_idx, column=1, value=row.get("sku") or "")
        sheet.cell(row=row_idx, column=2, value=row.get("category_name") or "")
        sheet.cell(row=row_idx, column=3, value=row.get("print_name") or row.get("name") or "")
        stock_cell = sheet.cell(
            row=row_idx,
            column=4,
            value=int(round(_resolve_client_price_quantity(row))),
        )
        price_cell = sheet.cell(
            row=row_idx,
            column=5,
            value=round(float(row.get("price") or 0), 2),
        )
        warehouse_cell = sheet.cell(
            row=row_idx,
            column=6,
            value=_resolve_client_price_warehouse(row),
        )
        stock_cell.number_format = '#,##0'
        price_cell.number_format = '#,##0.00'
        for styled_cell in (
            sheet.cell(row=row_idx, column=1),
            sheet.cell(row=row_idx, column=2),
            sheet.cell(row=row_idx, column=3),
            stock_cell,
            price_cell,
            warehouse_cell,
        ):
            styled_cell.border = copy(CLIENT_PRICE_GRID_BORDER)

    _refresh_client_template_formulas(sheet, last_data_row=last_data_row)
    workbook.save(target)
    return target


def _apply_client_template_row_style(
    sheet,
    *,
    row_idx: int,
    source_cells: list[Any],
    row_height: float | None,
) -> None:
    for column_idx, source_cell in enumerate(source_cells, start=1):
        target_cell = sheet.cell(row=row_idx, column=column_idx)
        target_cell._style = copy(source_cell._style)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format

    if row_height is not None:
        sheet.row_dimensions[row_idx].height = row_height


def _refresh_client_template_formulas(sheet, *, last_data_row: int) -> None:
    data_range = f"$A$1:$F${max(1, last_data_row)}"
    pattern = re.compile(r"\$A\$1:\$F\$\d+")
    for row in sheet.iter_rows(min_row=10, max_row=sheet.max_row, min_col=8, max_col=14):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = pattern.sub(data_range, cell.value)


def _resolve_client_price_quantity(row: dict[str, Any]) -> float:
    return float(row.get("row_quantity", row.get("quantity") or 0) or 0)


def _resolve_client_price_warehouse(row: dict[str, Any]) -> str:
    return str(
        row.get("row_warehouse_name")
        or row.get("warehouse_name")
        or row.get("warehouse_summary")
        or row.get("top_warehouse_name")
        or ""
    )


def _is_storage_location_file(file_path: Path) -> bool:
    target_name = file_path.name.lower()
    return "стелаж" in target_name or "стеллаж" in target_name


def _find_storage_mapping_sheet(sheet_names: list[str]) -> str | None:
    normalized_sheet_names = {name.casefold() for name in sheet_names}
    if "сопоставление" not in normalized_sheet_names:
        return None
    return next(
        (name for name in sheet_names if name.casefold() == "сопоставление"),
        "Сопоставление",
    )


def _select_excel_sheet(file_path: Path) -> tuple[str | int, bool]:
    if file_path.suffix.lower() == ".csv":
        return 0, False

    with pd.ExcelFile(file_path) as excel_file:
        sheet_names = [str(name) for name in excel_file.sheet_names]
    if _is_storage_location_file(file_path):
        mapping_sheet = _find_storage_mapping_sheet(sheet_names)
        if mapping_sheet:
            return mapping_sheet, True
        return 0, True
    return 0, False


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("-", "").replace(".", "")


def normalize_stock_sku(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value or "")).strip().casefold()


def _require_import_number(value: Any, *, label: str, source_row: int) -> float:
    number = _to_float(value)
    if not math.isfinite(number) or number < 0:
        raise ValueError(
            f"Строка {source_row}: {label} должно быть конечным неотрицательным числом."
        )
    return number


def read_stock_import_bundle(path: str | Path) -> dict[str, list[dict[str, Any]]]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        read_error: UnicodeDecodeError | None = None
        for encoding in ("utf-8", "utf-8-sig", "cp1251"):
            try:
                frame = pd.read_csv(file_path, encoding=encoding, keep_default_na=False)
                break
            except UnicodeDecodeError as exc:
                read_error = exc
        else:
            raise ValueError("Не удалось прочитать CSV: используйте UTF-8 или Windows-1251.") from read_error
        is_storage_mapping = False
    else:
        sheet_name, is_storage_mapping = _select_excel_sheet(file_path)
        if is_storage_mapping and sheet_name == 0:
            location_rows = _read_storage_location_layout(file_path)
            if location_rows:
                return {"stock_rows": [], "location_rows": location_rows}
        formula_book = load_workbook(file_path, read_only=True, data_only=False)
        try:
            formula_sheet = (
                formula_book.worksheets[sheet_name]
                if isinstance(sheet_name, int)
                else formula_book[sheet_name]
            )
            if any(
                cell.data_type == "f"
                for row in formula_sheet.iter_rows()
                for cell in row
            ):
                raise ValueError(
                    "Таблица импорта содержит формулы. Сохраните вычисленные значения перед импортом."
                )
        finally:
            formula_book.close()
        frame = pd.read_excel(file_path, sheet_name=sheet_name, keep_default_na=False)

    normalized_map: dict[str, str] = {}
    for original in frame.columns:
        normalized = _normalize_header(original)
        for canonical, aliases in COLUMN_ALIASES.items():
            if normalized in aliases:
                normalized_map[original] = canonical
                break

    frame = frame.rename(columns=normalized_map)
    if is_storage_mapping or (
        "sku" in frame.columns
        and ("rack" in frame.columns or "cell" in frame.columns)
        and "name" not in frame.columns
    ):
        location_rows = _read_storage_location_rows(frame)
        if not location_rows:
            raise ValueError("Импортируемая таблица пустая или не содержит валидных строк.")
        return {"stock_rows": [], "location_rows": location_rows}

    missing = [column for column in ("name", "quantity") if column not in frame.columns]
    if missing:
        raise ValueError(
            "В таблице не хватает обязательных колонок: " + ", ".join(missing)
        )

    has_rack_column = "rack" in frame.columns
    has_cell_column = "cell" in frame.columns
    cleaned_rows: list[dict[str, Any]] = []
    seen_stock_rows: set[tuple[str, str]] = set()
    for source_index, row in frame.iterrows():
        source_row = int(source_index) + 2
        name = str(row.get("name") or "").strip()
        if not name:
            continue

        sku = _clean_string(row.get("sku"))
        normalized_sku = normalize_stock_sku(sku)
        if not normalized_sku:
            raise ValueError(f"Строка {source_row}: артикул не может быть пустым.")
        warehouse_name = _clean_string(row.get("warehouse_name")) or "Основной склад"
        duplicate_key = (normalized_sku, normalize_stock_sku(warehouse_name))
        if duplicate_key in seen_stock_rows:
            raise ValueError(
                f"Строка {source_row}: дубликат артикула на одном складе в импортируемой таблице."
            )
        seen_stock_rows.add(duplicate_key)

        cleaned_row = {
            "sku": sku,
            "name": name,
            "print_name": _clean_string(row.get("print_name")) or name,
            "category_name": _clean_string(row.get("category_name")),
            "group_name": _clean_string(row.get("group_name")),
            "price": _require_import_number(
                row.get("price"), label="Цена", source_row=source_row
            ),
            "warehouse_name": warehouse_name,
            "quantity": _require_import_number(
                row.get("quantity"), label="Остаток", source_row=source_row
            ),
            "onec_key": _clean_string(row.get("onec_key")),
            "unit_key": _clean_string(row.get("unit_key")),
            "unit_name": _clean_string(row.get("unit_name")),
            "source_row": source_row,
        }
        if has_rack_column:
            cleaned_row["rack"] = _clean_string(row.get("rack"))
        if has_cell_column:
            cleaned_row["cell"] = _clean_string(row.get("cell"))
        cleaned_rows.append(cleaned_row)

    if not cleaned_rows:
        raise ValueError("Импортируемая таблица пустая или не содержит валидных строк.")
    return {"stock_rows": cleaned_rows, "location_rows": []}


def read_stock_import(path: str | Path) -> list[dict[str, Any]]:
    return read_stock_import_bundle(path)["stock_rows"]


def _read_storage_location_rows(frame: pd.DataFrame) -> list[dict[str, Any]]:
    if "sku" not in frame.columns:
        raise ValueError("В таблице сопоставления не хватает колонки: Артикул")
    if "rack" not in frame.columns and "cell" not in frame.columns:
        raise ValueError("В таблице сопоставления не хватает колонок: Стеллаж, Ячейка")

    cleaned_rows: list[dict[str, Any]] = []
    for _, row in frame.iterrows():
        sku = _clean_string(row.get("sku"))
        if not sku:
            continue
        cleaned_rows.append(
            {
                "sku": sku,
                "warehouse_name": _clean_string(row.get("warehouse_name")) or "Санкт-Петербург",
                "rack": _clean_string(row.get("rack")),
                "cell": _clean_string(row.get("cell")),
                "quantity": _to_float(row.get("quantity")),
            }
        )
    return cleaned_rows


def _read_storage_location_layout(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rack: str | None = None
    cell: str | None = None
    rows: list[dict[str, Any]] = []

    try:
        for raw_row in sheet.iter_rows(values_only=True):
            first_value = raw_row[0] if len(raw_row) > 0 else None
            second_value = raw_row[1] if len(raw_row) > 1 else None
            text = _clean_string(first_value)
            if not text:
                continue

            second_text = _clean_string(second_value)
            if _is_rack_marker(text) and second_text is None:
                rack = text
                cell = None
                continue

            if _is_cell_marker(text) and second_text is None:
                cell = _normalize_cell_label(text)
                continue

            third_text = _clean_string(raw_row[2] if len(raw_row) > 2 else None)
            fourth_text = _clean_string(raw_row[3] if len(raw_row) > 3 else None)
            if _is_headerless_storage_location_row(text, second_text, third_text, fourth_text):
                rows.append(
                    {
                        "sku": text,
                        "warehouse_name": "Санкт-Петербург",
                        "rack": fourth_text or rack,
                        "cell": _normalize_cell_label(third_text) if third_text else cell,
                        "quantity": _to_float(second_value),
                    }
                )
                continue

            if second_text is None:
                continue

            rows.append(
                {
                    "sku": text,
                    "warehouse_name": "Санкт-Петербург",
                    "rack": rack,
                    "cell": cell,
                    "quantity": _to_float(second_value),
                }
            )
    finally:
        workbook.close()

    return rows


def _is_headerless_storage_location_row(
    sku: str,
    quantity: str | None,
    cell: str | None,
    rack: str | None,
) -> bool:
    if not quantity or (not cell and not rack):
        return False
    normalized_sku = _normalize_header(sku)
    if normalized_sku in COLUMN_ALIASES["sku"]:
        return False
    normalized_cell = _normalize_header(cell)
    normalized_rack = _normalize_header(rack)
    if normalized_cell in COLUMN_ALIASES["cell"] or normalized_rack in COLUMN_ALIASES["rack"]:
        return False
    return True


def _is_rack_marker(value: str) -> bool:
    normalized = value.strip().casefold()
    return normalized.startswith("стелаж") or normalized.startswith("стеллаж")


def _is_cell_marker(value: str) -> bool:
    normalized = _normalize_cell_label(value)
    return bool(re.fullmatch(r"[A-ZА-Я]\d+", normalized, flags=re.IGNORECASE))


def _normalize_cell_label(value: str) -> str:
    return (
        value.strip()
        .upper()
        .replace("А", "A")
        .replace("В", "B")
        .replace("С", "C")
        .replace("Е", "E")
        .replace("Н", "H")
        .replace("К", "K")
        .replace("М", "M")
        .replace("О", "O")
        .replace("Р", "P")
        .replace("Т", "T")
        .replace("Х", "X")
    )


def _clean_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, float) and pd.isna(value):
        return 0.0
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return 0.0
    return float(text)
